"""Generated structural labels must be numbering/fields, never typed prefixes."""

import sys
import json
import subprocess
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))
from word_format.api import apply_plan, build_plan, verify_plan
from word_format.inspect import DocumentIndex
from word_format.numbering import read_heading_scheme, verify_heading_scheme
from word_format.numbering_audit import audit_numbering


SCHEME = {"levels": [
    {"format": "chineseCountingThousand", "text": "第%1章", "start": 1, "suffix": "space"},
    {"format": "decimal", "text": "%1.%2", "start": 1, "suffix": "space"},
    {"format": "decimal", "text": "%1.%2.%3", "start": 1, "suffix": "space"},
]}


class NumberingGuardTests(unittest.TestCase):
    def test_pipeline_blocks_typed_caption_and_cli_protects_source(self):
        root = Path(__file__).resolve().parents[2]
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            source, output = folder / "in.docx", folder / "out.docx"
            template = folder / "template.json"
            doc = Document()
            doc.add_paragraph("第一章 研究背景", style="Heading 1")
            doc.add_paragraph("图1 系统架构")
            doc.save(source)
            original = source.read_bytes()
            template.write_text(json.dumps({"schema_version": "1.0", "kind": "word-format-template",
                "id": "test", "version": "1", "origin": "user", "sources": [], "issues": [],
                "rules": [{"id": "numbering", "state": "set", "action": "heading.scheme.configure",
                           "target": {"type": "document"}, "params": SCHEME,
                           "evidence": [{"anchor": "test"}]}]}), encoding="utf-8")
            result = subprocess.run([sys.executable, str(root / "scripts/run_format_pipeline.py"),
                                     str(source), str(output), "--template", str(template)],
                                    capture_output=True)
            self.assertNotEqual(result.returncode, 0)
            report = json.loads((folder / "out-reports/numbering.json").read_text(encoding="utf-8"))
            self.assertEqual(report["status"], "failed")
            self.assertEqual(report["issues"][0]["kind"], "caption")
            result = subprocess.run([sys.executable, str(root / "scripts/check_numbering.py"),
                                     str(source), "--output", str(source)], capture_output=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(source.read_bytes(), original)

    def test_compact_multirun_prefixes_convert_without_destroying_title_format(self):
        doc = Document()
        for level, prefix, title in [(1, "第一章", "研究背景"), (2, "1.1", "研究目标"),
                                      (3, "1.1.1", "实验范围"), (1, "第二章", "方法")]:
            p = doc.add_paragraph(style=f"Heading {level}")
            p.add_run(prefix)
            p.add_run(title).bold = True
        with tempfile.TemporaryDirectory() as folder:
            source, output = Path(folder) / "in.docx", Path(folder) / "out.docx"
            doc.save(source)
            plan = build_plan(source, output, {"operations": [{"action": "heading.scheme.configure",
                "target": {"type": "document"}, "params": SCHEME}]})
            report = apply_plan(plan)
            self.assertEqual(verify_plan(plan, report)["status"], "passed")
            result = Document(output)
            self.assertEqual([p.text for p in result.paragraphs], ["研究背景", "研究目标", "实验范围", "方法"])
            self.assertTrue(all(p.runs[-1].bold for p in result.paragraphs))
            self.assertEqual(audit_numbering(result)["status"], "passed")
            result.paragraphs[1].runs[0].text = "1.1"
            actual = read_heading_scheme(DocumentIndex(result), None, None)
            self.assertFalse(verify_heading_scheme(actual, SCHEME, None))
            self.assertEqual(audit_numbering(result)["status"], "failed")

    def test_literal_labels_in_tables_headers_and_body_are_rejected(self):
        doc = Document()
        doc.add_paragraph("第一章 研究背景", style="Heading 1")
        doc.add_paragraph("（1）列表项")
        doc.add_paragraph("引用文献[1]。")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "图1-1 系统架构"
        doc.sections[0].footer.paragraphs[0].text = "1"
        result = audit_numbering(doc)
        self.assertEqual(result["status"], "failed")
        self.assertEqual({r["kind"] for r in result["issues"]},
                         {"heading", "list", "reference_or_citation", "caption", "page"})

    def test_fields_caches_and_real_data_are_not_typed_numbers(self):
        doc = Document()
        doc.add_paragraph("2026年，准确率86.4%，1.5倍行距，型号X123。")
        p = doc.add_paragraph("图")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), " SEQ Figure ")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "1"
        run.append(text)
        field.append(run)
        p._p.append(field)
        p.add_run(" 系统架构")
        p = doc.add_paragraph()
        for kind in ("begin", "separate", "end"):
            if kind == "end":
                p.add_run("[1]")
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
            p.add_run()._r.append(node)
        self.assertEqual(audit_numbering(doc)["status"], "passed")
        # A field elsewhere must not hide a typed caption prefix.
        p.add_run(" 图2 系统")
        self.assertEqual(audit_numbering(doc)["status"], "failed")


if __name__ == "__main__":
    unittest.main()
