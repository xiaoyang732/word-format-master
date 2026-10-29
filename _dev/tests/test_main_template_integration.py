"""Main-branch fidelity fixes must also hold through the template engine."""

import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))

from word_format import apply_plan, audit_document, build_plan, verify_plan
from word_format.contracts import FormatError
from word_format.inspect import DocumentIndex
from word_format.profiles import template_profile
from word_format.structures import append_field


class MainTemplateIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / "source.docx"
        self.output = Path(self.temp.name) / "output.docx"

    def test_profile_keeps_heading_levels_and_abstract_title_identity(self):
        doc = Document()
        doc.add_paragraph("题目：研究课题")
        doc.add_paragraph("摘要前的论文题名")
        doc.add_paragraph("摘 要")
        doc.add_paragraph("摘要文字。")
        doc.add_paragraph("关键词：系统；测试")
        for text, level, size in (("第一章 绪论", 1, 19), ("1.1 研究背景", 2, 15)):
            p = doc.add_heading(text, level)
            p.runs[0].font.size = Pt(size)
        doc.save(self.source)
        index = DocumentIndex(doc)
        self.assertEqual(index.records["p1"]["semantic_role"], "abstract_title")
        profile = template_profile(self.source)
        sizes = {
            rule["target"].get("level"): rule["params"]["value"]
            for rule in profile["rules"]
            if rule["target"].get("semantic_role") == "chapter_heading"
            and rule["action"] == "font.size.set"
        }
        self.assertEqual(sizes, {1: 19, 2: 15})

    def test_keyword_content_alias_keeps_main_value_and_language_contract(self):
        doc = Document()
        doc.add_paragraph("关键词：中文")
        doc.add_paragraph("Key words: English")
        doc.save(self.source)
        for alias in ("value", "content"):
            output = self.output.with_name(alias + ".docx")
            plan = build_plan(
                self.source,
                output,
                {
                    "operations": [
                        {
                            "action": "font.size.set",
                            "target": {
                                "type": "text_range",
                                "role": "keywords",
                                "language": "en",
                                "segment": alias,
                                "all": True,
                            },
                            "params": {"value": 21},
                        }
                    ]
                },
            )
            report = apply_plan(plan)
            self.assertEqual(verify_plan(plan, report)["status"], "passed")
            out = Document(output)
            self.assertEqual(out.paragraphs[1].runs[-1].font.size, Pt(21))
            self.assertIsNone(out.paragraphs[1].runs[0].font.size)
            self.assertIsNone(out.paragraphs[0].runs[0].font.size)

    def test_toc_new_and_legacy_appearance_are_independently_audited(self):
        doc = Document()
        doc.add_heading("第一章 绪论", 1)
        doc.save(self.source)
        for title_key in ("title_style", "title_format"):
            request = {
                "operations": [
                    {
                        "action": "toc.configure",
                        "target": {"type": "document"},
                        "params": {
                            "enabled": True,
                            title_key: {"font_size_pt": 18},
                            "field_format": {"font_size_pt": 11},
                            "break_format": {"space_after_pt": 3},
                        },
                    }
                ]
            }
            plan = build_plan(self.source, self.output, request)
            report = apply_plan(plan, overwrite=self.output.exists())
            self.assertEqual(verify_plan(plan, report)["status"], "passed")
            for style in ("WFM TOC Heading", "WFM TOC Field", "WFM TOC Break"):
                out = Document(self.output)
                p = next(p for p in out.paragraphs if p.style.name == style)
                if style.endswith("Break"):
                    p.paragraph_format.space_after = Pt(27)
                else:
                    p.runs[0].font.size = Pt(27)
                altered = self.output.with_name("altered.docx")
                out.save(altered)
                self.assertEqual(audit_document(altered, request)["status"], "failed", style)

    def test_conflicting_toc_aliases_do_not_silently_choose_one(self):
        Document().save(self.source)
        with self.assertRaises(FormatError):
            build_plan(
                self.source,
                self.output,
                {
                    "spec": {
                        "table_of_contents": {
                            "enabled": True,
                            "title_style": {"font_size_pt": 14},
                            "title_format": {"font_size_pt": 18},
                        }
                    }
                },
            )

    def test_body_first_page_override_preserves_footer_text_and_is_checked(self):
        doc = Document()
        doc.add_paragraph("封面")
        doc.sections[0].first_page_footer.paragraphs[0].add_run("保留的页脚内容")
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        doc.add_heading("第一章 绪论", 1)
        doc.add_paragraph("正文。")
        doc.save(self.source)
        request = {"spec": {"page_numbers": {"enabled": True, "show_on_body_first_page": False}}}
        plan = build_plan(self.source, self.output, request)
        report = apply_plan(plan)
        self.assertEqual(verify_plan(plan, report)["status"], "passed")
        out = Document(self.output)
        footer = out.sections[1].first_page_footer
        self.assertIn("保留的页脚内容", "".join(p.text for p in footer.paragraphs))
        append_field(footer.add_paragraph(), "PAGE")
        altered = self.output.with_name("incorrect-first-page.docx")
        out.save(altered)
        self.assertEqual(audit_document(altered, request)["status"], "failed")


if __name__ == "__main__":
    unittest.main()
