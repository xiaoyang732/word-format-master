"""Independent regressions for the cross-checked October review findings."""

import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
from json_output import write_json
from word_format import apply_plan, build_plan, verify_plan
from word_format.contracts import FormatError
from word_format.registry import REGISTRY
from word_format.properties import xml_value
from word_format import audit_document
from word_format import inspect_document
from word_format.inspect import DocumentIndex


class JsonSafetyTests(unittest.TestCase):
    def test_cli_collisions_preserve_inputs_and_do_not_create_outputs(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            source = folder / "source.docx"
            doc = Document()
            doc.add_paragraph("Source must survive")
            doc.save(source)
            original = source.read_bytes()
            alias = folder / "alias.json"
            os.link(source, alias)
            spec = folder / "spec.json"
            spec_text = json.dumps({"schema_version": "1.0", "body": {"font_size_pt": 12}})
            spec.write_text(spec_text, encoding="utf-8")
            output = folder / "result.docx"
            for report in (source, alias, folder / "sub" / ".." / "source.docx"):
                for script, args in (
                    ("analyze_docx.py", [source, "--output", report]),
                    ("analyze_citations.py", [source, "--output", report]),
                    ("apply_spec.py", [source, output, "--spec", spec, "--report", report]),
                    ("format_cli.py", ["inspect", source, "--output", report]),
                ):
                    with self.subTest(script=script, report=report):
                        result = subprocess.run([sys.executable, str(ROOT / "scripts" / script), *map(str, args)],
                                                capture_output=True, timeout=30)
                        self.assertNotEqual(result.returncode, 0)
                        self.assertNotIn(b"unrecognized arguments", result.stderr)
                        self.assertEqual(hashlib.sha256(source.read_bytes()).digest(), hashlib.sha256(original).digest())
                        self.assertTrue(zipfile.is_zipfile(source))
                        self.assertFalse(output.exists())
            requirements = folder / "requirements.txt"
            requirements.write_text("正文宋体", encoding="utf-8")
            result = subprocess.run([sys.executable, str(ROOT / "scripts/parse_requirements.py"),
                                     str(requirements), "--output", str(requirements)], capture_output=True, timeout=30)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(requirements.read_text(encoding="utf-8"), "正文宋体")
            result = subprocess.run([sys.executable, str(ROOT / "scripts/apply_spec.py"), str(source), str(output),
                                     "--spec", str(spec), "--report", str(spec)], capture_output=True, timeout=30)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(spec.read_text(encoding="utf-8"), spec_text)

    def test_atomic_json_write_preserves_previous_report_on_failure(self):
        with tempfile.TemporaryDirectory() as folder:
            report = Path(folder) / "report.json"
            report.write_text("old", encoding="utf-8")
            with patch("json_output.os.replace", side_effect=OSError("publish failed")):
                with self.assertRaises(OSError):
                    write_json(report, {"new": "中文"})
            self.assertEqual(report.read_text(encoding="utf-8"), "old")
            self.assertEqual(list(Path(folder).iterdir()), [report])
            write_json(report, {"new": "中文"})
            self.assertEqual(json.loads(report.read_text(encoding="utf-8")), {"new": "中文"})


class ReferenceSafetyTests(unittest.TestCase):
    def test_references_stop_at_plain_module_titles_and_protect_their_content(self):
        with tempfile.TemporaryDirectory() as folder:
            source, output = Path(folder) / "source.docx", Path(folder) / "output.docx"
            doc = Document()
            for text in ("参考文献", "[1] Reference entry", "致谢", "Thanks", "附录", "Appendix contents"):
                doc.add_paragraph(text)
            protected = [xml_value(p._p) for p in doc.paragraphs[2:]]
            doc.save(source)
            plan = build_plan(source, output, {"operations": [{"action": "reference.numbering.convert",
                "target": {"type": "document"}, "params": {"numbering_mode": "word-numbering"}}]})
            report = apply_plan(plan)
            self.assertEqual(verify_plan(plan, report)["status"], "passed")
            out = Document(output)
            self.assertEqual(out.paragraphs[1].style.name, "Bibliography")
            self.assertEqual(out.paragraphs[1].text, "Reference entry")
            self.assertEqual([xml_value(p._p) for p in out.paragraphs[2:]], protected)
            output.unlink()
            cap = REGISTRY["reference.numbering.convert"]
            def bad(index, ref, params):
                cap.handler(index, ref, params)
                index.document.paragraphs[2].style = "Bibliography"
            with patch.dict(REGISTRY, {cap.action: type(cap)(**{**cap.__dict__, "handler": bad})}):
                with self.assertRaises(FormatError):
                    apply_plan(plan)
            self.assertFalse(output.exists())


class PropertySemanticsTests(unittest.TestCase):
    def test_spacing_and_both_latin_font_channels_cannot_falsely_pass(self):
        with tempfile.TemporaryDirectory() as folder:
            source, output = Path(folder) / "source.docx", Path(folder) / "output.docx"
            doc = Document()
            style = doc.styles.add_style("Inherited", WD_STYLE_TYPE.PARAGRAPH)
            spacing = OxmlElement("w:spacing")
            spacing.set(qn("w:beforeAutospacing"), "1")
            style.element.get_or_add_pPr().append(spacing)
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:hAnsi"), "Courier New")
            style.element.get_or_add_rPr().append(fonts)
            for inherited in (False, True):
                p = doc.add_paragraph("ASCII and café", style=style if inherited else "Normal")
                spacing = OxmlElement("w:spacing")
                spacing.set(qn("w:before"), "0")
                if not inherited:
                    spacing.set(qn("w:beforeAutospacing"), "1")
                spacing.set(qn("w:afterLines"), "100")
                p._p.get_or_add_pPr().append(spacing)
                fonts = OxmlElement("w:rFonts")
                fonts.set(qn("w:ascii"), "Arial")
                if not inherited:
                    fonts.set(qn("w:hAnsi"), "Courier New")
                p.runs[0]._r.get_or_add_rPr().append(fonts)
            doc.save(source)
            request = {"operations": [{"action": action, "target": {"type": "paragraph", "all": True},
                       "params": {"value": value}} for action, value in
                       (("font.latin.set", "Arial"), ("paragraph.spacing_before.set", 0), ("paragraph.spacing_after.set", 0))]}
            self.assertEqual(audit_document(source, request)["status"], "failed")
            plan = build_plan(source, output, request)
            report = apply_plan(plan)
            self.assertTrue(all(check["changed"] for check in report["checks"][:2]))
            self.assertEqual(verify_plan(plan, report)["status"], "passed")
            out = Document(output)
            for p in out.paragraphs:
                fonts = p.runs[0]._r.find("./" + qn("w:rPr") + "/" + qn("w:rFonts"))
                self.assertEqual(fonts.get(qn("w:ascii")), "Arial")
                self.assertEqual(fonts.get(qn("w:hAnsi")), "Arial")
                spacing = p._p.pPr.find(qn("w:spacing"))
                for side in ("before", "after"):
                    self.assertEqual(spacing.get(qn("w:" + side)), "0")
                    self.assertEqual(spacing.get(qn("w:" + side + "Autospacing")), "0")
                    self.assertIsNone(spacing.get(qn("w:" + side + "Lines")))

    def test_orientation_swaps_existing_geometry_across_sections_and_round_trips(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "source.docx"
            doc = Document()
            doc.add_paragraph("Geometry")
            doc.sections[0].page_width, doc.sections[0].page_height = Mm(210), Mm(297)
            doc.add_section(WD_SECTION_START.NEW_PAGE)
            doc.save(source)
            for direction in ("landscape", "portrait"):
                output = Path(folder) / (direction + ".docx")
                plan = build_plan(source, output, {"operations": [{"action": "page.orientation.set",
                    "target": {"type": "section", "all": True}, "params": {"value": direction}}]})
                report = apply_plan(plan)
                self.assertEqual(verify_plan(plan, report)["status"], "passed")
                for section in Document(output).sections:
                    self.assertEqual(section.page_width > section.page_height, direction == "landscape")
                    self.assertAlmostEqual(max(section.page_width.mm, section.page_height.mm), 297, places=1)
                source = output


class ModuleContractTests(unittest.TestCase):
    def test_plain_modules_custom_outlines_occurrences_and_arabic_space(self):
        with tempfile.TemporaryDirectory() as folder:
            source, output = Path(folder) / "source.docx", Path(folder) / "output.docx"
            doc = Document()
            style = doc.styles.add_style("Custom chapter", WD_STYLE_TYPE.PARAGRAPH)
            outline = OxmlElement("w:outlineLvl")
            outline.set(qn("w:val"), "0")
            style.element.get_or_add_pPr().append(outline)
            for text in ("论文标题", "摘要", "中文摘要正文", "关键词：文档", "Abstract", "English abstract", "Key words: document"):
                doc.add_paragraph(text)
            for i in range(3):
                doc.add_paragraph("Custom chapter " + str(i), style=style)
                doc.add_paragraph("Body")
            index = DocumentIndex(doc)
            chapters = [r for r in index.records.values() if r.get("heading_level") == 1]
            self.assertEqual([r["module_occurrence"] for r in chapters], [1, 2, 3])
            self.assertEqual(index.records["p0"]["role"], "title")
            self.assertEqual(index.records["p3"]["role"], "keywords")
            doc.save(source)
            spec = {"title": {"font_size_pt": 22, "alignment": "center"},
                "body": {"font_size_pt": 12, "first_line_indent_chars": 2},
                "headings": [{"level": 1, "font_size_pt": 16, "bold": True, "first_line_indent_chars": 0}],
                "keywords": {"font_size_pt": 11, "first_line_indent_chars": 0},
                "lists": {"headings": {"level1": "arabic-space"}}}
            plan = build_plan(source, output, {"spec": spec})
            report = apply_plan(plan)
            self.assertEqual(verify_plan(plan, report)["status"], "passed")
            out = Document(output)
            self.assertEqual(out.paragraphs[0].runs[0].font.size.pt, 22)
            for i in (1, 4):
                self.assertEqual(out.paragraphs[i].runs[0].font.size.pt, 16)
                self.assertTrue(out.paragraphs[i].runs[0].bold)
                self.assertIsNone(out.paragraphs[i]._p.find("./" + qn("w:pPr") + "/" + qn("w:numPr")))
            for i in (3, 6):
                self.assertEqual(out.paragraphs[i].runs[0].font.size.pt, 11)
            self.assertEqual(sum(p.style.name == "Heading 1" for p in out.paragraphs), 3)
            markers = [n.get(qn("w:val")) for n in out.part.numbering_part.element.iter(qn("w:lvlText"))]
            self.assertIn("%1", markers)
            inspected = inspect_document(output)
            fmt = next(r["format"] for r in inspected["objects"] if r["kind"] == "paragraph")
            expected = {c.property for c in REGISTRY.values() if not c.structural and "paragraph" in c.targets}
            self.assertEqual(set(fmt), expected)


if __name__ == "__main__":
    unittest.main()
