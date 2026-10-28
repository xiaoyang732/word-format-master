"""Regression fixtures for template identity, requirement scope and real TOC results."""

import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))

from word_format import apply_plan, build_plan, verify_plan
from word_format.contracts import FormatError, sha256
from word_format.inspect import DocumentIndex
from word_format.profiles import profile_plan, requirements_profile, template_profile
from word_format.requirement_rules import parse_blocks
from word_format.structures import append_field
from word_format.toc import imported_fields, result_status
from verify_output import verify_structure


class TemplateFidelityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "source.docx"
        self.output = self.root / "output.docx"

    def fixture(self, path, target=False):
        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(12)
        normal.font.bold = False
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

        def add(text, size=12, bold=False, font="宋体"):
            p = doc.add_paragraph(text)
            r = p.runs[0]
            r.font.size = Pt(13 if target else size)
            r.bold = not bold if target else bold
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体" if target else font)
            return p

        add("学校课程论文", 36, True, "黑体")
        add("题 目：" + ("目标题目" if target else "样例题目"), 14, True, "黑体")
        add("学生姓名：" + ("目标作者" if target else "样例作者"), 14, True, "黑体")
        add("二级学院：计算机学院", 14, True, "黑体")
        add("专 业：软件工程", 14, True, "黑体")
        add("指导教师：教师", 14, True, "黑体")
        add("二〇二六年五月", 18, True)
        add("目标系统研究" if target else "样例系统研究", 16, False, "黑体")
        add("摘  要", 16, False, "黑体")
        add("目标摘要内容" if target else "样例摘要内容")
        p = add("关键词：", 12, False, "黑体")
        p.add_run("系统；测试").bold = False
        add("Target Research" if target else "Sample Research", 16, True)
        add("Abstract", 16, True)
        add("Target abstract text." if target else "Sample abstract text.")
        p = add("Key words:", 12, True)
        p.add_run("system; test").bold = False
        p = doc.add_heading("第1章 绪论", 1)
        p.runs[0].font.name = "Times New Roman"
        p.runs[0]._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
        p.runs[0].font.size = Pt(15)
        p.runs[0].bold = True
        add("目标正文" if target else "样例正文")
        doc.save(path)

    def test_template_semantic_profile_round_trip(self):
        sample = self.root / "sample.docx"
        self.fixture(sample)
        self.fixture(self.source, target=True)
        source_hash, sample_hash = sha256(self.source), sha256(sample)
        profile = template_profile(sample)
        for identity, size in (("cover_title", 36), ("cover_student", 14), ("cover_college", 14), ("cover_date", 18)):
            self.assertTrue(any(r["target"].get("semantic_role") == identity and r["action"] == "font.size.set" and r["params"]["value"] == size for r in profile["rules"]))
        decisions = {i["id"]: {"decision": "ignore", "reason": "Synthetic fixture: leave unspecified properties and empty headers unchanged"} for i in profile["issues"]}
        plan = profile_plan(self.source, self.output, profile, resolutions=decisions)
        report = apply_plan(plan)
        self.assertEqual(verify_plan(plan, report)["status"], "passed")
        index = DocumentIndex(Document(self.output))
        for r in index.records.values():
            if r.get("semantic_role") == "abstract_heading":
                self.assertEqual(index.font_values(r["id"], "font_size_pt"), [16])
                self.assertEqual(index.font_values(r["id"], "bold"), [r["language"] == "en"])
            if r.get("semantic_role") == "keywords":
                from word_format.spec import keyword_text_ranges
                label = keyword_text_ranges(r["text"])["label"]
                self.assertEqual(index.font_values(r["id"], "bold", *label), [r["language"] == "en"])
        self.assertEqual(sha256(self.source), source_hash)
        self.assertEqual(sha256(sample), sample_hash)
        self.assertIn("target_properties", report["profile_context"]["coverage"])
        self.assertNotIn("样例", "".join(p.text for p in index.document.paragraphs))

    def test_empty_abstract_is_evidence_to_review_not_a_body_default(self):
        doc = Document()
        doc.add_paragraph("摘 要")
        p = doc.add_paragraph()
        pr = p._p.get_or_add_pPr()
        rp = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "32")
        rp.append(sz)
        pr.append(rp)
        doc.add_paragraph("关键词：")
        doc.save(self.source)
        profile = template_profile(self.source)
        self.assertTrue(any("Empty abstract" in i["reason"] for i in profile["issues"]))
        self.assertFalse(any(r["target"].get("semantic_role") == "abstract_body" for r in profile["rules"]))

    def test_overflowing_template_indent_requires_decision(self):
        from docx.shared import Mm

        doc = Document()
        p = doc.add_paragraph("学校课程论文")
        p.paragraph_format.first_line_indent = Mm(220)
        doc.add_paragraph("摘 要")
        doc.save(self.source)
        profile = template_profile(self.source)
        self.assertTrue(any("exceeds usable page width" in i["reason"] for i in profile["issues"]))
        self.assertFalse(any(r["target"].get("semantic_role") == "cover_title" and r["action"] == "paragraph.first_line_indent.set" for r in profile["rules"]))

    def test_paragraph_spacing_in_lines_survives_template_round_trip(self):
        doc = Document()
        doc.add_heading("第1章 研究", 1)
        p = doc.add_paragraph("正文")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:beforeLines"), "50")
        spacing.set(qn("w:afterLines"), "100")
        spacing.set(qn("w:before"), "999")  # stale cached physical measurement
        p._p.get_or_add_pPr().append(spacing)
        doc.save(self.source)
        plan = build_plan(self.source, self.output, {"spec": {"body": {"space_before_pt": {"value": 0.5, "unit": "lines"}, "space_after_pt": {"value": 1, "unit": "lines"}}}})
        report = apply_plan(plan)
        self.assertEqual(verify_plan(plan, report)["status"], "passed")
        from word_format.properties import read_paragraph
        out = Document(self.output)
        self.assertEqual(read_paragraph(out, out.paragraphs[1]._p, "space_before_pt"), {"value": .5, "unit": "lines"})
        self.assertEqual(read_paragraph(out, out.paragraphs[1]._p, "space_after_pt"), {"value": 1, "unit": "lines"})

    def test_requirements_split_scopes_and_do_not_inherit_across_paragraphs(self):
        path = self.root / "requirements.txt"
        path.write_text("正文文字内容字型一律采用宋体，标题加黑，章节题目采用小三号字，内容采用小四号字汉字宋体和小四号Times New Roman英文。\n页眉：宋体，小五，居中。\n[4] Atlanta, Georgia, 1995.\n（十二）优秀毕业设计（论文）简介的规格\n正文五号，参考文献小五。\n上3cm，下2cm。", encoding="utf-8")
        profile = requirements_profile(path)
        rules = profile["rules"]
        self.assertTrue(any(r["action"] == "font.bold.set" and r["params"]["value"] and r["target"].get("role") == "heading" for r in rules))
        self.assertTrue(any(r["action"] == "font.east_asia.set" and r["params"]["value"] == "宋体" and r["target"].get("role") == "heading" for r in rules))
        self.assertTrue(any(r["action"] == "font.size.set" and r["params"]["value"] == 12 and r["target"].get("role") == "body" for r in rules))
        self.assertTrue(any(r["target"].get("role") == "header" for r in rules))
        self.assertFalse(any(r["params"].get("value") in {"Georgia", 10.5, 9, 30, 20} and r["target"].get("role") != "header" for r in rules))
        self.assertTrue(any(i.get("applicability") for i in profile["issues"]))

    def test_fixed_spacing_and_negated_bold_do_not_become_size_or_font_family(self):
        from parse_requirements import parse_requirements

        old = parse_requirements("一级标题宋体，三号，不加黑。正文固定22磅。")
        self.assertFalse(old["spec_patch"]["headings"][0]["bold"])
        self.assertEqual(old["spec_patch"]["body"]["line_spacing"], {"kind": "exact", "value_pt": 22})
        self.assertNotIn("font_size_pt", old["spec_patch"]["body"])
        rules, issues = parse_blocks([{"anchor": "line:1", "text": "正文宋体，小4号，固定22磅，不加黑。一级标题宋体，三号，加黑。"}])
        self.assertEqual(issues, [])
        self.assertTrue(any(r["action"] == "paragraph.line_spacing.set" and r["params"] == {"kind": "exact", "value": 22} for r in rules))
        self.assertFalse(any(r["action"] == "font.size.set" and r["params"]["value"] == 22 for r in rules))
        self.assertFalse(any(r["action"] == "font.east_asia.set" and r["params"]["value"] == "黑体" for r in rules))

    def toc_plan(self):
        return build_plan(self.source, self.output, {"operations": [{"action": "toc.configure", "target": {"type": "document"}, "params": {"enabled": True, "max_heading_level": 3}}]})

    def test_toc_infers_outline_without_changing_font_and_replaces_error(self):
        doc = Document()
        doc.add_paragraph("目录")
        append_field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u', cached_text="错误！未找到目录项。")
        for text in ("第1章 绪论", "1.1 研究背景", "第二章：介绍相关研究。"):
            p = doc.add_paragraph(text)
            p.runs[0]._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
        doc.save(self.source)
        plan = self.toc_plan()
        report = apply_plan(plan)
        self.assertEqual(verify_plan(plan, report)["status"], "passed")
        out = Document(self.output)
        status = result_status(out, 3)
        self.assertTrue(status["valid"])
        self.assertEqual(status["heading_count"], 2)
        self.assertNotIn("错误", status["text"])
        self.assertEqual(imported_fields(out), [])
        for p in out.paragraphs:
            if p.text in {"第1章 绪论", "1.1 研究背景", "第二章：介绍相关研究。"}:
                self.assertEqual(p.style.name, "Normal")
                self.assertEqual(p.runs[0]._r.rPr.rFonts.get(qn("w:eastAsia")), "宋体")

    def test_custom_outline_style_and_depth_limit(self):
        doc = Document()
        s = doc.styles.add_style("学校章节", WD_STYLE_TYPE.PARAGRAPH)
        node = OxmlElement("w:outlineLvl")
        node.set(qn("w:val"), "0")
        s.element.get_or_add_pPr().append(node)
        doc.add_paragraph("绪论", style=s)
        doc.add_heading("超出深度", 4)
        doc.save(self.source)
        report = apply_plan(self.toc_plan())
        self.assertEqual(report["status"], "passed")
        self.assertEqual(result_status(Document(self.output), 3)["heading_count"], 1)

    def test_empty_error_and_missing_results_are_rejected(self):
        doc = Document()
        doc.add_heading("第1章 绪论", 1)
        doc.add_heading("第2章 测试", 1)
        doc.save(self.source)
        apply_plan(self.toc_plan())
        for text in ("", "错误！未找到目录项。", "Error! No table of contents entries found.", "第1章 绪论", "目录将在 Word 中更新"):
            out = Document(self.output)
            field = next(p for p in out.paragraphs if p.style.name == "WFM TOC Field")
            for node in field._p.iter(qn("w:t")):
                node.text = text
            path = self.root / "bad.docx"
            out.save(path)
            self.assertFalse(result_status(out, 3)["valid"], text)
            self.assertEqual(verify_structure(path, {"table_of_contents": {"enabled": True}})["status"], "failed", text)

    def test_no_headings_never_publishes_output(self):
        doc = Document()
        doc.add_paragraph("正文内容。")
        doc.save(self.source)
        with self.assertRaises(FormatError):
            apply_plan(self.toc_plan())
        self.assertFalse(self.output.exists())

    def test_toc_never_discards_prose_sharing_a_field_paragraph(self):
        doc = Document()
        p = doc.add_paragraph("保留正文")
        append_field(p, 'TOC \\o "1-3"', cached_text="旧目录")
        doc.add_heading("第1章 绪论", 1)
        doc.save(self.source)
        with self.assertRaises(FormatError):
            apply_plan(self.toc_plan())
        self.assertFalse(self.output.exists())

    def test_toc_outline_side_effect_cannot_change_the_wrong_level(self):
        from unittest.mock import patch
        from word_format.structures import apply_table_of_contents

        for inherited in (False, True):
            doc = Document()
            p = doc.add_paragraph("第1章 绪论")
            if inherited:
                node = OxmlElement("w:outlineLvl")
                node.set(qn("w:val"), "0")
                p._p.get_or_add_pPr().append(node)
            doc.save(self.source)

            def broken(document, tokens):
                result = apply_table_of_contents(document, tokens)
                heading = next(p for p in document.paragraphs if p.text == "第1章 绪论" and p.style.name == "Normal")
                heading._p.pPr.find(qn("w:outlineLvl")).set(qn("w:val"), "1")
                return result

            with patch("word_format.structures.apply_table_of_contents", side_effect=broken):
                with self.assertRaises(FormatError):
                    apply_plan(self.toc_plan())
            self.assertFalse(self.output.exists())

    def test_toc_never_discards_an_unrelated_page_field(self):
        doc = Document()
        p = doc.add_paragraph()
        append_field(p, 'TOC \\o "1-3"', cached_text="旧目录")
        append_field(p, 'PAGE', cached_text="1")
        doc.add_heading("第1章 绪论", 1)
        doc.save(self.source)
        with self.assertRaises(FormatError):
            apply_plan(self.toc_plan())
        self.assertFalse(self.output.exists())


if __name__ == "__main__":
    unittest.main()
