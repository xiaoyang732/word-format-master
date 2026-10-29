"""Evidence extraction, semantic mapping and deterministic profile round trips."""

import copy
from dataclasses import replace
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))

from word_format import apply_plan, verify_plan
from word_format.contracts import FormatError, sha256
from word_format.profiles import profile_plan, requirements_profile, template_profile
from word_format.requirement_sources import extract_requirements
from word_format.registry import REGISTRY


class ProfileTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.target = self.root / "待修改论文.docx"
        self.output = self.root / "排版后.docx"
        doc = Document()
        doc.add_heading("第1章 目标研究", 1)
        doc.add_paragraph("目标正文内容，保留姓名与数据。")
        doc.save(self.target)

    def requirements(self, value):
        path = self.root / "要求.txt"
        path.write_text(value, encoding="utf-8")
        return requirements_profile(path)

    def test_requirements_round_trip_and_evidence(self):
        profile = self.requirements("A4；页边距2.5cm。正文宋体，小四，1.5倍行距，首行缩进2字符，不加粗。一级标题黑体，三号，居中。")
        self.assertEqual(profile["issues"], [])
        original = sha256(self.target)
        plan = profile_plan(self.target, self.output, profile)
        report = apply_plan(plan)
        self.assertEqual(verify_plan(plan, report)["status"], "passed")
        out = Document(self.output)
        self.assertEqual(out.paragraphs[1].text, "目标正文内容，保留姓名与数据。")
        self.assertEqual(out.paragraphs[1].runs[0].font.size.pt, 12)
        self.assertFalse(out.paragraphs[1].runs[0].bold)
        self.assertAlmostEqual(out.sections[0].top_margin.mm, 25, delta=.03)
        self.assertEqual(sha256(self.target), original)
        self.assertIn("profile_context", report)
        self.assertTrue(report["profile_context"]["mapping"][0]["evidence"][0]["anchor"].startswith("line:"))

    def test_partial_unknown_negative_and_conflict_do_not_execute(self):
        for text in ("正文宋体，字符间距加宽2磅。", "正文不要使用宋体。", "正文小四。正文五号。", "正文小四，图片除外。"):
            with self.subTest(text=text):
                profile = self.requirements(text)
                self.assertTrue(profile["issues"])
                review = profile_plan(self.target, self.output, profile)
                self.assertEqual(review["status"], "needs_clarification")
                self.assertFalse(self.output.exists())

    def test_explicit_resolution_keeps_conflicting_evidence(self):
        profile = self.requirements("正文小四。正文五号。")
        issue = profile["issues"][0]
        decisions = {issue["id"]: {"decision": "choose", "candidate": 1, "reason": "用户明确选择五号"}}
        plan = profile_plan(self.target, self.output, profile, resolutions=decisions)
        report = apply_plan(plan)
        self.assertEqual(Document(self.output).paragraphs[1].runs[0].font.size.pt, 10.5)
        self.assertEqual(len(report["profile_context"]["decisions"][0]["issue"]["candidates"]), 2)

    def test_docx_table_header_and_textbox_anchors(self):
        path = self.root / "学校要求.docx"
        doc = Document()
        row = doc.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text = "正文"
        row.cells[1].text = "宋体，小四，首行缩进2字符"
        doc.sections[0].header.paragraphs[0].text = "A4"
        from docx.oxml import OxmlElement
        box = OxmlElement("w:txbxContent")
        p = OxmlElement("w:p")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "一级标题黑体三号"
        r.append(t)
        p.append(r)
        box.append(p)
        doc.add_paragraph()._p.append(box)
        doc.save(path)
        extracted = extract_requirements(path)
        self.assertTrue(any("#row:" in b["anchor"] and "正文" in b["text"] for b in extracted["blocks"]))
        self.assertTrue(any("header" in b["anchor"] for b in extracted["blocks"]))
        self.assertTrue(any("一级标题" in b["text"] for b in extracted["blocks"]))
        profile = requirements_profile(path)
        self.assertTrue(any(r["action"] == "font.size.set" and r["params"]["value"] == 12 for r in profile["rules"]))

    def sample(self, multiple=False, mixed=False):
        path = self.root / "样例.docx"
        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(12)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        doc.add_heading("第1章 样例标题", 1)
        doc.add_paragraph("样例专属正文不能复制")
        if mixed:
            doc.add_paragraph("另一个正文").runs[0].font.size = Pt(14)
        doc.sections[0].top_margin = Mm(22)
        if multiple:
            section = doc.add_section(WD_SECTION_START.NEW_PAGE)
            section.top_margin = Mm(32)
            doc.add_heading("第2章 样例实验", 1)
            doc.add_paragraph("样例实验正文")
        doc.save(path)
        return path

    def test_docx_mixed_text_and_image_requirements_are_not_silently_complete(self):
        from PIL import Image
        image = self.root / "要求截图.png"
        Image.new("RGB", (16, 16), "white").save(image)
        source = self.root / "图文要求.docx"
        doc = Document()
        doc.add_paragraph("正文小四。")
        doc.add_picture(str(image))
        doc.save(source)
        profile = requirements_profile(source)
        self.assertTrue(profile["rules"])
        self.assertTrue(any("Non-text" in issue["reason"] for issue in profile["issues"]))
        self.assertEqual(profile_plan(self.target, self.output, profile)["status"], "needs_clarification")

    def test_template_effective_values_content_and_all_sections(self):
        sample = self.sample(multiple=True)
        doc = Document(self.target)
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        doc.add_heading("第2章 目标实验", 1)
        doc.add_paragraph("目标实验数据")
        doc.save(self.target)
        profile = template_profile(sample)
        self.assertEqual(len(profile["snapshot"]["sections"]), 2)
        odd_even_rules = [r for r in profile["rules"] if r["action"] == "header.odd_even_different.set"]
        self.assertEqual(len(odd_even_rules), 1)
        self.assertEqual(odd_even_rules[0]["target"], {"type": "document"})
        # Missing fonts in default heading styles are explicit unresolved evidence.
        decisions = {i["id"]: {"decision": "ignore", "reason": "样例未定义该属性，保留目标值"} for i in profile["issues"]}
        review = profile_plan(self.target, self.output, profile, resolutions=decisions)
        self.assertEqual(review["issues"][0]["id"], "section-mapping")
        plan = profile_plan(self.target, self.output, profile, resolutions=decisions,
                            mapping={"sections": {"1": 1, "2": 2}})
        report = apply_plan(plan)
        self.assertEqual(verify_plan(plan, report)["status"], "passed")
        out = Document(self.output)
        self.assertAlmostEqual(out.sections[0].top_margin.mm, 22, delta=.03)
        self.assertAlmostEqual(out.sections[1].top_margin.mm, 32, delta=.03)
        self.assertNotIn("样例", "".join(p.text for p in out.paragraphs))
        self.assertEqual(out.paragraphs[1].runs[0].font.size.pt, 12)

    def test_mixed_sample_property_not_silently_replaced_by_majority(self):
        profile = template_profile(self.sample(mixed=True))
        self.assertTrue(any(i.get("property") == "font_size_pt" for i in profile["issues"]))
        self.assertFalse(any(r["target"].get("role") == "body" and r["action"] == "font.size.set" for r in profile["rules"]))

    def test_inherited_unsupported_properties_are_explicit(self):
        sample = self.sample()
        doc = Document(sample)
        doc.styles["Normal"].font.underline = True
        doc.save(sample)
        profile = template_profile(sample)
        self.assertTrue(any("u" in i.get("properties", []) for i in profile["issues"]))
        self.assertIn("sections_xml", profile["snapshot"])

    def test_rule_mapping_uses_target_section_without_remapping_it(self):
        doc = Document(self.target)
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        doc.add_paragraph("第二节正文")
        doc.save(self.target)
        profile = template_profile(self.sample())
        decisions = {i["id"]: {"decision": "ignore", "reason": "保留不迁移的属性"} for i in profile["issues"]}
        top_rule = next(r for r in profile["rules"] if r["action"] == "page.margin_top.set")
        plan = profile_plan(self.target, self.output, profile, resolutions=decisions,
                            mapping={"sections": {"1": 2}, "rules": {top_rule["id"]: {"type": "section", "number": 1}}})
        report = apply_plan(plan)
        self.assertEqual(verify_plan(plan, report)["status"], "passed")
        out = Document(self.output)
        self.assertAlmostEqual(out.sections[0].top_margin.mm, 22, delta=.03)
        self.assertAlmostEqual(out.sections[1].top_margin.mm, 25.4, delta=.03)

    def test_low_confidence_requires_explicit_mapping(self):
        profile = self.requirements("正文小四。")
        rule = profile["rules"][0]
        rule["evidence"][0]["confidence"] = .5
        review = profile_plan(self.target, self.output, profile)
        self.assertEqual(review["status"], "needs_clarification")
        self.assertIn("Low-confidence", review["issues"][0]["reason"])
        plan = profile_plan(self.target, self.output, profile,
                            mapping={"rules": {rule["id"]: rule["target"]}})
        self.assertEqual(verify_plan(plan, apply_plan(plan))["status"], "passed")

    def test_malformed_profile_or_mapping_returns_contract_error(self):
        profile = self.requirements("正文小四。")
        for mapping in ([], {"rules": {"r1": []}}, {"sections": {"1": 1}}):
            with self.subTest(mapping=mapping), self.assertRaises(FormatError):
                profile_plan(self.target, self.output, profile, mapping=mapping)
        for change in ({"rules": [None]}, {"coverage": []},
                       {"source": {**profile["source"], "kind": "other"}}):
            with self.subTest(change=change), self.assertRaises(FormatError):
                profile_plan(self.target, self.output, {**profile, **change})

    def test_profile_source_change_during_apply_cannot_publish(self):
        profile = self.requirements("正文小四。")
        plan = profile_plan(self.target, self.output, profile)
        cap = REGISTRY["font.size.set"]

        def changing_source(index, ref, params):
            cap.handler(index, ref, params)
            Path(profile["source"]["path"]).write_text("正文五号。", encoding="utf-8")

        with patch.dict(REGISTRY, {cap.action: replace(cap, handler=changing_source)}):
            with self.assertRaises(FormatError):
                apply_plan(plan)
        self.assertFalse(self.output.exists())

    def test_pdf_text_has_page_evidence_and_blank_page_requires_ocr(self):
        path = self.root / "要求.pdf"
        text = (b"BT /F1 12 Tf 50 750 Td (Body font size 12pt) Tj ET\n"
                b"q 10 0 0 10 100 100 cm BI /W 1 /H 1 /CS /RGB /BPC 8 ID \xff\xff\xff EI Q")
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Count 2 /Kids [3 0 R 5 0 R] >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 7 0 R >> >> /Contents 4 0 R >>",
            b"<< /Length " + str(len(text)).encode() + b" >>\nstream\n" + text + b"\nendstream",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 6 0 R >>",
            b"<< /Length 0 >>\nstream\n\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        ]
        data = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for i, body in enumerate(objects, 1):
            offsets.append(len(data))
            data.extend(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
        xref = len(data)
        data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
        for offset in offsets[1:]:
            data.extend(f"{offset:010} 00000 n \n".encode())
        data.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
        path.write_bytes(data)
        profile = requirements_profile(path)
        self.assertEqual(profile["snapshot"]["blocks"][0]["anchor"], "page:1")
        self.assertTrue(any(r["action"] == "font.size.set" and r["params"]["value"] == 12 for r in profile["rules"]))
        self.assertTrue(any(i.get("anchor") == "page:2" and "OCR" in i["reason"] for i in profile["issues"]))
        self.assertTrue(any(i.get("anchor") == "page:1" and "images" in i["reason"] for i in profile["issues"]))
        self.assertEqual(profile_plan(self.target, self.output, profile)["status"], "needs_clarification")
        self.assertFalse(self.output.exists())

    def test_cli_direct_request_cannot_replace_profile_authority(self):
        profile = self.requirements("正文小四。")
        plan = profile_plan(self.target, self.output, profile)
        request = self.root / "request.json"
        # A legacy spec wrapper must retain the same source binding after expansion.
        plan["request"].pop("operations")
        plan["request"]["spec"] = {"schema_version": "1.0", "body": {"font_size_pt": 12}}
        request.write_text(json.dumps(plan["request"]), encoding="utf-8")
        authority = Path(profile["source"]["path"])
        before = authority.read_bytes()
        for args in (["plan", self.target, self.output], ["audit", self.target]):
            result = subprocess.run([sys.executable, str(ROOT / "scripts/format_cli.py"), *map(str, args),
                                     "--request", str(request), "--output", str(authority)], capture_output=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(authority.read_bytes(), before)
        from word_format import build_plan
        expanded = build_plan(self.target, self.output, plan["request"])
        self.assertEqual(expanded["request"]["profile_context"]["source"], profile["source"])

    def test_profile_source_changes_and_output_aliases_are_rejected(self):
        profile = self.requirements("正文小四。")
        plan = profile_plan(self.target, self.output, profile)
        Path(profile["source"]["path"]).write_text("正文五号。", encoding="utf-8")
        with self.assertRaises(FormatError):
            apply_plan(plan)
        self.assertFalse(self.output.exists())
        sample = self.sample()
        profile = template_profile(sample)
        with self.assertRaises(FormatError):
            profile_plan(self.target, sample, profile)

    def test_cli_profile_plan_json_output_cannot_destroy_authority(self):
        profile = self.requirements("正文小四。")
        profile_path = self.root / "规范.json"
        profile_path.write_text(json.dumps(profile, ensure_ascii=False), encoding="utf-8")
        source = Path(profile["source"]["path"])
        before = source.read_bytes()
        result = subprocess.run([sys.executable, str(ROOT / "scripts/format_cli.py"), "plan", str(self.target),
                                 str(self.output), "--profile", str(profile_path), "--output", str(source)], capture_output=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(source.read_bytes(), before)

    def test_cli_requirements_profile_plan_apply_verify_round_trip(self):
        authority = self.root / "学校要求.txt"
        authority.write_text("A4；正文宋体，小四，首行缩进2字符。", encoding="utf-8")
        profile_path = self.root / "格式档案.json"
        plan_path = self.root / "计划.json"
        report_path = self.root / "验收.json"
        verification_path = self.root / "复检.json"
        commands = [
            ["profile", "requirements", authority, "--output", profile_path],
            ["plan", self.target, self.output, "--profile", profile_path, "--output", plan_path],
            ["apply", "--plan", plan_path, "--report", report_path],
            ["verify", "--plan", plan_path, "--report", report_path, "--output", verification_path],
        ]
        for command in commands:
            result = subprocess.run([sys.executable, str(ROOT / "scripts/format_cli.py"), *map(str, command)],
                                    capture_output=True, timeout=30)
            self.assertEqual(result.returncode, 0, result.stderr.decode("utf-8"))
        self.assertEqual(json.loads(verification_path.read_text(encoding="utf-8"))["status"], "passed")
        self.assertIn("profile_context", json.loads(report_path.read_text(encoding="utf-8")))
        self.assertEqual(Document(self.output).paragraphs[1].text, "目标正文内容，保留姓名与数据。")

    def test_dotx_read_without_rewriting_template(self):
        sample = self.sample()
        dotx = self.root / "模板.dotx"
        with zipfile.ZipFile(sample) as src, zipfile.ZipFile(dotx, "w") as dst:
            for item in src.infolist():
                blob = src.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    blob = blob.replace(b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml")
                dst.writestr(item, blob)
        before = sha256(dotx)
        profile = template_profile(dotx)
        self.assertTrue(profile["rules"])
        self.assertEqual(sha256(dotx), before)


if __name__ == "__main__":
    unittest.main()
