"""Content-scope checks for thesis quality reports."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
from thesis_quality import analyze_thesis  # noqa: E402
from verify_output import content_review_scope  # noqa: E402
from word_format import apply_plan, build_plan, verify_plan  # noqa: E402


class ThesisQualityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / "draft.docx"

    def make_thesis(self, *, body_text="正文包含引用[1,2]。"):
        document = Document()
        document.add_paragraph("论文封面标题", style="Title")
        document.add_paragraph("摘要", style="Heading 1")
        document.add_paragraph("中文摘要内容")
        document.add_paragraph("关键词：论文格式")
        document.add_paragraph("Abstract", style="Heading 1")
        document.add_paragraph("English abstract")
        document.add_paragraph("Key words: formatting")
        document.add_paragraph("目录", style="Heading 1")
        document.add_paragraph("第一章 绪论", style="Heading 1")
        document.add_paragraph(body_text)
        document.add_paragraph("参考文献", style="Heading 1")
        document.add_paragraph("[1] Known entry metadata.")
        document.save(self.source)

    def test_body_length_and_citation_gaps_are_reported_separately(self):
        self.make_thesis()
        report = analyze_thesis(
            self.source,
            minimum_chinese_characters=5,
            maximum_chinese_characters=100,
        )
        self.assertEqual(report["status"], "failed")
        self.assertGreaterEqual(
            report["measurement_scope"]["body"]["chinese_han_characters"], 5
        )
        self.assertEqual(report["checks"]["body_length"]["status"], "passed")
        self.assertEqual(
            report["checks"]["numeric_citation_mapping"]["citations_without_matching_reference"],
            [2],
        )
        self.assertEqual(
            report["checks"]["bibliography_source_verification"]["status"],
            "manual_review_required",
        )
        self.assertEqual(
            report["checks"]["bibliography_source_verification"]["entries"][0]["source_verification"],
            "not_verified",
        )

    def test_no_word_target_does_not_claim_length_pass_or_fail(self):
        self.make_thesis(body_text="正文。")
        report = analyze_thesis(self.source)
        self.assertEqual(report["status"], "review_required")
        self.assertEqual(report["checks"]["body_length"]["status"], "not_requested")
        self.assertEqual(report["checks"]["academic_claims"]["status"], "not_assessed")

    def test_author_date_style_bracket_year_is_not_mistaken_for_numeric_reference_id(self):
        document = Document()
        document.add_paragraph("论文标题", style="Title")
        document.add_paragraph("第一章 绪论", style="Heading 1")
        document.add_paragraph("研究方法已有讨论[2020]。")
        document.add_paragraph("参考文献", style="Heading 1")
        document.add_paragraph("张三. 研究方法. 学术出版社, 2020.")
        document.save(self.source)
        report = analyze_thesis(self.source)
        self.assertEqual(
            report["checks"]["numeric_citation_mapping"]["status"],
            "not_assessable",
        )

    def test_absent_body_heading_cannot_silently_pass_requested_word_target(self):
        document = Document()
        document.add_paragraph("Only front matter")
        document.save(self.source)
        report = analyze_thesis(self.source, minimum_chinese_characters=10)
        self.assertEqual(report["checks"]["body_length"]["status"], "not_assessable")

    def test_format_verifier_always_discloses_content_review_boundary(self):
        self.assertEqual(content_review_scope()["status"], "not_assessed")
        self.assertEqual(content_review_scope()["bibliography_metadata_and_sources"], "not_assessed")

    def test_auto_numbered_references_are_included_in_citation_mapping(self):
        document = Document()
        document.add_paragraph("论文标题", style="Title")
        document.add_paragraph("第一章 绪论", style="Heading 1")
        document.add_paragraph("正文引用[1,2]。")
        document.add_paragraph("参考文献", style="Heading 1")
        document.add_paragraph("Author. First source.")
        document.add_paragraph("Author. Second source.")
        document.save(self.source)
        output = Path(self.temp.name) / "numbered.docx"
        plan = build_plan(
            self.source,
            output,
            {"spec": {"references": {"numbering_mode": "word-numbering"}}},
        )
        report = apply_plan(plan)
        self.assertEqual(verify_plan(plan, report)["status"], "passed")

        quality = analyze_thesis(output)
        mapping = quality["checks"]["numeric_citation_mapping"]
        self.assertEqual(mapping["status"], "passed")
        self.assertEqual(mapping["listed_reference_ids"], [1, 2])
        self.assertEqual(mapping["citations_without_matching_reference"], [])


if __name__ == "__main__":
    unittest.main()
