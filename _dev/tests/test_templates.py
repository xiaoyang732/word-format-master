"""Behavior regressions for the unified template path and document preservation."""

from __future__ import annotations

import copy
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))
from word_format.api import apply_plan, verify_plan
from word_format.contracts import FormatError, json_hash
from word_format.inspect import DocumentIndex
from word_format.profiles import requirements_profile, template_profile
from word_format.templates import compile_template, load_template, template_plan, validate_template


class TemplateTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source, self.output = self.root / "source.docx", self.root / "output.docx"

    def template(self, rules):
        return {
            "schema_version": "1.0",
            "kind": "word-format-template",
            "id": "test",
            "version": "1",
            "origin": "user",
            "sources": [],
            "issues": [],
            "rules": [
                {"state": "set", "id": f"r{i}", "evidence": [{"anchor": "user instruction"}], **r}
                for i, r in enumerate(rules)
            ],
        }

    def test_body_format_excludes_images_and_equations_and_preserves_media(self):
        image = self.root / "image.png"
        Image.new("RGB", (20, 40), "blue").save(image)
        doc = Document()
        doc.add_paragraph("Unchanged content")
        doc.add_picture(str(image))
        p = doc.add_paragraph()
        p._p.append(OxmlElement("m:oMath"))
        doc.save(self.source)
        template = self.template(
            [
                {
                    "action": "paragraph.line_spacing.set",
                    "target": {"type": "paragraph", "role": "body", "all": True},
                    "params": {"kind": "exact", "value": 31, "unit": "pt"},
                }
            ]
        )
        plan = template_plan(self.source, self.output, template)
        report = apply_plan(plan)
        out = Document(self.output)
        self.assertEqual(out.paragraphs[0].paragraph_format.line_spacing, Pt(31))
        self.assertIsNone(out.paragraphs[1].paragraph_format.line_spacing)
        self.assertIsNone(out.paragraphs[2].paragraph_format.line_spacing)
        for name in ["word/media/image1.png"]:
            with zipfile.ZipFile(self.source) as before, zipfile.ZipFile(self.output) as after:
                self.assertEqual(before.read(name), after.read(name))
        result = verify_plan(plan, report)
        self.assertEqual(result["execution_status"], "passed")
        self.assertEqual(result["status"], "pending")

    def test_user_template_does_not_fill_font_or_change_shared_style(self):
        doc = Document()
        doc.styles["Normal"].font.name = "Courier New"
        doc.styles["Normal"].font.size = Pt(17)
        doc.add_paragraph("Body")
        doc.save(self.source)
        template = self.template(
            [
                {
                    "action": "paragraph.spacing_after.set",
                    "target": {"type": "paragraph", "role": "body", "all": True},
                    "params": {"value": 9, "unit": "pt"},
                }
            ]
        )
        apply_plan(template_plan(self.source, self.output, template))
        out = Document(self.output)
        self.assertEqual(out.styles["Normal"].font.name, "Courier New")
        self.assertEqual(out.styles["Normal"].font.size, Pt(17))
        self.assertEqual(doc.styles["Normal"].element.xml, out.styles["Normal"].element.xml)

    def test_deleted_rule_cannot_be_hidden_by_rehashing_plan(self):
        doc = Document()
        doc.add_paragraph("Body")
        doc.save(self.source)
        template = self.template(
            [
                {
                    "action": "font.size.set",
                    "target": {"type": "paragraph", "role": "body", "all": True},
                    "params": {"value": 18, "unit": "pt"},
                }
            ]
        )
        plan = template_plan(self.source, self.output, template)
        plan["operations"] = []
        plan["request"]["operations"] = []
        plan["plan_sha256"] = json_hash({k: v for k, v in plan.items() if k != "plan_sha256"})
        with self.assertRaisesRegex(FormatError, "coverage"):
            apply_plan(plan)

    def test_native_heading_scheme_and_unnumbered_references_are_idempotent(self):
        doc = Document()
        doc.add_heading("第1章 绪论", 1)
        doc.add_paragraph("Chapter prose")
        doc.add_heading("1.1 Motivation", 2)
        doc.add_heading("参考文献", 1)
        doc.add_paragraph("[1] Source information")
        doc.save(self.source)
        template = self.template(
            [
                {
                    "action": "heading.scheme.configure",
                    "target": {"type": "document"},
                    "params": {
                        "levels": [
                            {"format": "decimal", "text": "第%1章", "start": 1, "suffix": "space"},
                            {"format": "decimal", "text": "%1.%2", "start": 1, "suffix": "space"},
                        ]
                    },
                }
            ]
        )
        plan = template_plan(self.source, self.output, template)
        report = apply_plan(plan)
        self.assertEqual(verify_plan(plan, report)["execution_status"], "passed")
        out = Document(self.output)
        self.assertEqual(out.paragraphs[0].text, "绪论")
        self.assertEqual(out.paragraphs[2].text, "Motivation")
        self.assertEqual(
            out.paragraphs[3]._p.xpath("./w:pPr/w:numPr/w:numId")[0].get(qn("w:val")), "0"
        )
        count = len(out.part.numbering_part.element)
        second = self.root / "second.docx"
        apply_plan(template_plan(self.output, second, template))
        out2 = Document(second)
        self.assertEqual(count, len(out2.part.numbering_part.element))
        self.assertEqual([p.text for p in out.paragraphs], [p.text for p in out2.paragraphs])

    def test_language_and_keyword_segments_are_independent(self):
        doc = Document()
        doc.add_paragraph("摘要")
        doc.add_paragraph("中文内容")
        doc.add_paragraph("关键词：中文")
        doc.add_paragraph("Abstract")
        doc.add_paragraph("English prose")
        doc.add_paragraph("Key words: English")
        doc.save(self.source)
        template = self.template(
            [
                {
                    "action": "font.bold.set",
                    "target": {
                        "type": "text_range",
                        "role": "keywords",
                        "language": "en",
                        "segment": "label",
                        "all": True,
                    },
                    "params": {"value": True},
                },
                {
                    "action": "font.size.set",
                    "target": {
                        "type": "paragraph",
                        "semantic_role": "abstract_body",
                        "language": "zh",
                        "all": True,
                    },
                    "params": {"value": 19, "unit": "pt"},
                },
            ]
        )
        apply_plan(template_plan(self.source, self.output, template))
        out = Document(self.output)
        self.assertEqual(out.paragraphs[1].runs[0].font.size, Pt(19))
        self.assertIsNone(out.paragraphs[4].runs[0].font.size)
        self.assertTrue(out.paragraphs[5].runs[0].bold)
        self.assertFalse(bool(out.paragraphs[5].runs[-1].bold))
        self.assertFalse(bool(out.paragraphs[2].runs[0].bold))

    def test_multiple_sources_conflict_without_last_wins(self):
        profiles = []
        for i, size in enumerate([12, 18]):
            source = self.root / f"requirements{i}.txt"
            source.write_text(f"正文 {size}pt", encoding="utf-8")
            profiles.append(requirements_profile(source))
        template = compile_template(profiles)
        self.assertTrue(any("Sources disagree" in i["reason"] for i in template["issues"]))

    def test_template_source_is_bound_and_protected(self):
        doc = Document()
        doc.add_paragraph("Body")
        doc.save(self.source)
        template = self.template(
            [
                {
                    "action": "font.size.set",
                    "target": {"type": "paragraph", "role": "body", "all": True},
                    "params": {"value": 18, "unit": "pt"},
                }
            ]
        )
        path = self.root / "template.json"
        path.write_text(json.dumps(template), encoding="utf-8")
        value, source = load_template(path)
        plan = template_plan(self.source, self.output, value, template_source=source)
        path.write_text(json.dumps({**template, "version": "2"}), encoding="utf-8")
        with self.assertRaisesRegex(FormatError, "changed"):
            apply_plan(plan)

    def test_builtin_uses_explicit_template_without_changing_image_style(self):
        doc = Document()
        doc.add_heading("第一章 Chapter", 1)
        doc.add_paragraph("Body")
        image = self.root / "img.png"
        Image.new("RGB", (20, 20), "red").save(image)
        doc.add_picture(str(image))
        doc.save(self.source)
        template, source = load_template("thesis-standard")
        plan = template_plan(self.source, self.output, template, template_source=source)
        self.assertNotIn("status", plan)
        apply_plan(plan)
        out = Document(self.output)
        self.assertEqual(out.paragraphs[-1]._p.pPr, doc.paragraphs[-1]._p.pPr)
        self.assertEqual(out.styles["Normal"].element.xml, doc.styles["Normal"].element.xml)

    def test_empty_abstract_placeholder_is_reported_and_not_guessed(self):
        doc = Document()
        doc.add_paragraph("摘要")
        doc.add_paragraph()
        doc.add_paragraph("关键词：")
        doc.save(self.source)
        profile = template_profile(self.source)
        self.assertTrue(any("Empty abstract" in issue["reason"] for issue in profile["issues"]))

    def test_cover_layout_is_data_and_has_no_fabricated_personal_information(self):
        from word_format.cover import add_standard_thesis_cover

        path = Path(__file__).resolve().parents[2] / "templates/builtin/thesis-standard/cover.json"
        layout = json.loads(path.read_text(encoding="utf-8"))
        title = next(block for block in layout["blocks"] if block["field"] == "title_cn")
        title["tokens"]["font_size_pt"] = 23
        doc = Document()
        add_standard_thesis_cover(doc, title_cn="真实题目", layout=layout, author_name="真实作者")
        self.assertEqual(doc.paragraphs[0].runs[0].font.size, Pt(23))
        text = "".join(p.text for p in doc.paragraphs) + "".join(
            c.text for t in doc.tables for r in t.rows for c in r.cells
        )
        self.assertIn("真实作者", text)
        for fabricated in ("张三", "2022010101", "二〇二六年五月", "RAG"):
            self.assertNotIn(fabricated, text)
        self.assertEqual(len(doc.tables[0].rows), 1)
        self.assertTrue(doc.tables[0]._tbl.xpath("./w:tblPr/w:tblLayout"))

    def test_requirement_synopsis_scope_does_not_leak_to_headers(self):
        doc = Document()
        doc.add_paragraph("正文 12pt")
        doc.add_paragraph("优秀毕业设计（论文）简介撰写要求")
        doc.add_paragraph("正文 10pt")
        doc.sections[0].header.paragraphs[0].text = "正文 15pt"
        doc.save(self.source)
        profile = requirements_profile(self.source)
        blocks = profile["snapshot"]["blocks"]
        self.assertEqual(
            [b["document_kind"] for b in blocks], ["thesis", "summary", "summary", "thesis"]
        )
        self.assertEqual(profile["coverage"]["excluded_blocks"], 2)

    def test_sample_target_coverage_requires_explicit_preservation(self):
        from word_format.contracts import sha256

        sample = self.root / "sample.docx"
        Document().save(sample)
        doc = Document()
        doc.add_paragraph("Body")
        doc.save(self.source)
        template = self.template([])
        template["sources"] = [{"kind": "template", "path": str(sample), "sha256": sha256(sample)}]
        review = template_plan(self.source, self.output, template)
        self.assertEqual(review["status"], "needs_clarification")
        oid = review["issues"][0]["target_ids"][0]
        plan = template_plan(
            self.source,
            self.output,
            template,
            mapping={"preserve": {oid: "No corresponding sample content"}},
        )
        self.assertNotIn("status", plan)
        result = verify_plan(plan, apply_plan(plan))
        self.assertEqual(result["status"], "partial")

    def test_invalid_template_shapes_raise_contract_errors(self):
        template = self.template(
            [{"action": "font.bold.set", "target": [], "params": {"value": True}}]
        )
        with self.assertRaises(FormatError):
            validate_template(template)
        template["rules"][0]["target"] = {"type": "paragraph", "role": "body", "all": True}
        template["issues"] = [{"id": "r0", "reason": "duplicate"}]
        with self.assertRaisesRegex(FormatError, "Duplicate"):
            validate_template(template)

    def test_cached_toc_labels_are_not_chapter_headings(self):
        doc = Document()
        doc.add_heading("目录", 1)
        from docx.enum.style import WD_STYLE_TYPE

        doc.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
        doc.add_paragraph("第1章 绪论4", style="toc 1")
        doc.add_paragraph("参考文献20", style="toc 1")
        doc.add_heading("绪论", 1)
        index = DocumentIndex(doc)
        self.assertEqual(index.records["p1"]["role"], "toc")
        self.assertIsNone(index.records["p1"]["heading_level"])
        self.assertEqual(index.records["p2"]["module"], "toc")
        self.assertEqual(index.records["p3"]["semantic_role"], "chapter_heading")

    def test_builtin_creates_missing_heading_style_and_preserves_content_controls(self):
        doc = Document()
        doc.styles.element.remove(doc.styles["Heading 2"].element)
        doc.add_heading("第一章 绪论", 1)
        doc.add_paragraph("1.1 目标")
        control = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        paragraph = OxmlElement("w:p")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "保留控件内文字"
        run.append(text)
        paragraph.append(run)
        content.append(paragraph)
        control.append(content)
        doc.element.body.insert(2, control)
        from lxml import etree

        original = etree.tostring(control)
        doc.save(self.source)
        template, authority = load_template("thesis-standard")
        plan = template_plan(self.source, self.output, template, template_source=authority)
        report = apply_plan(plan)
        result = verify_plan(plan, report)
        self.assertEqual(result["execution_status"], "passed")
        self.assertEqual(result["status"], "partial")
        output = Document(self.output)
        self.assertEqual(etree.tostring(output.element.body.find(qn("w:sdt"))), original)
        self.assertEqual(
            next(p for p in output.paragraphs if p.text == "目标").style.name, "Heading 2"
        )

    def test_reference_numbering_never_allocates_suppression_id_zero(self):
        doc = Document()
        numbering = doc.part.numbering_part.element
        for child in list(numbering):
            numbering.remove(child)
        doc.add_heading("参考文献", 1)
        doc.add_paragraph("[1] Example source")
        doc.save(self.source)
        template = self.template(
            [
                {
                    "action": "reference.numbering.convert",
                    "target": {"type": "document"},
                    "params": {"numbering_mode": "word-numbering"},
                }
            ]
        )
        plan = template_plan(self.source, self.output, template)
        report = apply_plan(plan)
        self.assertEqual(verify_plan(plan, report)["execution_status"], "passed")
        paragraph = Document(self.output).paragraphs[-1]
        self.assertGreater(
            int(paragraph._p.xpath("./w:pPr/w:numPr/w:numId")[0].get(qn("w:val"))), 0
        )


if __name__ == "__main__":
    unittest.main()
