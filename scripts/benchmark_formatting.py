#!/usr/bin/env python3
"""Reproduce the review's 96-paragraph thesis workload in an isolated directory."""

import argparse
import cProfile
import io
import json
from pathlib import Path
import pstats
import sys
import tempfile
import time

from docx import Document


def synthetic_thesis():
    doc = Document()
    for text in ("基于深度学习的文档版式确定性优化研究", "摘要", "本文研究Word文档确定性版式引擎的性能与正确性问题。" * 8,
                 "关键词：版式引擎；Word；确定性", "Abstract",
                 "This thesis studies deterministic Word formatting engines and their performance. " * 6,
                 "Key words: formatting; Word; deterministic"):
        doc.add_paragraph(text)
    for ch in range(1, 4):
        doc.add_paragraph(f"第{ch}章 系统设计与实现", style="Heading 1")
        for sec in range(1, 4):
            doc.add_paragraph(f"{ch}.{sec} 核心模块分析", style="Heading 2")
            for i in range(8):
                doc.add_paragraph(f"第{ch}章第{sec}节正文段落{i+1}：为了验证字符级投影与属性克隆的开销，这里构造具有一定长度的学术论文正文内容，包含中文与English混合文本（2026）。" * 2)
    for text in ("参考文献", "[1] 张三. 论文排版规范研究[J]. 计算机学报, 2025, 48(1): 1-12.",
                 "[2] 李四. 自动排版技术综述[J]. 软件学报, 2026, 37(2): 20-35.", "致谢", "感谢导师与实验室同学在研究过程中的指导与帮助。"):
        doc.add_paragraph(text)
    return doc


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output", type=Path)
    parser.add_argument("--profile", action="store_true")
    args = parser.parse_args()
    sys.path.insert(0, str(args.project_root.resolve() / "scripts"))
    from word_format import build_plan, apply_plan, verify_plan
    from cli_output import configure_cli_output
    configure_cli_output()
    with tempfile.TemporaryDirectory() as folder:
        source, destination = Path(folder) / "source.docx", Path(folder) / "result.docx"
        doc = synthetic_thesis()
        doc.save(source)
        spec = json.loads((args.project_root / "assets/presets.json").read_text(encoding="utf-8"))["presets"][0]
        plan = build_plan(source, destination, {"spec": spec})
        profile = cProfile.Profile()
        start = time.perf_counter()
        if args.profile:
            profile.enable()
        report = apply_plan(plan)
        if args.profile:
            profile.disable()
        elapsed = time.perf_counter() - start
        result = {"paragraphs": len(doc.paragraphs), "operations": len(plan["operations"]),
                  "apply_seconds": round(elapsed, 3), "profiled": args.profile,
                  "verification": verify_plan(plan, report)["status"]}
        if args.profile:
            summary = io.StringIO()
            stats = pstats.Stats(profile, stream=summary).sort_stats("cumulative")
            stats.print_stats(15)
            result.update(function_calls=stats.total_calls, profile=summary.getvalue())
    if args.output:
        from json_output import write_json
        write_json(args.output, result)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
