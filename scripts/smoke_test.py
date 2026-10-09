#!/usr/bin/env python3
"""Run a structural DOCX round-trip smoke test for the formatting engine."""

from __future__ import annotations

import copy
import json
import hashlib
import os
import subprocess
import sys
import tempfile
import threading
import urllib.error
import urllib.request
from http.server import ThreadingHTTPServer
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from analyze_docx import analyze_docx
from apply_spec import SUPPORTED_APPLICATION_PATHS, append_field, apply_document, validate_spec
from citation_workflow import build_citation_task
from dashboard_session import build_ai_handoff, prepare_dashboard_session, public_dashboard_session, record_visual_verification
from render_docx import resolve_renderer
from runtime_detection import detect_renderers
from serve_dashboard import DashboardHandler, _cleanup_analysis_tasks
from config import MAX_ANALYSIS_TASKS, utc_now
from verify_output import validate_visual_report


ROOT = Path(__file__).resolve().parent.parent


def assert_dashboard_api_surface() -> None:
    server = ThreadingHTTPServer(("127.0.0.1", 0), DashboardHandler)
    server.dashboard_session = None
    server.dashboard_session_lock = threading.Lock()
    server.dashboard_runtime_lock = threading.Lock()
    server.dashboard_runtime_job = {"status": "idle", "percent": 0}
    server.dashboard_analysis_tasks = {}
    server.dashboard_analysis_lock = threading.Lock()
    worker = threading.Thread(target=server.serve_forever, daemon=True)
    worker.start()
    base_url = f"http://127.0.0.1:{server.server_address[1]}"
    try:
        with urllib.request.urlopen(f"{base_url}/api/capabilities", timeout=5) as response:
            capabilities = json.loads(response.read().decode("utf-8"))
            assert response.headers["X-Word-Format-API-Version"] == "1"
        assert capabilities["api_version"] == "1"
        untrusted_request = urllib.request.Request(
            f"{base_url}/api/capabilities",
            headers={"Host": "example.invalid"},
        )
        try:
            urllib.request.urlopen(untrusted_request, timeout=5)
        except urllib.error.HTTPError as exc:
            assert exc.code == 403
        else:
            raise AssertionError("dashboard accepted an untrusted Host header")
        request = urllib.request.Request(
            f"{base_url}/api/apply",
            data=b"{}",
            method="POST",
            headers={"Content-Type": "application/json"},
        )
        try:
            urllib.request.urlopen(request, timeout=5)
        except urllib.error.HTTPError as exc:
            assert exc.code == 404
        else:
            raise AssertionError("direct /api/apply endpoint must remain unavailable")
    finally:
        server.shutdown()
        server.server_close()
        worker.join(timeout=5)


def main() -> int:
    queue_server = type("QueueServer", (), {})()
    queue_server.dashboard_analysis_lock = threading.Lock()
    queue_server.dashboard_analysis_tasks = {
        f"task-{index:02d}": {
            "status": "completed",
            "created_at": utc_now(),
            "completed_at": f"{index:02d}",
        }
        for index in range(MAX_ANALYSIS_TASKS)
    }
    _cleanup_analysis_tasks(queue_server, reserve_slot=True)
    assert len(queue_server.dashboard_analysis_tasks) == MAX_ANALYSIS_TASKS - 1
    assert "task-00" not in queue_server.dashboard_analysis_tasks

    presets = json.loads((ROOT / "assets" / "presets.json").read_text(encoding="utf-8"))
    thesis_std = next(
        item for item in presets["presets"] if item["id"] == "thesis-standard"
    )
    assert thesis_std["page"]["size"] == "A4"
    assert thesis_std["page"]["margin_top_mm"] == 25.4
    assert thesis_std["page"]["margin_bottom_mm"] == 25.4
    assert thesis_std["page"]["margin_left_mm"] == 25.0
    assert thesis_std["headers_footers"]["preserve_existing"] is True
    assert thesis_std["headers_footers"]["header"]["enabled"] is False
    assert thesis_std["body"]["font_east_asia"] == "宋体"
    assert thesis_std["body"]["font_size_pt"] == 12
    assert thesis_std["body"]["line_spacing"]["value"] == 1.5
    assert {item["level"] for item in thesis_std["headings"]} >= {1, 2, 3}
    assert thesis_std["headings"][0]["font_size_pt"] == 16
    assert thesis_std["lists"]["numbered"]["style"] == "decimal-period"
    assert thesis_std["captions"]["figure"]["position"] == "below"
    assert thesis_std["captions"]["table"]["position"] == "above"
    assert thesis_std["table_of_contents"]["max_heading_level"] == 3
    assert thesis_std["references"]["citation_system"] == "GB/T 7714-2015"
    assert "GB/T 7713.1" in thesis_std["authority"]["override_note"]
    base = thesis_std
    spec = json.loads(json.dumps(base))
    spec["headers_footers"]["preserve_existing"] = False
    spec["headers_footers"]["header"] = {
        "enabled": True,
        "text": "Formatting smoke test",
        "alignment": "center",
        "font_latin": "Arial",
        "font_size_pt": 9,
    }
    spec["page_numbers"]["format"] = "page-x-of-y"
    spec["page"]["size"] = "Letter"
    spec["captions"]["figure"]["position"] = "above"
    spec["captions"]["table"]["position"] = "above"
    spec["lists"]["numbered"] = {
        "style": "decimal-fullwidth-parentheses",
        "start": 1,
        "left_indent_mm": 7.4,
        "hanging_indent_mm": 4.5,
    }
    spec["table_of_contents"] = {
        "enabled": True,
        "title": "Contents",
        "max_heading_level": 2,
        "page_break_after": True,
    }

    with tempfile.TemporaryDirectory(prefix="word-format-master-smoke-") as temp:
        runtime = detect_renderers(Path(temp) / "runtime")
        assert "platform" not in runtime
        assert set(runtime) >= {"word", "libreoffice", "preferred", "renderer_ready", "pdf_rasterizer", "visual_ready", "official_download_page"}
        assert runtime["pdf_rasterizer"]["available"] is True
        assert runtime["official_download_page"].startswith("https://www.libreoffice.org/")
        renderer_snapshot = {
            "word": {"id": "word", "name": "Microsoft Word", "available": True, "path": "WINWORD.EXE", "method": "installed executable"},
            "libreoffice": {"id": "libreoffice", "name": "LibreOffice", "available": True, "path": "soffice", "method": "installed executable"},
            "preferred": "word",
            "renderer_ready": True,
            "pdf_rasterizer": {"id": "poppler", "name": "Poppler", "available": True, "path": "poppler"},
            "visual_ready": True,
        }
        assert resolve_renderer(renderer_snapshot, "auto")[0] == "word"
        assert resolve_renderer(renderer_snapshot, "libreoffice")[0] == "libreoffice"
        try:
            resolve_renderer({**renderer_snapshot, "word": {"name": "Microsoft Word", "available": False}}, "word")
            raise AssertionError("explicit unavailable renderer was silently replaced")
        except RuntimeError as exc:
            assert "Microsoft Word" in str(exc)
        source = Path(temp) / "source.docx"
        output = Path(temp) / "formatted.docx"
        cleaned = Path(temp) / "cleaned.docx"
        document = Document()
        document.add_heading("Formatting Test", level=1)
        document.add_heading("Methods", level=2)
        document.add_paragraph("Body text used for structural verification.")
        for item in range(1, 6):
            document.add_paragraph(f"Numbered item {item}", style="List Number")
        document.add_table(rows=2, cols=2)
        document.add_paragraph("表 1 参数状态")
        figure = document.add_paragraph()
        figure._p.append(OxmlElement("w:drawing"))
        document.add_paragraph("图 1 排版流程")
        document.add_heading("参考文献", level=1)
        document.add_paragraph("[1] A. Author, Example reference, 2026.")
        append_field(document.sections[0].footer.paragraphs[0], "PAGE")
        document.save(source)

        blank_template = Path(temp) / "blank-template.docx"
        blank_template_document = Document()
        blank_template_document.add_paragraph("Template body")
        blank_template_document.save(blank_template)
        blank_template_spec = analyze_docx(blank_template)["inferred_spec"]
        assert blank_template_spec["body"]["font_size_pt"] > 0
        assert blank_template_spec["body"]["font_latin"]
        assert blank_template_spec["body"]["font_east_asia"]
        assert blank_template_spec["body"]["first_line_indent_mm"] == 0
        assert blank_template_spec["body"]["line_spacing"]
        assert blank_template_spec["headers_footers"]["preserve_existing"] is False
        assert blank_template_spec["headers_footers"]["header"] == {"enabled": False, "text": ""}
        assert blank_template_spec["headers_footers"]["footer"] == {"enabled": False, "text": ""}
        assert blank_template_spec["page_numbers"] == {"enabled": False}
        assert blank_template_spec["table_of_contents"] == {"enabled": False}

        module_source = Path(temp) / "module-template.docx"
        module_document = Document()
        module_document.add_paragraph("论文题目", style="Title")
        module_document.add_paragraph("作者与学院信息")
        module_document.add_paragraph("摘要")
        module_document.add_paragraph("摘要正文。")
        module_document.add_paragraph("关键词：格式、模板")
        module_document.add_paragraph("目录")
        module_document.add_heading("第1章 绪论", level=1)
        module_document.add_paragraph("第一章正文。")
        module_document.add_heading("参考文献", level=1)
        module_document.add_paragraph("[1] A. Author, Example reference, 2026.")
        module_document.add_heading("附录 A", level=1)
        module_document.add_paragraph("附录内容。")
        module_document.save(module_source)
        module_analysis = analyze_docx(module_source)
        module_structure = module_analysis["inferred_spec"]["document_structure"]
        assert [item["kind"] for item in module_structure["ordered_modules"]] == [
            "cover", "abstract", "keywords", "toc", "chapters", "references", "appendix"
        ]
        assert module_structure["order_errors"] == []
        module_spec = json.loads(json.dumps(module_analysis["inferred_spec"]))
        module_target = Path(temp) / "module-target.docx"
        module_target_document = Document()
        module_target_document.add_paragraph("论文题目", style="Title")
        module_target_document.add_paragraph("摘要")
        module_target_document.add_paragraph("摘要正文。")
        module_target_document.add_paragraph("关键词：格式")
        module_target_document.add_paragraph("目录")
        module_target_document.add_heading("第1章 绪论", level=1)
        module_target_document.add_paragraph("第一章正文。")
        module_target_document.add_heading("参考文献", level=1)
        module_target_document.add_paragraph("[1] A. Author, Example reference, 2026.")
        module_target_document.save(module_target)
        try:
            apply_document(module_target, Path(temp) / "module-output.docx", module_spec)
        except ValueError as exc:
            assert "appendix" in str(exc)
        else:
            raise AssertionError("Missing module must prevent publication")
        assert not (Path(temp) / "module-output.docx").exists()

        # Test dual abstract (Chinese + English) document structure and TOC placement
        dual_abstract_source = Path(temp) / "dual-abstract-source.docx"
        dual_abstract_doc = Document()
        dual_abstract_doc.add_paragraph("双语毕业论文题目", style="Title")
        dual_abstract_doc.add_paragraph("作者与指导教师信息")
        dual_abstract_doc.add_paragraph("摘要")
        dual_abstract_doc.add_paragraph("本文研究了基于深度学习的文档智能排版系统。")
        dual_abstract_doc.add_paragraph("关键词：深度学习、文档排版、智能格式化")
        dual_abstract_doc.add_paragraph("ABSTRACT")
        dual_abstract_doc.add_paragraph("This paper investigates an intelligent document formatting system based on deep learning.")
        dual_abstract_doc.add_paragraph("KEY WORDS: Deep Learning, Document Formatting, Smart Layout")
        dual_abstract_doc.add_heading("第1章 绪论", level=1)
        dual_abstract_doc.add_paragraph("第一章详细研究背景。")
        dual_abstract_doc.add_heading("参考文献", level=1)
        dual_abstract_doc.add_paragraph("[1] B. Scholar, Dual Abstract Research, 2026.")
        dual_abstract_doc.save(dual_abstract_source)

        dual_analysis = analyze_docx(dual_abstract_source)
        dual_structure = dual_analysis["inferred_spec"]["document_structure"]
        assert [item["kind"] for item in dual_structure["ordered_modules"]] == [
            "cover", "abstract", "keywords", "abstract", "keywords", "chapters", "references"
        ]
        assert dual_structure["order_errors"] == []

        # Apply TOC to dual abstract document
        dual_spec = json.loads(json.dumps(thesis_std))
        dual_spec["table_of_contents"] = {
            "enabled": True,
            "title": "目录",
            "max_heading_level": 3,
            "page_break_after": True,
        }
        dual_output = Path(temp) / "dual-abstract-output.docx"
        dual_report = apply_document(dual_abstract_source, dual_output, dual_spec, verification={"visual_enabled": False})
        assert dual_report["verification"]["structure"]["status"] == "passed"

        # Verify TOC is inserted AFTER English keywords and BEFORE Chapter 1
        reopened_dual = Document(dual_output)
        dual_texts = [p.text.strip() for p in reopened_dual.paragraphs if p.text.strip()]
        toc_heading_idx = dual_texts.index("目录")
        en_keywords_idx = next(i for i, t in enumerate(dual_texts) if t.startswith("KEY WORDS"))
        ch1_idx = next(i for i, p in enumerate(reopened_dual.paragraphs) if p.style and p.style.name == "Heading 1" and "绪论" in p.text)
        assert en_keywords_idx < toc_heading_idx < ch1_idx
        # Verify page break exists before TOC heading on the last page of the abstract
        toc_heading_p = next(p for p in reopened_dual.paragraphs if p.style and p.style.name == "WFM TOC Heading")
        prev_p = toc_heading_p._p.getprevious()
        assert prev_p is not None and any(node.get(qn("w:type")) == "page" for node in prev_p.iter(qn("w:br")))

        # Test Chinese abstract only document TOC placement
        cn_only_source = Path(temp) / "cn-only-source.docx"
        cn_only_doc = Document()
        cn_only_doc.add_paragraph("中文论文题目", style="Title")
        cn_only_doc.add_paragraph("作者信息")
        cn_only_doc.add_paragraph("摘要")
        cn_only_doc.add_paragraph("本文研究了中文摘要单独存在时的目录排版。")
        cn_only_doc.add_paragraph("关键词：中文摘要、目录生成")
        cn_only_doc.add_heading("第1章 绪论", level=1)
        cn_only_doc.add_paragraph("正文第一章。")
        cn_only_doc.add_heading("参考文献", level=1)
        cn_only_doc.add_paragraph("[1] C. Author, Chinese Only Abstract, 2026.")
        cn_only_doc.save(cn_only_source)

        cn_only_output = Path(temp) / "cn-only-output.docx"
        cn_only_report = apply_document(cn_only_source, cn_only_output, dual_spec, verification={"visual_enabled": False})
        assert cn_only_report["verification"]["structure"]["status"] == "passed"

        reopened_cn = Document(cn_only_output)
        cn_texts = [p.text.strip() for p in reopened_cn.paragraphs if p.text.strip()]
        cn_toc_idx = cn_texts.index("目录")
        cn_kw_idx = next(i for i, t in enumerate(cn_texts) if t.startswith("关键词"))
        cn_ch1_idx = next(i for i, p in enumerate(reopened_cn.paragraphs) if p.style and p.style.name == "Heading 1" and "绪论" in p.text)
        assert cn_kw_idx < cn_toc_idx < cn_ch1_idx

        centered_header_template = Path(temp) / "centered-header-template.docx"
        centered_header_document = Document()
        centered_header_document.add_paragraph("Template body")
        centered_header = centered_header_document.sections[0].header
        centered_header.paragraphs[0].alignment = 0
        centered_header_paragraph = centered_header.add_paragraph("Centered template header")
        centered_header_paragraph.alignment = 1
        centered_header_document.save(centered_header_template)
        centered_header_spec = analyze_docx(centered_header_template)["inferred_spec"]
        assert centered_header_spec["headers_footers"]["header"]["enabled"] is True
        assert centered_header_spec["headers_footers"]["header"]["text"] == "Centered template header"
        assert centered_header_spec["headers_footers"]["header"]["alignment"] == "center"

        tab_centered_header_template = Path(temp) / "tab-centered-header-template.docx"
        tab_centered_header_document = Document()
        tab_centered_header_document.add_paragraph("Template body")
        tab_centered_header = tab_centered_header_document.sections[0].header.paragraphs[0]
        tab_centered_header.alignment = 0
        tab_centered_header.paragraph_format.tab_stops.add_tab_stop(
            Inches(3.25), WD_TAB_ALIGNMENT.CENTER,
        )
        tab_centered_header.add_run().add_tab()
        tab_centered_header.add_run("Visually centered template header")
        tab_centered_header_document.save(tab_centered_header_template)
        tab_centered_header_spec = analyze_docx(tab_centered_header_template)["inferred_spec"]
        assert tab_centered_header_spec["headers_footers"]["header"]["alignment"] == "center"

        citation_task = build_citation_task(Document(source), filename=source.name)
        body_anchor = next(item for item in citation_task["paragraphs"] if item["text"].startswith("Body text"))
        spec["citations"] = {
            "placements": [{
                "paragraph_index": body_anchor["paragraph_index"],
                "paragraph_sha256": body_anchor["paragraph_sha256"],
                "reference_ids": ["ref-1"],
                "citation_text": "[1]",
                "insert": "before-terminal-punctuation",
                "confidence": 0.98,
                "reason": "Smoke-test evidence",
            }]
        }

        report = apply_document(source, output, spec, verification={"visual_enabled": False})
        result = analyze_docx(output)
        reopened = Document(output)

        paragraph_styles = {paragraph.text: paragraph.style.name for paragraph in reopened.paragraphs}
        assert paragraph_styles["表 1 参数状态"] == "Table Caption"
        assert paragraph_styles["图 1 排版流程"] == "Figure Caption"
        assert paragraph_styles["A. Author, Example reference, 2026."] == "Bibliography"
        assert any(paragraph.text == "Body text used for structural verification[1]." for paragraph in reopened.paragraphs)
        numbered_paragraphs = [paragraph for paragraph in reopened.paragraphs if paragraph.text.startswith("Numbered item")]
        assert len(numbered_paragraphs) == 5
        assert all(paragraph._p.pPr.find(qn("w:numPr")) is not None for paragraph in numbered_paragraphs)
        numbering = reopened.part.numbering_part.element
        assert any(node.get(qn("w:val")) == "（%1）" for node in numbering.findall(".//" + qn("w:lvlText")))
        # Multi-level heading assertions
        h1_p = next(p for p in reopened.paragraphs if p.text == "Formatting Test")
        h2_p = next(p for p in reopened.paragraphs if p.text == "Methods")
        ref_p = next(p for p in reopened.paragraphs if p.text == "参考文献")
        assert h1_p._p.pPr.find(qn("w:numPr")) is not None
        assert h1_p._p.pPr.find(qn("w:numPr")).find(qn("w:ilvl")).get(qn("w:val")) == "0"
        assert h2_p._p.pPr.find(qn("w:numPr")) is not None
        assert h2_p._p.pPr.find(qn("w:numPr")).find(qn("w:ilvl")).get(qn("w:val")) == "1"
        ref_num_pr = ref_p._p.pPr.find(qn("w:numPr")) if ref_p._p.pPr is not None else None
        assert ref_num_pr is None or (ref_num_pr.find(qn("w:numId")) is not None and ref_num_pr.find(qn("w:numId")).get(qn("w:val")) == "0")
        assert any(node.get(qn("w:val")) == "第%1章" for node in numbering.findall(".//" + qn("w:lvlText")))
        assert "Bibliography" in [style.name for style in reopened.styles]
        body_children = list(reopened.element.body)
        table_caption = next(paragraph._p for paragraph in reopened.paragraphs if paragraph.text == "表 1 参数状态")
        figure_caption = next(paragraph._p for paragraph in reopened.paragraphs if paragraph.text == "图 1 排版流程")
        figure_object = next(paragraph._p for paragraph in reopened.paragraphs if paragraph._p.find(qn("w:drawing")) is not None)
        assert body_children.index(table_caption) < body_children.index(reopened.tables[0]._tbl)
        assert body_children.index(figure_caption) < body_children.index(figure_object)
        assert abs(reopened.sections[0].page_width.mm - 215.9) < 0.2
        assert result["document"]["fields"]["types"].get("PAGE", 0) == 1
        assert result["document"]["fields"]["types"].get("NUMPAGES", 0) == 1
        assert result["document"]["fields"]["types"].get("TOC", 0) >= 1
        assert any('TOC \\o "1-2" \\h \\z \\u' in code for code in result["document"]["fields"]["codes"])
        toc_check = next(
            check for check in report["verification"]["structure"]["checks"]
            if check["name"] == "table-of-contents"
        )
        assert "headings=3" in toc_check["detail"]
        assert any(
            part["kind"] == "header" and part["text"] == "Formatting smoke test"
            for part in result["document"]["headers_footers"]
        )
        update = reopened.settings.element.find(qn("w:updateFields"))
        assert update is not None and update.get(qn("w:val")) in {"true", "1"}
        assert any("PAGE field" in change for change in report["changes"])
        assert any("Cleared" in change and "header/footer" in change for change in report["changes"])
        assert any("decimal-fullwidth-parentheses" in change for change in report["changes"])
        assert any("Word TOC field" in change for change in report["changes"])
        assert any("AI-approved citation" in change for change in report["changes"])
        assert report["verification"]["structure"]["status"] == "passed"
        assert report["verification"]["visual"]["status"] == "skipped"

        direct_visual_output = Path(temp) / "direct-visual.docx"
        direct_visual = apply_document(
            source,
            direct_visual_output,
            spec,
            verification={
                "visual_enabled": True,
                "visual_model": "current",
                "visual_execution_mode": "dashboard-direct",
            },
        )
        assert direct_visual["verification"]["visual"]["status"] == "skipped"
        assert "本地下载模式" in direct_visual["verification"]["visual"]["reason"]

        # Test heading first-line indent control for H2 and H3
        indent_heading_spec = json.loads(json.dumps(spec))
        indent_heading_spec["headings"] = [
            {"level": 1, "font_size_pt": 16, "bold": True, "alignment": "center", "first_line_indent_chars": 0},
            {"level": 2, "font_size_pt": 14, "bold": True, "alignment": "left", "first_line_indent_chars": 2},
            {"level": 3, "font_size_pt": 12, "bold": True, "alignment": "left", "first_line_indent_chars": 2},
        ]
        indent_output = Path(temp) / "indent-headings-output.docx"
        indent_report = apply_document(source, indent_output, indent_heading_spec)
        assert indent_report["verification"]["structure"]["status"] == "passed"
        indent_doc = Document(str(indent_output))
        h2_style = indent_doc.styles["Heading 2"]
        h2_ind = h2_style._element.find(qn("w:pPr")).find(qn("w:ind"))
        assert h2_ind is not None
        assert h2_ind.get(qn("w:firstLineChars")) == "200"
        h3_style = indent_doc.styles["Heading 3"]
        h3_ind = h3_style._element.find(qn("w:pPr")).find(qn("w:ind"))
        assert h3_ind is not None
        assert h3_ind.get(qn("w:firstLineChars")) == "200"

        # Test preserving existing header when applying default thesis-standard preset
        existing_hdr_source = Path(temp) / "existing-hdr-source.docx"
        existing_hdr_output = Path(temp) / "existing-hdr-output.docx"
        existing_hdr_doc = Document()
        existing_hdr_doc.add_heading("第1章 绪论", level=1)
        existing_hdr_doc.add_paragraph("正文测试段落。")
        existing_hdr_doc.sections[0].header.paragraphs[0].text = "东南大学本科毕业设计（论文）"
        existing_hdr_doc.save(existing_hdr_source)

        # Apply standard preset (preserve_existing: true, header.enabled: false)
        apply_report = apply_document(existing_hdr_source, existing_hdr_output, thesis_std)
        applied_analysis = analyze_docx(existing_hdr_output)
        assert any(
            "东南大学本科毕业设计（论文）" in part["text"]
            for part in applied_analysis["document"]["headers_footers"]
        )

        no_heading_source = Path(temp) / "no-heading-source.docx"
        no_heading_output = Path(temp) / "no-heading-output.docx"
        no_heading_document = Document()
        no_heading_document.add_paragraph("A visually bold title that is still Normal style").runs[0].bold = True
        no_heading_document.add_paragraph("Body text without semantic heading styles.")
        no_heading_document.save(no_heading_source)
        no_heading_spec = json.loads(json.dumps(spec))
        no_heading_spec.pop("citations", None)
        try:
            apply_document(no_heading_source, no_heading_output, no_heading_spec)
        except ValueError as exc:
            assert "TOC" in str(exc) or "no non-empty Heading" in str(exc)
        else:
            raise AssertionError("Empty TOC must prevent publication")
        assert not no_heading_output.exists()

        removal_spec = {
            "headers_footers": {
                "header": {"enabled": False},
                "footer": {"enabled": False},
            },
            "page_numbers": {"enabled": False},
            "table_of_contents": {"enabled": False},
        }
        removal_report = apply_document(output, cleaned, removal_spec)
        cleaned_analysis = analyze_docx(cleaned)
        assert cleaned_analysis["document"]["fields"]["types"].get("PAGE", 0) == 0
        assert cleaned_analysis["document"]["fields"]["types"].get("NUMPAGES", 0) == 0
        assert cleaned_analysis["document"]["fields"]["types"].get("TOC", 0) == 0
        assert not any(
            part["text"] == "Formatting smoke test"
            for part in cleaned_analysis["document"]["headers_footers"]
        )
        assert any("Removed managed page-number" in change for change in removal_report["changes"])

        template_cleared = Path(temp) / "template-cleared.docx"
        template_clear_spec = copy.deepcopy(blank_template_spec)
        # This case exercises clearing absent template features, not requiring a
        # cover-only template's module layout on a chapter/reference document.
        template_clear_spec.pop("document_structure", None)
        template_clear_report = apply_document(output, template_cleared, template_clear_spec)
        template_cleared_analysis = analyze_docx(template_cleared)
        assert template_cleared_analysis["document"]["fields"]["types"].get("PAGE", 0) == 0
        assert template_cleared_analysis["document"]["fields"]["types"].get("NUMPAGES", 0) == 0
        assert not any(
            part["text"].strip()
            for part in template_cleared_analysis["document"]["headers_footers"]
        )
        assert any("Cleared" in change and "header/footer" in change for change in template_clear_report["changes"])

        session_output = Path(temp) / "session-formatted.docx"
        session = prepare_dashboard_session(source, session_output)
        public_session = public_dashboard_session(session)
        assert "data" not in public_session["input"]
        assert validate_spec({"body": {"font_size_pt": "invalid"}})
        assert session["citation_context"]["status"] == "ready"
        assert session["citation_context"]["paragraph_count"] >= 1
        assert session["citation_context"]["reference_count"] >= 1
        imported_template_session = prepare_dashboard_session(source, Path(temp) / "imported-template-session.docx")
        imported_template_spec = json.loads(json.dumps(blank_template_spec))
        imported_template_spec["mode"] = "parameterized"
        imported_template_spec["template_required"] = False
        imported_template_handoff = build_ai_handoff(
            imported_template_session,
            {"spec": imported_template_spec, "verification": {"visual_enabled": False}},
            skill_dir=ROOT,
            registry=SUPPORTED_APPLICATION_PATHS,
            renderers=renderer_snapshot,
        )
        assert imported_template_handoff["execution_mode"] == "parameterized"
        invalid_template_session = prepare_dashboard_session(source, Path(temp) / "invalid-template-session.docx")
        invalid_template_spec = json.loads(json.dumps(imported_template_spec))
        invalid_template_spec["template_required"] = True
        try:
            build_ai_handoff(
                invalid_template_session,
                {"spec": invalid_template_spec, "verification": {"visual_enabled": False}},
                skill_dir=ROOT,
                registry=SUPPORTED_APPLICATION_PATHS,
                renderers=renderer_snapshot,
            )
            raise AssertionError("parameterized template incorrectly entered official-template mode")
        except ValueError as exc:
            assert "template_required" in str(exc)
        blocked_session = prepare_dashboard_session(source, Path(temp) / "blocked-session.docx")
        try:
            build_ai_handoff(
                blocked_session,
                {"spec": spec, "verification": {"visual_enabled": True, "render_method": "auto"}},
                skill_dir=ROOT,
                registry=SUPPORTED_APPLICATION_PATHS,
                renderers={
                    "word": {"name": "Microsoft Word", "available": False},
                    "libreoffice": {"name": "LibreOffice", "available": False},
                    "preferred": None,
                    "visual_ready": False,
                },
            )
            raise AssertionError("visual verification was enabled without a renderer")
        except ValueError as exc:
            assert "无法启用视觉验收" in str(exc)
        handoff = build_ai_handoff(
            session,
            {
                "spec": spec,
                "clear_direct_font_formatting": True,
                "citation_request": True,
                "verification": {"visual_enabled": True, "visual_model": "custom", "custom_model_id": "image-test", "render_method": "libreoffice"},
            },
            skill_dir=ROOT,
            registry=SUPPORTED_APPLICATION_PATHS,
            renderers=renderer_snapshot,
        )
        assert handoff["execution_mode"] == "parameterized"
        assert any(item["path"] == "body.font_latin" for item in handoff["application_methods"])
        assert any(item["path"] == "lists.numbered.style" for item in handoff["application_methods"])
        assert any(item["path"] == "table_of_contents.max_heading_level" for item in handoff["application_methods"])
        assert handoff["citation_context"]["task"]["paragraphs"]
        assert handoff["citation_request"] is True
        assert "Analyze citation_context.task" in " ".join(handoff["next_steps"])
        assert handoff["tools"]["apply"]["path"].endswith("scripts\\apply_spec.py") or handoff["tools"]["apply"]["path"].endswith("scripts/apply_spec.py")
        assert "--handoff" in handoff["tools"]["apply"]["arguments"]
        assert handoff["visual_verification"]["status"] == "pending"
        assert handoff["verification"]["render_method"] == "libreoffice"
        assert handoff["verification"]["custom_model_id"] == "image-test"
        assert handoff["tools"]["verify_visual"]["path"].endswith("scripts\\render_docx.py") or handoff["tools"]["verify_visual"]["path"].endswith("scripts/render_docx.py")
        assert handoff["tools"]["verify_visual"]["renderer_selection"]["fallback_policy"] == "no-fallback"
        assert handoff["tools"]["verify_visual"]["renderer_selection"]["selected_method"] == "libreoffice"
        assert handoff["tools"]["verify_visual"]["report_contract"]["page_count"]
        try:
            build_ai_handoff(
                session,
                {"spec": spec, "verification": {"visual_enabled": False}},
                skill_dir=ROOT,
                registry=SUPPORTED_APPLICATION_PATHS,
                renderers=renderer_snapshot,
            )
            raise AssertionError("a submitted dashboard session was overwritten")
        except ValueError as exc:
            assert "不能重复提交" in str(exc)

        handoff_path = Path(temp) / "confirmed-handoff.json"
        handoff_path.write_text(json.dumps(handoff, ensure_ascii=False), encoding="utf-8")
        cli_command = [
            sys.executable,
            str(ROOT / "scripts" / "apply_spec.py"),
            str(source),
            str(session_output),
            "--handoff",
            str(handoff_path),
        ]
        # Exercise the CLI under the same legacy pipe encoding as Windows CI;
        # the CLI must select UTF-8 itself, and callers must decode it explicitly.
        cli_environment = {**os.environ, "PYTHONUTF8": "0", "PYTHONIOENCODING": "cp1252"}
        confirmed_run = subprocess.run(cli_command, capture_output=True, text=True, encoding="utf-8", env=cli_environment, check=False)
        assert confirmed_run.returncode == 0, confirmed_run.stderr
        confirmed_report = json.loads(confirmed_run.stdout)
        assert confirmed_report["output"] == str(session_output)
        assert session_output.is_file()
        direct_spec_run = subprocess.run(
            [
                sys.executable,
                str(ROOT / "scripts" / "apply_spec.py"),
                str(source),
                str(Path(temp) / "blocked-direct-spec.docx"),
                "--spec",
                str(handoff_path),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            env=cli_environment,
            check=False,
        )
        assert direct_spec_run.returncode != 0
        tampered_spec = json.loads(json.dumps(handoff["spec"]))
        tampered_spec["body"]["font_size_pt"] = 99
        tampered_spec_path = Path(temp) / "tampered-spec.json"
        tampered_spec_path.write_text(json.dumps(tampered_spec, ensure_ascii=False), encoding="utf-8")
        tampered_run = subprocess.run(
            [*cli_command, "--spec", str(tampered_spec_path)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            env=cli_environment,
            check=False,
        )
        assert tampered_run.returncode == 2
        assert "only add validated citations" in tampered_run.stderr

        session_output.write_bytes(output.read_bytes())
        try:
            record_visual_verification(
                session,
                {
                    "session_id": session["id"],
                    "output_sha256": hashlib.sha256(session_output.read_bytes()).hexdigest(),
                    "report": {"model": {"id": "other-model", "supports_image_input": False}, "status": "skipped"},
                },
            )
            raise AssertionError("visual verification accepted a different custom model")
        except ValueError as exc:
            assert "自定义模型 ID" in str(exc)
        try:
            record_visual_verification(
                session,
                {
                    "session_id": session["id"],
                    "output_sha256": hashlib.sha256(session_output.read_bytes()).hexdigest(),
                    "report": {
                        "model": {"id": "image-test", "supports_image_input": True},
                        "renderer": {"method": "libreoffice", "page_files": ["page-1.png"]},
                        "status": "passed",
                        "reviewed_pages": [1],
                    },
                },
            )
            raise AssertionError("visual verification accepted a passed report without page_count")
        except ValueError as exc:
            assert "渲染页数" in str(exc)
        try:
            record_visual_verification(
                session,
                {
                    "session_id": session["id"],
                    "output_sha256": hashlib.sha256(session_output.read_bytes()).hexdigest(),
                    "page_count": 2,
                    "report": {
                        "model": {"id": "image-test", "supports_image_input": True},
                        "renderer": {"method": "word", "page_files": ["page-1.png", "page-2.png"]},
                        "status": "passed",
                        "reviewed_pages": [1, 2],
                    },
                },
            )
            raise AssertionError("visual verification accepted a renderer that contradicted the user selection")
        except ValueError as exc:
            assert "渲染器" in str(exc)
        render_dir = Path(temp) / "rendered-pages"
        render_dir.mkdir()
        page_files = []
        for index in (1, 2):
            page_path = render_dir / f"page-{index}.png"
            page_path.write_bytes(f"smoke-page-{index}".encode("ascii"))
            page_files.append(page_path)
        manifest = {
            "schema_version": "1.0",
            "source": str(session_output),
            "source_sha256": hashlib.sha256(session_output.read_bytes()).hexdigest(),
            "selected_method": "libreoffice",
            "page_count": 2,
            "pages": [
                {
                    "number": index,
                    "path": str(page_path),
                    "sha256": hashlib.sha256(page_path.read_bytes()).hexdigest(),
                }
                for index, page_path in enumerate(page_files, start=1)
            ],
        }
        manifest_path = render_dir / "render-manifest.json"
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
        visual_result = record_visual_verification(
            session,
            {
                "session_id": session["id"],
                "output_sha256": hashlib.sha256(session_output.read_bytes()).hexdigest(),
                "page_count": 2,
                "report": {
                    "model": {"id": "image-test", "supports_image_input": True},
                    "renderer": {"method": "libreoffice", "manifest_path": str(manifest_path), "page_files": [str(path) for path in page_files]},
                    "status": "passed",
                    "reviewed_pages": [1, 2],
                    "findings": [],
                },
            },
        )
        assert visual_result["status"] == "passed"
        assert session["visual_verification"]["output_sha256"] == hashlib.sha256(session_output.read_bytes()).hexdigest()
        skipped_visual = validate_visual_report(
            {"model": {"id": "text-only-test", "supports_image_input": False}, "status": "passed"},
            page_count=2,
        )
        assert skipped_visual["status"] == "skipped"
        incomplete_visual = validate_visual_report(
            {"model": {"id": "image-test", "supports_image_input": True}, "status": "passed", "reviewed_pages": [1]},
            page_count=2,
        )
        assert incomplete_visual["status"] == "failed"

        # Test declaration module detection and heuristic heading safety
        from document_structure import classify_role, is_heading_candidate

        assert is_heading_candidate("深度学习 计算机视觉 自然语言处理") is False
        assert classify_role("深度学习 计算机视觉 自然语言处理", "Normal", first=False) == "body"
        assert classify_role("关键词：深度学习、计算机视觉", "Normal", first=False) == "keywords"
        assert classify_role("第1章 绪论", "Normal", first=False) == "heading"
        assert is_heading_candidate("第1章 绪论") is True

        # Test declaration module detection
        decl_source = Path(temp) / "decl-source.docx"
        decl_doc = Document()
        decl_doc.add_paragraph("论文题目", style="Title")
        decl_doc.add_paragraph("诚信声明")
        decl_doc.add_paragraph("本人郑重声明所呈交的学位论文是本人独立完成的研究成果。")
        decl_doc.add_paragraph("摘要")
        decl_doc.add_paragraph("摘要正文。")
        decl_doc.add_paragraph("关键词：声明、排版")
        decl_doc.add_heading("第1章 绪论", level=1)
        decl_doc.add_paragraph("正文。")
        decl_doc.add_heading("参考文献", level=1)
        decl_doc.add_paragraph("[1] Test Ref.")
        decl_doc.save(decl_source)

        decl_analysis = analyze_docx(decl_source)
        decl_structure = decl_analysis["inferred_spec"]["document_structure"]
        assert [item["kind"] for item in decl_structure["ordered_modules"]] == [
            "cover", "declaration", "abstract", "keywords", "chapters", "references"
        ]
        assert decl_structure["order_errors"] == []

        # Test three-line table application
        table_source = Path(temp) / "table-source.docx"
        table_output = Path(temp) / "table-output.docx"
        table_doc = Document()
        table_doc.add_heading("第1章 绪论", level=1)
        t = table_doc.add_table(rows=3, cols=3)
        t.rows[0].cells[0].text = "Header 1"
        t.rows[0].cells[1].text = "Header 2"
        t.rows[0].cells[2].text = "Header 3"
        t.rows[1].cells[0].text = "Data A"
        t.rows[1].cells[1].text = "Data B"
        t.rows[1].cells[2].text = "Data C"
        table_doc.save(table_source)

        tbl_spec = copy.deepcopy(thesis_std)
        tbl_spec["tables"] = {
            "three_line_table": True,
            "repeat_header_row": True,
            "alignment": "center",
            "font_size_pt": 10.5,
            "font_east_asia": "宋体",
            "font_latin": "Times New Roman",
        }
        apply_document(table_source, table_output, tbl_spec, verification={"visual_enabled": False})
        reopened_table_doc = Document(str(table_output))
        assert len(reopened_table_doc.tables) == 1
        out_table = reopened_table_doc.tables[0]
        out_tblPr = out_table._tbl.tblPr
        tbl_borders = out_tblPr.find(qn("w:tblBorders"))
        assert tbl_borders is not None
        assert tbl_borders.find(qn("w:top")).get(qn("w:sz")) == "12"
        assert tbl_borders.find(qn("w:bottom")).get(qn("w:sz")) == "12"
        assert tbl_borders.find(qn("w:insideH")).get(qn("w:val")) == "none"
        h_tcPr = out_table.rows[0].cells[0]._tc.tcPr
        assert h_tcPr.find(qn("w:tcBorders")).find(qn("w:bottom")).get(qn("w:sz")) == "6"
        assert out_table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
        assert out_tblPr.find(qn("w:jc")).get(qn("w:val")) == "center"

        assert presets["template_workflows"] == []
        assert len(presets["presets"]) >= 1

    assert_dashboard_api_surface()

    print("DOCX smoke test passed: lists, citations, page size, captions, TOC, references, headers, fields, verification, AI handoff, three-level headings, declaration module, three-line tables")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
