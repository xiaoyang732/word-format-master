#!/usr/bin/env python3
"""Perform deterministic post-application DOCX checks."""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn

from config import LIST_NUMBERING_STYLES
from document_structure import detect_module_spans, validate_module_order

EXPECTED_LIST_MARKERS = LIST_NUMBERING_STYLES


def _check(checks: list[dict[str, Any]], name: str, passed: bool, detail: str) -> None:
    checks.append({"name": name, "status": "passed" if passed else "failed", "detail": detail})


def verify_structure(path: str | Path, spec: dict[str, Any], *, check_format: bool = True) -> dict[str, Any]:
    checks: list[dict[str, Any]] = []
    errors: list[str] = []
    try:
        document = Document(str(path))
        _check(checks, "docx-readable", True, "python-docx reopened the generated package")
    except Exception as exc:  # pragma: no cover - defensive boundary for corrupted packages
        _check(checks, "docx-readable", False, str(exc))
        return {"status": "failed", "checks": checks, "errors": [str(exc)]}

    requested_structure = spec.get("document_structure")
    if isinstance(requested_structure, dict) and requested_structure.get("ordered_modules"):
        paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
        target_items = [
            {
                "index": index,
                "text": paragraph.text.strip(),
                "style_name": paragraph.style.name if paragraph.style else "",
            }
            for index, paragraph in enumerate(paragraphs)
        ]
        actual_modules = detect_module_spans(target_items)
        expected_kinds = [str(item.get("kind") or item.get("id") or "") for item in requested_structure.get("ordered_modules", [])]
        actual_kinds = [str(item.get("kind") or item.get("id") or "") for item in actual_modules]
        OPTIONAL_STANDARD_MODULES = {"declaration", "symbols", "acknowledgements", "appendix"}
        required = [kind for kind in expected_kinds if kind != "chapters"]
        missing = [kind for kind in required if kind not in actual_kinds]
        unexpected = [kind for kind in actual_kinds if kind not in expected_kinds and kind not in OPTIONAL_STANDARD_MODULES]
        order_errors = validate_module_order(actual_modules)
        passed = not missing and not unexpected and not order_errors
        detail = f"expected={expected_kinds}, actual={actual_kinds}, missing={missing}, unexpected={unexpected}, order_errors={order_errors}"
        _check(checks, "document-module-structure", passed, detail)
        if missing:
            errors.append(f"required document modules are missing: {', '.join(missing)}")
        if unexpected:
            errors.append(f"document modules are not defined by the template: {', '.join(unexpected)}")
        errors.extend(order_errors)

    page = spec.get("page", {})
    expected_size = {"A4": (210.0, 297.0), "Letter": (215.9, 279.4)}.get(str(page.get("size")))
    if expected_size and document.sections:
        section = document.sections[0]
        actual = (section.page_width.mm, section.page_height.mm)
        expected = expected_size if page.get("orientation") != "landscape" else expected_size[::-1]
        passed = all(abs(left - right) <= 0.4 for left, right in zip(actual, expected))
        _check(checks, "page-geometry", passed, f"actual={actual!r}, expected={expected!r}")
        if not passed:
            errors.append("page geometry does not match the requested specification")

    numbered = spec.get("lists", {}).get("numbered", {})
    if numbered:
        list_paragraphs = 0
        num_ids: set[str] = set()
        for paragraph in document.paragraphs:
            style_name = paragraph.style.name.lower() if paragraph.style else ""
            if (
                style_name.startswith("heading")
                or style_name.startswith("标题")
                or style_name in {"title", "subtitle", "toc", "wfm toc heading", "bibliography", "references", "reference", "参考文献"}
            ):
                continue
            ppr = paragraph._p.pPr
            num_pr = ppr.find(qn("w:numPr")) if ppr is not None else None
            if num_pr is not None:
                list_paragraphs += 1
                num_id = num_pr.find(qn("w:numId"))
                if num_id is not None and num_id.get(qn("w:val")):
                    num_ids.add(num_id.get(qn("w:val")))
        expected_marker = EXPECTED_LIST_MARKERS.get(str(numbered.get("style")))
        marker_matches = False
        try:
            numbering = document.part.numbering_part.element
        except (KeyError, NotImplementedError):
            numbering = None
        if numbering is not None:
            for num in numbering.findall(qn("w:num")):
                if num.get(qn("w:numId")) not in num_ids:
                    continue
                abstract_ref = num.find(qn("w:abstractNumId"))
                abstract_id = abstract_ref.get(qn("w:val")) if abstract_ref is not None else None
                abstract = next(
                    (item for item in numbering.findall(qn("w:abstractNum")) if item.get(qn("w:abstractNumId")) == abstract_id),
                    None,
                )
                level = abstract.find(qn("w:lvl")) if abstract is not None else None
                fmt = level.find(qn("w:numFmt")) if level is not None else None
                text = level.find(qn("w:lvlText")) if level is not None else None
                actual_marker = (
                    fmt.get(qn("w:val")) if fmt is not None else None,
                    text.get(qn("w:val")) if text is not None else None,
                )
                marker_matches = marker_matches or actual_marker == expected_marker
        passed = list_paragraphs == 0 or marker_matches
        detail = f"{list_paragraphs} paragraph(s) carry Word numPr; marker definition matched={marker_matches}"
        _check(checks, "numbered-lists", passed, detail)
        if not passed:
            errors.append("numbered-list OOXML marker does not match the requested style")

    page_numbers = spec.get("page_numbers", {})
    if page_numbers.get("enabled"):
        location = str(page_numbers.get("location") or "footer")
        containers = [getattr(section, location) for section in document.sections]
        if page_numbers.get("show_on_first_page", True):
            containers.extend(
                getattr(section, f"first_page_{location}")
                for section in document.sections
                if section.different_first_page_header_footer
            )
        page_field_counts = []
        seen_parts: set[int] = set()
        for container in containers:
            identity = id(container._element)
            if identity in seen_parts:
                continue
            seen_parts.add(identity)
            page_field_counts.append(sum(
                1
                for paragraph in container.paragraphs
                for node in paragraph._p.iter(qn("w:instrText"))
                if node.text and re.search(r"\bPAGE\b", node.text, re.I)
            ))
        passed = bool(page_field_counts) and all(count == 1 for count in page_field_counts)
        _check(checks, "page-numbers", passed, f"PAGE fields per target {location} part={page_field_counts}")
        if not passed:
            errors.append(f"target {location} parts must contain exactly one PAGE field each")

    placements = spec.get("citations", {}).get("placements", [])
    if placements:
        paragraphs = [
            paragraph for paragraph in document.paragraphs
            if not paragraph.style or paragraph.style.name not in {"WFM TOC Heading", "WFM TOC Field", "WFM TOC Break"}
        ]
        citation_failures = []
        for placement in placements:
            index = int(placement.get("paragraph_index", -1))
            text = str(placement.get("citation_text") or "")
            if index < 0 or index >= len(paragraphs) or text not in paragraphs[index].text:
                citation_failures.append(index)
        passed = not citation_failures
        _check(checks, "citation-placements", passed, f"checked {len(placements)} placement(s)")
        if citation_failures:
            errors.append(f"citation text missing at paragraph indexes {citation_failures}")

    toc = spec.get("table_of_contents", {})
    if toc:
        managed = {
            paragraph.style.name
            for paragraph in document.paragraphs
            if paragraph.style and paragraph.style.name in {"WFM TOC Heading", "WFM TOC Field", "WFM TOC Break"}
        }
        all_toc_codes = [
            " ".join(node.text.split())
            for paragraph in document.paragraphs
            for node in paragraph._p.iter(qn("w:instrText"))
            if node.text and node.text.strip()
        ]
        toc_codes = [
            " ".join(node.text.split())
            for paragraph in document.paragraphs
            if paragraph.style and paragraph.style.name == "WFM TOC Field"
            for node in paragraph._p.iter(qn("w:instrText"))
            if node.text and re.search(r"\bTOC\b", node.text, re.I)
        ]
        update = document.settings.element.find(qn("w:updateFields"))
        update_enabled = update is not None and update.get(qn("w:val")) in {"true", "1"}
        if toc.get("enabled"):
            expected_level = int(toc.get("max_heading_level", 3))
            expected_switch = f'\\o "1-{expected_level}"'
            field_matches = any(expected_switch in code for code in toc_codes)
            heading_matches = [
                paragraph
                for paragraph in document.paragraphs
                if paragraph.text.strip()
                and paragraph.style
                and re.fullmatch(r"Heading ([1-9])", paragraph.style.name, re.I)
                and int(re.fullmatch(r"Heading ([1-9])", paragraph.style.name, re.I).group(1)) <= expected_level
            ]
            title = str(toc.get("title") or "目录")
            title_matches = any(
                paragraph.style and paragraph.style.name == "WFM TOC Heading" and paragraph.text == title
                for paragraph in document.paragraphs
            )
            break_matches = any(
                paragraph.style and paragraph.style.name == "WFM TOC Break"
                and any(node.get(qn("w:type")) == "page" for node in paragraph._p.iter(qn("w:br")))
                and paragraph._p.getprevious() is not None
                and any(n.text and "TOC " in n.text for n in paragraph._p.getprevious().iter(qn("w:instrText")))
                for paragraph in document.paragraphs
            )
            break_ok = break_matches if toc.get("page_break_after", True) else not break_matches
            passed = bool(toc_codes) and field_matches and bool(heading_matches) and title_matches and break_ok and update_enabled
            _check(
                checks,
                "table-of-contents",
                passed,
                f"TOC fields={len(toc_codes)}, level_switch={field_matches}, headings={len(heading_matches)}, title={title_matches}, "
                f"page_break={break_matches}, updateFields={update_enabled}",
            )
            if not passed:
                if not heading_matches:
                    errors.append(
                        f"table of contents has no non-empty Heading 1-{expected_level} paragraphs; "
                        "assign real Word Heading styles before accepting the document"
                    )
                else:
                    errors.append("managed Word TOC field, title, page break, or updateFields setting is missing")
        else:
            passed = not managed and not toc_codes
            _check(checks, "table-of-contents-removed", passed, f"managed styles={sorted(managed)}, managed TOC fields={len(toc_codes)}, other TOC fields={len(all_toc_codes) - len(toc_codes)}")
            if not passed:
                errors.append("managed Word TOC content was not removed")

    if check_format:
        try:
            from word_format.inspect import DocumentIndex
            from word_format.spec import spec_request
            from word_format.registry import REGISTRY
            from word_format.selectors import resolve_targets
            index=DocumentIndex(document)
            for operation in spec_request(index,spec)["operations"]:
                cap=REGISTRY[operation["action"]]
                if cap.structural: continue
                for ref in resolve_targets(index,operation["target"],cap.targets):
                    actual=cap.reader(index,ref,cap)
                    passed=cap.verifier(actual,operation["params"],cap)
                    detail=f"target={ref['id']}, expected={operation['params']!r}, actual={actual!r}"
                    _check(checks,"format:"+cap.action,passed,detail)
                    if not passed: errors.append(detail)
        except (ValueError,TypeError,KeyError) as exc:
            errors.append("Format verification failed: "+str(exc))
    status = "passed" if not errors else "failed"
    return {"status": status, "checks": checks, "errors": errors}


def validate_visual_report(report: dict[str, Any], page_count: int | None = None) -> dict[str, Any]:
    """Validate a model-produced report without deciding whether a model is multimodal."""
    if not isinstance(report, dict):
        return {"status": "failed", "errors": ["视觉验收报告必须是对象"]}
    model = report.get("model") if isinstance(report.get("model"), dict) else {}
    if model.get("supports_image_input") is not True:
        reason = (
            "所选模型不支持图片输入"
            if model.get("supports_image_input") is False
            else "未确认所选模型是否支持图片输入"
        )
        return {"status": "skipped", "reason": reason, "model": model}
    status = str(report.get("status") or "")
    if status == "skipped":
        return {"status": "skipped", "reason": str(report.get("reason") or "模型未执行视觉验收"), "model": model}
    if status not in {"passed", "failed"}:
        return {"status": "failed", "errors": ["视觉验收状态必须是 passed、failed 或 skipped"]}
    pages = report.get("reviewed_pages")
    if page_count is not None:
        if isinstance(page_count, bool) or not isinstance(page_count, int) or page_count < 1:
            return {"status": "failed", "errors": ["渲染页数必须是大于 0 的整数"]}
        if not isinstance(pages, list):
            return {"status": "failed", "errors": ["视觉验收报告必须列出已检查页面"]}
        try:
            reviewed_pages = sorted({int(item) for item in pages})
        except (TypeError, ValueError):
            return {"status": "failed", "errors": ["已检查页面必须是页码整数"]}
        if reviewed_pages != list(range(1, page_count + 1)):
            return {"status": "failed", "errors": ["视觉验收报告必须覆盖每个渲染页面"]}
    return {"status": status, "model": model, "reviewed_pages": pages or [], "findings": report.get("findings", [])}


def main() -> int:
    if len(sys.argv) != 3:
        print("usage: verify_output.py OUTPUT.docx SPEC.json", file=sys.stderr)
        return 2
    result = verify_structure(sys.argv[1], json.loads(Path(sys.argv[2]).read_text(encoding="utf-8")))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
