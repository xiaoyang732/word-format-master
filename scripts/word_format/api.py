"""One inspect/plan/apply/verify API for direct AI requests and legacy handoffs."""

from __future__ import annotations

import copy
import json
import os
import uuid
from collections import Counter
from pathlib import Path

from docx import Document

from .contracts import FormatError, MAX_OPERATIONS, SCHEMA_VERSION, json_hash, output_path, require_keys, sha256, source_path, text_value
from .inspect import DocumentIndex, inspect_document
from .registry import REGISTRY, capabilities, equal_value
from .selectors import resolve_targets
from .verification import atomic_projection, check_page_geometry, structural_snapshot, verify_structural_preservation


def build_plan(input_path,output,request,*,_document=None):
    source=source_path(input_path)
    destination=output_path(source,output)
    require_keys(request,{"schema_version","origin","request_text","operations","spec","verification_spec","notes","clear_direct_font_formatting","verification"},label="request")
    if request.get("schema_version",SCHEMA_VERSION)!=SCHEMA_VERSION: raise FormatError("Unsupported request schema_version")
    if text_value(request.get("origin","direct"), "origin") not in {"direct","dashboard","legacy"}: raise FormatError("origin must be direct, dashboard or legacy")
    if "verification" in request and not isinstance(request["verification"], dict): raise FormatError("verification must be an object")
    if "spec" in request and "operations" in request: raise FormatError("Choose operations or legacy spec, not both")
    document=_document or Document(source)
    index=DocumentIndex(document)
    if "spec" in request:
        from .spec import spec_request
        expanded=spec_request(index,request["spec"],request.get("clear_direct_font_formatting",False))
        expanded["origin"]=request.get("origin","direct")
        if "verification" in request: expanded["verification"]=request["verification"]
        if "request_text" in request: expanded["request_text"]=request["request_text"]
        request=expanded
    operations=request.get("operations")
    if not isinstance(operations,list) or len(operations)>MAX_OPERATIONS:
        raise FormatError(f"operations must be an array with at most {MAX_OPERATIONS} entries")
    if any(isinstance(o,dict) and o.get("action")=="section.break.next_page.insert" for o in operations) and any(
        not isinstance(o,dict) or o.get("action")!="section.break.next_page.insert" for o in operations):
        raise FormatError("Insert section boundaries in a separate plan, then inspect the new DOCX before formatting", "needs_clarification")
    resolved=[];conflicts={}; assignments=[]; normalized=[]
    for i,item in enumerate(operations):
        require_keys(item,{"id","action","target","params"},{"action","target","params"},f"operation {i+1}")
        text_value(item["action"], "action")
        if not isinstance(item["target"], dict): raise FormatError("target must be an object")
        text_value(item["target"].get("type"), "target.type")
        cap=REGISTRY.get(item["action"])
        if cap is None: raise FormatError(f"Unknown operation {item['action']!r}","unsupported")
        params=cap.validator(item["params"])
        if not cap.structural and item["target"].get("type")=="text_range" and "text_range" not in cap.targets:
            raise FormatError("Paragraph properties cannot target a character range", "unsupported")
        targets=resolve_targets(index,item["target"],cap.targets)
        if cap.property=="style" and "style:"+params["name"] not in index.records:
            targets=[{**r,"new_style_name":params["name"]} for r in targets]
        for ref in targets:
            # Conflicting assignments are rejected before any mutation. Zero
            # opposite indent defaults from snapshots are collapsed by the adapter.
            scope=params.get("role") if cap.property=="caption_position" else params.get("variant") if cap.property.startswith("header:") else None
            key=(ref["id"],cap.property,ref.get("start"),ref.get("end"),scope)
            if key in conflicts and not equal_value(conflicts[key],params):
                raise FormatError(f"Conflicting operations for {ref['id']} / {cap.property}","needs_clarification")
            conflicts[key]=params
            if not cap.structural or cap.property.startswith("header:"):
                for prior_ref,prior_cap,prior_params in assignments:
                    if prior_ref["id"]!=ref["id"]: continue
                    if cap.property.startswith("header:") and prior_params.get("variant")!=params.get("variant"): continue
                    overlap=max(prior_ref.get("start",0),ref.get("start",0))<min(prior_ref.get("end",10**12),ref.get("end",10**12))
                    if overlap and prior_cap.property==cap.property and not equal_value(prior_params,params):
                        raise FormatError(f"Overlapping assignments for {ref['id']} / {cap.property}","needs_clarification")
                    if {prior_cap.property,cap.property} in ({"first_line_indent","hanging_indent"},{"header:first_line_indent","header:hanging_indent"}):
                        raise FormatError("First-line and hanging indents are mutually exclusive; request one", "needs_clarification")
                assignments.append((ref,cap,params))
        oid=item.get("id",f"op-{i+1:04d}")
        if not isinstance(oid,str) or any(o["id"]==oid for o in resolved): raise FormatError("Operation IDs must be unique strings")
        resolved.append({"id":oid,"action":cap.action,"params":params,"targets":targets,"effects":list(cap.effects)})
        normalized.append({"id":oid,"action":cap.action,"target":copy.deepcopy(item["target"]),"params":params})
    normalized_request={**request,"schema_version":SCHEMA_VERSION,"operations":normalized}
    plan={"schema_version":SCHEMA_VERSION,"kind":"word-format-plan","origin":request.get("origin","direct"),
          "source":{"path":str(source),"sha256":sha256(source)},"output_path":str(destination),
          "request":normalized_request,"operations":resolved,
          "preservation":{"unrequested_properties":"preserve","unselected_objects":"preserve","content":"preserve except declared structural effects"}}
    plan["plan_sha256"]=json_hash(plan)
    return plan


def validate_plan(plan):
    require_keys(plan,{"schema_version","kind","origin","source","output_path","request","operations","preservation","plan_sha256"},
                 {"schema_version","kind","source","output_path","request","operations","plan_sha256"},"plan")
    if plan["kind"]!="word-format-plan" or plan["schema_version"]!=SCHEMA_VERSION: raise FormatError("Unsupported plan")
    actual={k:v for k,v in plan.items() if k!="plan_sha256"}
    if json_hash(actual)!=plan["plan_sha256"]: raise FormatError("Plan changed after creation; build a new plan")
    source=source_path(plan["source"]["path"])
    if sha256(source)!=plan["source"]["sha256"]: raise FormatError("Source changed after planning; inspect and plan again")
    regenerated=build_plan(source,plan["output_path"],plan["request"])
    if regenerated["operations"]!=plan["operations"]: raise FormatError("Resolved targets no longer match the request")
    return source,output_path(source,plan["output_path"])


def _read(cap,index,ref):
    if ref["kind"]=="style" and ref["id"] not in index.elements: return None
    return cap.reader(index,ref,cap)


def _preservation(index,refs,cap):
    return atomic_projection(index.document,index,refs,cap)


def _format_verification(document,checks):
    index=DocumentIndex(document)
    failures=[];output=[]
    for check in checks:
        cap=REGISTRY[check["action"]]
        ref=check["output_target"]
        if ref["id"] not in index.records:
            failures.append(f"Target missing: {ref['id']}");continue
        actual=_read(cap,index,ref)
        passed=cap.verifier(actual,check["params"],cap)
        output.append({**check,"after":actual,"status":"passed" if passed else "failed"})
        if not passed: failures.append(f"{check['operation_id']} {ref['id']}: requested={check['params']!r}, actual={actual!r}")
    return {"status":"failed" if failures else "passed","checks":output,"errors":failures}


def _remap_checks(document,index,checks):
    final=DocumentIndex(document)
    by_element={node:oid for oid,node in final.elements.items()}
    for check in checks:
        ref=check["target"]
        node=index.elements.get(ref["id"])
        if ref["kind"]=="document": oid="document"
        elif ref["kind"]=="style": oid=ref["id"]
        else: oid=by_element.get(node)
        if oid is None: raise FormatError(f"An operation removed a requested target {ref['id']}")
        check["output_target"]={**ref,"id":oid}


def apply_plan(plan,*,overwrite=False):
    source,destination=validate_plan(plan)
    if destination.exists() and not overwrite: raise FormatError("Output already exists; choose a new filename or explicitly request overwrite")
    original_hash=sha256(source)
    document=Document(source)
    index=DocumentIndex(document)
    checks=[];changes=[];notes=[]
    for item in plan["operations"]:
        cap=REGISTRY[item["action"]]
        local=[]
        before_projection=_preservation(index,item["targets"],cap) if not cap.structural else None
        structural_before=structural_snapshot(index,cap,item["params"],item["targets"]) if cap.structural else None
        for ref in item["targets"]:
            before=_read(cap,index,ref)
            if before is not None and ref["kind"]!="style" and not cap.structural and cap.verifier(before,item["params"],cap):
                local.append({"operation_id":item["id"],"action":item["action"],"target":ref,"params":item["params"],"before":before,"changed":False})
                continue
            result=cap.handler(index,ref,item["params"])
            if result: notes.extend(result)
            local.append({"operation_id":item["id"],"action":item["action"],"target":ref,"params":item["params"],"before":before,"changed":True})
        if before_projection is not None:
            after_projection=_preservation(index,item["targets"],cap)
            bad=[p for p in before_projection if before_projection[p]!=after_projection.get(p)]
            if set(before_projection)!=set(after_projection) or bad:
                raise FormatError(f"{item['id']} changed unrelated properties or package parts: {bad}")
        if structural_before is not None:
            verify_structural_preservation(index,cap,item["params"],item["targets"],structural_before)
        checks.extend(local)
        changes.append({"id":item["id"],"action":item["action"],"matched":len(local),"modified":sum(c["changed"] for c in local),"effects":item["effects"]})
    check_page_geometry(document)
    _remap_checks(document,index,checks)
    destination.parent.mkdir(parents=True,exist_ok=True)
    temporary=destination.with_name(f".{destination.stem}.{uuid.uuid4().hex}.docx")
    try:
        document.save(temporary)
        reopened=Document(temporary)
        format_result=_format_verification(reopened,checks)
        from verify_output import verify_structure
        structure=verify_structure(temporary,plan["request"].get("verification_spec",{}),check_format=False)
        if format_result["status"]!="passed" or structure["status"]!="passed":
            raise FormatError("Output verification failed: "+"; ".join(format_result["errors"]+structure["errors"]))
        if sha256(source)!=original_hash: raise FormatError("Source changed during execution; output not published")
        if destination.exists() and not overwrite: raise FormatError("Output appeared during execution; output not replaced")
        os.replace(temporary,destination)
    finally:
        if temporary.exists(): temporary.unlink()
    verification=plan["request"].get("verification",{})
    result={"schema_version":SCHEMA_VERSION,"status":"passed" if any(c["changed"] for c in checks) else "no_change",
            "input":str(source),"output":str(destination),"output_sha256":sha256(destination),"plan_sha256":plan["plan_sha256"],
            "operations":changes,"checks":format_result["checks"],"changes":notes or [f"Applied {c['action']} to {c['matched']} target(s)" for c in changes],
            "verification":{"format":format_result,"structure":structure,"preservation":{"status":"passed","method":"per-operation package projection; structural changes declared separately"},
                            "visual":{"status":"pending" if verification.get("visual_enabled") else "skipped","reason":"Awaiting explicit rendering and page review" if verification.get("visual_enabled") else "Visual review not requested"}},
            "warnings":["Fields need the selected Word/LibreOffice renderer to evaluate; a structural pass does not assert pagination."]}
    return result


def _validate_report_checks(plan, report):
    """A successful recheck must cover the entire plan, including no-ops."""
    if not isinstance(report, dict): raise FormatError("Report must be an object")
    if report.get("status") not in {"passed", "no_change"}:
        raise FormatError("Report must be a successful apply result")
    checks = report.get("checks")
    if not isinstance(checks, list): raise FormatError("Report checks must be an array covering the complete plan")
    fields = ("operation_id", "action", "target", "params")
    expected = Counter(json_hash({"operation_id": item["id"], "action": item["action"],
                                  "target": ref, "params": item["params"]})
                       for item in plan["operations"] for ref in item["targets"])
    shifts_ids = any(item["action"] in {"toc.configure", "heading.numbering.configure",
                                              "section.break.next_page.insert", "caption.position.configure"}
                     for item in plan["operations"])
    actual = Counter()
    for check in checks:
        if not isinstance(check, dict) or any(field not in check for field in fields):
            raise FormatError("Each report check must identify its operation, action, source target and parameters")
        actual[json_hash({field: check[field] for field in fields})] += 1
        ref = check.get("output_target")
        target = check["target"]
        if not isinstance(ref, dict) or not isinstance(target, dict) or not isinstance(ref.get("id"), str):
            raise FormatError("Each report check must have a valid output_target")
        # Structural operations may shift paragraph IDs. All other selector
        # metadata must still belong to the source-bound resolved target.
        if {k: v for k, v in ref.items() if k != "id"} != {k: v for k, v in target.items() if k != "id"}:
            raise FormatError("Report output_target does not match its source target")
        if not shifts_ids and ref["id"] != target["id"]:
            raise FormatError("Report retargeted an operation without a structural ID shift")
    if actual != expected:
        raise FormatError("Report checks do not cover the complete plan; missing, duplicate or changed checks")
    return checks


def verify_plan(plan,report):
    validate_plan(plan)
    checks = _validate_report_checks(plan, report)
    if report.get("plan_sha256")!=plan["plan_sha256"]: raise FormatError("Report does not belong to this plan")
    output=Path(plan["output_path"])
    if not output.is_file() or sha256(output)!=report.get("output_sha256"): raise FormatError("Output changed after execution")
    result=_format_verification(Document(output),checks)
    from verify_output import verify_structure
    structure = verify_structure(output, plan["request"].get("verification_spec", {}), check_format=False)
    status = "passed" if result["status"] == structure["status"] == "passed" else "failed"
    return {"schema_version":SCHEMA_VERSION,"status":status,"format":result,"structure":structure,"output":str(output)}


def audit_document(input_path, request):
    """Read requested effective values using the same selectors and verifiers."""
    source = source_path(input_path)
    source_hash = sha256(source)
    document = Document(source)
    # Planning resolves/validates the request but performs no writes. This
    # internal destination is never created or exposed as an audit output.
    plan = build_plan(source, source.with_name(f".{source.stem}.audit-only.docx"), request, _document=document)
    index = DocumentIndex(document)
    checks = []
    for item in plan["operations"]:
        cap = REGISTRY[item["action"]]
        for ref in item["targets"]:
            actual = _read(cap, index, ref)
            passed = cap.verifier(actual, item["params"], cap)
            record = index.records.get(ref["id"], {})
            checks.append({"operation_id": item["id"], "action": item["action"], "target": ref,
                           "text": record.get("text", ""), "expected": item["params"], "actual": actual,
                           "status": "passed" if passed else "failed"})
    from verify_output import verify_structure
    structure = verify_structure(source, plan["request"].get("verification_spec", {}), check_format=False)
    if sha256(source) != source_hash: raise FormatError("Source changed during audit; inspect and audit again")
    failed = sum(check["status"] == "failed" for check in checks)
    status = "failed" if failed or structure["status"] == "failed" else "passed" if checks else "no_change"
    return {"schema_version": SCHEMA_VERSION, "kind": "word-format-audit", "status": status,
            "source": {"path": str(source), "sha256": source_hash}, "checks": checks,
            "summary": {"checked": len(checks), "passed": len(checks) - failed, "failed": failed},
            "structure": structure, "visual": {"status": "skipped", "reason": "Read-only audit does not render pages"},
            "warnings": ["Only explicitly requested properties are checked; structural agreement does not prove pagination."]}


def apply_spec(input_path,output,spec,*,clear_direct=False,verification=None,overwrite=False,origin="legacy"):
    request={"schema_version":SCHEMA_VERSION,"origin":origin,"spec":spec,"clear_direct_font_formatting":clear_direct,"verification":verification or {}}
    plan=build_plan(input_path,output,request)
    return apply_plan(plan,overwrite=overwrite)
