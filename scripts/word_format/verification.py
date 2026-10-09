"""Independent output assertions and change-boundary checks."""

from __future__ import annotations

import copy
import re
from collections import Counter

from docx.oxml.ns import qn

from .contracts import FormatError
from .properties import FONT_TAGS, PARAGRAPH_TAGS, sanitize_properties, xml_value


def field_codes(document):
    return [" ".join(n.text.split()) for part in document.part.package.parts if hasattr(part,"element")
            for n in part.element.iter(qn("w:instrText")) if n.text]


def read_complex_state(index,name):
    doc=index.document
    if name=="headers_footers":
        state={"first_page":[s.different_first_page_header_footer for s in doc.sections],
               "odd_even":doc.settings.odd_and_even_pages_header_footer,"managed":{}}
        for kind in ("header","footer"):
            state["managed"][kind]=[]
            for part in doc.part.package.parts:
                if str(part.partname).startswith("/word/"+kind) and hasattr(part,"element"):
                    for p in part.element.iter(qn("w:p")):
                        ps=p.find("./"+qn("w:pPr")+"/"+qn("w:pStyle"))
                        if ps is not None and ps.get(qn("w:val"))=="WFM"+kind.title():
                            state["managed"][kind].append("".join(n.text or "" for n in p.iter(qn("w:t"))))
        return state
    if name=="page_numbers":
        parts={}
        for kind in ("header","footer"):
            for i,s in enumerate(doc.sections):
                for prefix in ("","first_page_","even_page_"):
                    container=getattr(s,prefix+kind)
                    # Access only existing parts; inspection must not create them.
                    effective=container
                    while effective is not None and not effective._has_definition:
                        effective=effective._prior_headerfooter
                    if effective is not None:
                        codes=[n.text or "" for n in effective._element.iter(qn("w:instrText"))]
                        parts[f"{i}:{prefix}{kind}"]=sum(bool(re.search(r"\bPAGE\b",c,re.I)) for c in codes)
        starts=[s._sectPr.find(qn("w:pgNumType")) for s in doc.sections]
        managed=sum(p.find("./"+qn("w:pPr")+"/"+qn("w:pStyle")) is not None and
                    p.find("./"+qn("w:pPr")+"/"+qn("w:pStyle")).get(qn("w:val"))=="WFMPageNumber" and
                    any(re.search(r"\bPAGE\b",n.text or "",re.I) for n in p.iter(qn("w:instrText")))
                    for part in doc.part.package.parts if hasattr(part,"element") for p in part.element.iter(qn("w:p")))
        return {"parts":parts,"first_page":[s.different_first_page_header_footer for s in doc.sections],"odd_even":doc.settings.odd_and_even_pages_header_footer,
                "managed":managed,"starts":[n.get(qn("w:start")) if n is not None else None for n in starts],"codes":field_codes(doc)}
    if name=="table_of_contents":
        managed=[p for p in doc.paragraphs if p.style and p.style.name.startswith("WFM TOC")]
        trailing=[p for p in managed if p.style.name=="WFM TOC Break" and p._p.getprevious() is not None and any(
            n.tag==qn("w:instrText") and n.text and "TOC " in n.text for n in p._p.getprevious().iter())]
        return {"titles":[p.text for p in managed if p.style.name=="WFM TOC Heading"],
                "codes":[n.text or "" for p in managed for n in p._p.iter(qn("w:instrText"))],
                "trailing_break":bool(trailing),"heading_count":sum(p.style.name.startswith("Heading ") and bool(p.text.strip()) for p in doc.paragraphs)}
    if name in {"heading_numbering","numbered","reference_numbering"}:
        from apply_spec import _get_numbering_element, _num_format_for_num_id
        numbering=_get_numbering_element(doc)
        records=[]
        for p in doc.paragraphs:
            np=p._p.find("./"+qn("w:pPr")+"/"+qn("w:numPr"))
            ni=np.find(qn("w:numId")) if np is not None else None
            il=np.find(qn("w:ilvl")) if np is not None else None
            nid=ni.get(qn("w:val")) if ni is not None else None
            marker=None;start=None
            if numbering is not None and nid not in {None,"0"}:
                num=next((n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId"))==nid),None)
                abstract_ref=num.find(qn("w:abstractNumId")) if num is not None else None
                aid=abstract_ref.get(qn("w:val")) if abstract_ref is not None else None
                abstract=next((n for n in numbering.findall(qn("w:abstractNum")) if n.get(qn("w:abstractNumId"))==aid),None)
                lvl=next((n for n in abstract.findall(qn("w:lvl")) if n.get(qn("w:ilvl"))==(il.get(qn("w:val")) if il is not None else "0")),None) if abstract is not None else None
                marker_node=lvl.find(qn("w:lvlText")) if lvl is not None else None
                start_node=lvl.find(qn("w:start")) if lvl is not None else None
                marker=marker_node.get(qn("w:val")) if marker_node is not None else None
                start=start_node.get(qn("w:val")) if start_node is not None else None
            records.append({"text":p.text,"style":p.style.name,"num_id":nid,"level":int(il.get(qn("w:val")))+1 if il is not None else 1,"marker":marker,"start":start})
        return records
    if name=="caption_position":
        from apply_spec import _neighboring_content_element, _is_figure_paragraph
        states=[]
        for p in doc.paragraphs:
            if re.match(r"^\s*(?:图|表|Figure|Fig\.|Table)\s*\d",p.text,re.I):
                before=_neighboring_content_element(doc,p._p,-1)
                after=_neighboring_content_element(doc,p._p,1)
                states.append({"text":p.text,"before_table":before is not None and before.tag==qn("w:tbl"),"after_table":after is not None and after.tag==qn("w:tbl"),
                               "before_figure":before is not None and _is_figure_paragraph(before),"after_figure":after is not None and _is_figure_paragraph(after)})
        return states
    if name=="citations":
        return [p.text for p in doc.paragraphs if not p.style.name.startswith("WFM TOC")]
    raise FormatError(f"Missing structural reader for {name}")


def verify_complex_state(actual,params,name):
    if name=="headers_footers":
        if "different_first_page" in params and any(v!=params["different_first_page"] for v in actual["first_page"]): return False
        if "different_odd_even" in params and actual["odd_even"]!=params["different_odd_even"]: return False
        for kind in ("header","footer"):
            c=params.get(kind,{})
            if c.get("enabled") and str(c.get("text","")).strip():
                if not actual["managed"][kind] or any(t!=c.get("text","") for t in actual["managed"][kind]): return False
            elif "enabled" in c and any(actual["managed"][kind]): return False
        return True
    if name=="page_numbers":
        if not params.get("enabled"):
            return actual["managed"]==0
        kind=params.get("location","footer")
        keys=[f"{i}:{kind}" for i in range(len(actual["starts"]))]
        if params.get("show_on_first_page",True):
            keys.extend(f"{i}:first_page_{kind}" for i,enabled in enumerate(actual["first_page"]) if enabled)
        if actual["odd_even"]: keys.extend(f"{i}:even_page_{kind}" for i in range(len(actual["starts"])))
        counts=[actual["parts"].get(k,0) for k in keys]
        return bool(counts) and all(c==1 for c in counts) and actual["starts"][0]==str(params.get("start",1)) and (params.get("format")!="page-x-of-y" or any("NUMPAGES" in c for c in actual["codes"]))
    if name=="table_of_contents":
        if not params.get("enabled"): return not actual["titles"] and not actual["codes"]
        return actual["heading_count"]>0 and params.get("title","目录") in actual["titles"] and any(
            f'\\o "1-{params.get("max_heading_level",3)}"' in c for c in actual["codes"]) and actual["trailing_break"]==params.get("page_break_after",True)
    if name=="heading_numbering":
        for level in (1,2,3):
            mode=params.get(f"level{level}","keep")
            if mode=="keep": continue
            candidates=[r for r in actual if r["style"]==f"Heading {level}" and r["num_id"] not in {None,"0"}]
            # No headings of this level is a no-op, not permission to invent one.
            if any(r["marker"]!=({1:"第%1章" if mode=="chinese" else "%1、",2:"%1.%2",3:"%1.%2.%3"}[level]) for r in candidates): return False
        return True
    if name=="numbered":
        from config import LIST_NUMBERING_STYLES
        items=[r for r in actual if r["num_id"] not in {None,"0"} and not r["style"].startswith("Heading") and r["style"]!="Bibliography"]
        return all(r["marker"]==LIST_NUMBERING_STYLES[params.get("style","decimal-period")][1] and r["start"]==str(params.get("start",1)) for r in items)
    if name=="reference_numbering":
        items=[r for r in actual if r["style"]=="Bibliography"]
        return all(r["marker"]=="[%1]" if params["numbering_mode"]=="word-numbering" else r["num_id"] in {None,"0"} for r in items)
    if name=="caption_position":
        label=params.get("label") or ("图" if params["role"]=="figure" else "表")
        items=[r for r in actual if r["text"].strip().startswith(label)]
        target_key=("after_" if params["position"]=="above" else "before_")+params["role"]
        return all(r[target_key] for r in items)
    if name=="citations":
        return all(0<=p["paragraph_index"]<len(actual) and p["citation_text"] in actual[p["paragraph_index"]] for p in params["placements"])
    return False


def atomic_projection(document,index,refs,cap):
    """Compare all package parts, masking only the requested properties/range.

    Text runs are flattened to character tokens so harmless splitting does not
    conceal bookmark/hyperlink/object changes or adjacent-character formatting.
    """
    targets={index.elements.get(r["id"]):r for r in refs if index.elements.get(r["id"]) is not None}
    new_style_names={r["id"][6:] for r in refs if r["kind"]=="style" and r.get("xml_sha256") is None}
    if cap.property=="style":
        new_style_names={r["new_style_name"] for r in refs if "new_style_name" in r}

    def paragraph_projection(node,ref):
        prop=cap.property
        pr=node.find(qn("w:pPr"))
        sanitized=sanitize_properties(pr,prop) if prop in PARAGRAPH_TAGS else xml_value(pr) if pr is not None and len(pr) else ""
        if prop in FONT_TAGS and pr is not None and pr.find(qn("w:rPr")) is not None:
            clone=copy.deepcopy(pr);rpr=clone.find(qn("w:rPr"))
            masked=sanitize_properties(rpr,prop);clone.remove(rpr)
            if masked: clone.append(_masked_node("rPr",masked))
            sanitized=xml_value(clone) if len(clone) else ""
        if prop=="style":
            clone=copy.deepcopy(pr)
            if clone is not None:
                st=clone.find(qn("w:pStyle"))
                if st is not None: clone.remove(st)
            sanitized=xml_value(clone) if clone is not None and len(clone) else ""
        offset=0
        def content(child):
            nonlocal offset
            if child.tag==qn("w:r"):
                rpr=child.find(qn("w:rPr"));tokens=[]
                for sub in child:
                    if sub.tag==qn("w:rPr"): continue
                    if sub.tag==qn("w:t"):
                        for char in sub.text or "":
                            selected=ref.get("start",0)<=offset<ref.get("end",10**12)
                            fmt=sanitize_properties(rpr,prop) if selected and prop in FONT_TAGS else (xml_value(rpr) if rpr is not None and len(rpr) else "")
                            tokens.append(("char",char,fmt,tuple(sorted(child.attrib.items()))));offset+=1
                    else: tokens.append(("object",xml_value(sub),xml_value(rpr)))
                return tokens
            if child.tag==qn("w:pPr"): return []
            if len(child):
                return [("container",child.tag,tuple(sorted(child.attrib.items())),[item for sub in child for item in content(sub)])]
            return [("node",xml_value(child))]
        return (tuple(sorted(node.attrib.items())),sanitized,[item for child in node for item in content(child)])

    def tree(node):
        if node.tag==qn("w:style") and node.find(qn("w:name")) is not None:
            if node.find(qn("w:name")).get(qn("w:val")) in new_style_names: return None
        ref=targets.get(node)
        if ref is not None:
            if ref["kind"] in {"paragraph","text_range"}: return ("selected-paragraph",paragraph_projection(node,ref))
            if ref["kind"]=="style":
                clone=copy.deepcopy(node)
                for tag in ("rPr","pPr"):
                    pr=clone.find(qn("w:"+tag))
                    if pr is not None:
                        masked=sanitize_properties(pr,cap.property)
                        clone.remove(pr)
                        if masked: clone.append(_masked_node(tag,masked))
                return xml_value(clone)
            clone=copy.deepcopy(node)
            if cap.property.startswith("page:"):
                attr=cap.property[5:];tag="pgSz" if attr in {"w","h","orient"} else "pgMar"
                pr=clone.find(qn("w:"+tag))
                if pr is not None:
                    pr.attrib.pop(qn("w:"+attr),None)
                    if not pr.attrib: clone.remove(pr)
            elif cap.property=="table_alignment":
                pr=clone.find(qn("w:tblPr"));jc=pr.find(qn("w:jc")) if pr is not None else None
                if jc is not None: pr.remove(jc)
            elif cap.property=="repeat_header":
                pr=clone.find("./"+qn("w:tr")+"/"+qn("w:trPr"))
                hd=pr.find(qn("w:tblHeader")) if pr is not None else None
                if hd is not None: pr.remove(hd)
                if pr is not None and not len(pr) and not pr.attrib: pr.getparent().remove(pr)
            return xml_value(clone)
        # Newly created styles from assigning a name are not unrelated changes.
        if node.tag==qn("w:style") and node.find(qn("w:name")) is not None:
            name=node.find(qn("w:name")).get(qn("w:val"))
            if name in new_style_names: return None
        return (node.tag,tuple(sorted(node.attrib.items())),node.text or "",tuple(v for child in node if (v:=tree(child)) is not None))
    result={}
    for part in document.part.package.parts:
        name=str(part.partname)
        result[name]=tree(part.element) if hasattr(part,"element") else part.blob
    return result


def _masked_node(tag,value):
    from lxml import etree
    node=etree.Element("masked-"+tag)
    node.text=value
    return node


def structural_snapshot(index,cap,params,refs):
    """Protect body content/objects and unrelated parts across structural writes."""
    from .inspect import paragraph_text
    paragraphs=[]
    for p in index.document.element.body.iter(qn("w:p")):
        ps=p.find("./"+qn("w:pPr")+"/"+qn("w:pStyle"))
        style=ps.get(qn("w:val"),"") if ps is not None else ""
        text=paragraph_text(p)
        if cap.property=="table_of_contents" and (style.startswith("WFMTOC") or re.fullmatch(r"\s*(?:目\s*录|table\s+of\s+contents|contents)\s*",text,re.I)):
            continue
        clone=copy.deepcopy(p)
        if cap.property=="citations":
            objects=[xml_value(n) for n in clone.iter() if n.tag in {qn("w:bookmarkStart"),qn("w:bookmarkEnd"),qn("w:drawing"),qn("w:instrText"),qn("w:fldChar")}]
            paragraphs.append((p,text,repr(objects)))
            continue
        if cap.property in {"heading_numbering","reference_numbering","numbered"}:
            pr=clone.find(qn("w:pPr"))
            if pr is not None:
                for child in list(pr):
                    if child.tag in {qn("w:numPr"),qn("w:pStyle")} or (cap.property=="numbered" and child.tag==qn("w:ind")):
                        pr.remove(child)
                if not len(pr) and not pr.attrib: clone.remove(pr)
        for n in clone.iter(qn("w:t")): n.text=""
        paragraphs.append((p,text,xml_value(clone)))
    allowed={"/word/document.xml"}
    if cap.property in {"heading_numbering","reference_numbering","numbered"}: allowed.update({"/word/numbering.xml","/word/styles.xml"})
    if cap.property in {"headers_footers","page_numbers","table_of_contents"}: allowed.update({"/word/styles.xml","/word/settings.xml"})
    parts={str(p.partname):p.blob for p in index.document.part.package.parts
           if str(p.partname) not in allowed and not (cap.property in {"headers_footers","page_numbers"} and re.fullmatch(r"/word/(header|footer)\d+\.xml",str(p.partname)))}
    return paragraphs,parts


def verify_structural_preservation(index,cap,params,refs,before):
    from .inspect import paragraph_text
    after=structural_snapshot(index,cap,params,refs)
    by_node={p:(text,xml) for p,text,xml in after[0]}
    for node,text,xml in before[0]:
        if node not in by_node: raise FormatError("Structural operation removed unrelated body content")
        new_text,new_xml=by_node[node]
        if xml!=new_xml: raise FormatError("Structural operation changed an unrelated body object or format")
        if text==new_text: continue
        if cap.property in {"heading_numbering","reference_numbering"} and text.endswith(new_text):
            removed=text[:len(text)-len(new_text)] if new_text else text
            if re.fullmatch(r"\s*(?:第[一二三四五六七八九十百千零〇0-9]+章|\d+(?:\.\d+){0,2}[、.．]?|\[\d+\]|［\d+］)\s*",removed): continue
        if cap.property=="citations":
            placements=params.get("placements",[])
            old_index=next((i for i,p in enumerate(index.document.paragraphs) if p._p is node),None)
            matches=[p for p in placements if p["paragraph_index"]==old_index]
            if len(matches)==1 and new_text.replace(matches[0]["citation_text"],"",1).strip()==text.strip(): continue
        raise FormatError("Structural operation changed body text outside its declared transformation")
    if set(by_node)-{p for p,_,_ in before[0]}:
        raise FormatError("Structural operation inserted unrelated body paragraphs")
    for name,blob in before[1].items():
        if after[1].get(name)!=blob: raise FormatError(f"Structural operation changed unrelated part: {name}")


def check_page_geometry(document):
    for i,s in enumerate(document.sections,1):
        if s.page_width is None or s.page_height is None:
            raise FormatError(f"Section {i} has no dimensions")
        if s.left_margin+s.right_margin+(s.gutter or 0)>=s.page_width or s.top_margin+s.bottom_margin>=s.page_height:
            raise FormatError(f"Section {i} margins leave no content area")
