"""Section-scoped header operations and independent OOXML readback."""

import copy
import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.section import Section

from .contracts import FormatError, require_keys, text_value
from .properties import (element, read_font, read_paragraph,
                         write_font, write_paragraph, xml_value, FONT_TAGS)

VARIANTS = {"default": "header", "first": "first_page_header", "even": "even_page_header"}


def variant_params(params, validator):
    variant = params.get("variant", "default")
    if variant not in VARIANTS:
        raise FormatError("Header variant must be default, first or even")
    return {**validator({k: v for k, v in params.items() if k != "variant"}), "variant": variant}


def content_params(params):
    require_keys(params, {"mode", "text", "style_name", "prefix", "variant"}, {"mode"}, "header content")
    variant = params.get("variant", "default")
    if variant not in VARIANTS:
        raise FormatError("Header variant must be default, first or even")
    if params["mode"] == "text":
        require_keys(params, {"mode", "text", "variant"}, {"mode", "text"}, "literal header")
        return {"mode": "text", "text": text_value(params["text"], "text", True), "variant": variant}
    if params["mode"] != "chapter_title":
        raise FormatError("Header mode must be text or chapter_title")
    require_keys(params, {"mode", "style_name", "prefix", "variant"}, {"mode", "style_name"}, "chapter header")
    name = text_value(params["style_name"], "style_name")
    if any(c in name for c in ('"', '\\', '\n', '\r')):
        raise FormatError("Unsafe STYLEREF style name")
    return {"mode": "chapter_title", "style_name": name,
            "prefix": text_value(params.get("prefix", ""), "prefix", True), "variant": variant}


def section(index, ref):
    return Section(index.elements[ref["id"]], index.document.part)


def effective_root(container):
    while container is not None and not container._has_definition:
        container = container._prior_headerfooter
    return container._element if container is not None else None


def detach(container):
    source = container
    while source is not None and not source._has_definition:
        source = source._prior_headerfooter
    root = source._element if source is not None else None
    relationships = list(source.part.rels.values()) if source is not None else []
    saved = copy.deepcopy(root) if root is not None else None
    # Always create a new part: two explicit references may share the same rId.
    if container._has_definition:
        container._drop_definition()
    container._add_definition()
    if saved is None:
        container._element.attrib.clear()
        container._element[:] = [OxmlElement("w:p")]
    if saved is not None:
        container._element.attrib.clear()
        container._element.attrib.update(saved.attrib)
        ids = {rel.rId: container.part.relate_to(rel.target_ref if rel.is_external else rel.target_part,
                                               rel.reltype, rel.is_external) for rel in relationships}
        for node in saved.iter():
            for attr, value in list(node.attrib.items()):
                if attr.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}") and value in ids:
                    node.set(attr, ids[value])
        container._element[:] = list(saved)


def editable_header(index, ref, variant):
    current = section(index, ref)
    sections = list(index.document.sections)
    pos = next(i for i, s in enumerate(sections) if s._sectPr is current._sectPr)
    # Freeze the next inherited header before editing this section's definition.
    if pos + 1 < len(sections):
        following = getattr(sections[pos + 1], VARIANTS[variant])
        if following.is_linked_to_previous:
            detach(following)
    container = getattr(current, VARIANTS[variant])
    detach(container)
    return container


def set_content(index, ref, params):
    if params["mode"] == "chapter_title":
        name = params["style_name"]
        if name not in index.document.styles or index.document.styles[name].type != 1:
            raise FormatError("STYLEREF requires an existing paragraph style", "needs_clarification")
        sid = index.document.styles[name].style_id
        if not any(p.find("./" + qn("w:pPr") + "/" + qn("w:pStyle")) is not None and
                   p.find("./" + qn("w:pPr") + "/" + qn("w:pStyle")).get(qn("w:val")) == sid
                   for p in index.document.element.body.iter(qn("w:p"))):
            raise FormatError("STYLEREF style has no body heading to reference", "needs_clarification")
    container = editable_header(index, ref, params["variant"])
    original = next(container._element.iter(qn("w:p")), None)
    saved_pr = copy.deepcopy(original.find(qn("w:pPr"))) if original is not None else None
    first_run = next(original.iter(qn("w:r")), None) if original is not None else None
    saved_rpr = copy.deepcopy(first_run.find(qn("w:rPr"))) if first_run is not None else None
    for child in list(container._element):
        container._element.remove(child)
    p = container.add_paragraph()
    if saved_pr is not None: p._p.insert(0, saved_pr)
    p._p.set("{urn:word-format-master}managed-header", "true")
    if params["mode"] == "text":
        p.add_run(params["text"])
    else:
        p.add_run(params["prefix"])
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), f'STYLEREF "{params["style_name"]}"')
        field.set(qn("w:dirty"), "true")
        field.append(OxmlElement("w:r"))
        p._p.append(field)
    if saved_rpr is not None:
        for run in p._p.iter(qn("w:r")): run.insert(0, copy.deepcopy(saved_rpr))


def set_link(index, ref, params):
    container = editable_header(index, ref, params["variant"])
    if params["value"]:
        if section(index, ref)._sectPr is index.document.sections[0]._sectPr:
            raise FormatError("The first section has no preceding header", "needs_clarification")
        container.is_linked_to_previous = True


def set_first_page(index, ref, params):
    section(index, ref).different_first_page_header_footer = params["value"]


def set_odd_even(index, ref, params):
    index.document.settings.odd_and_even_pages_header_footer = params["value"]


def set_format(index, ref, params, name):
    container = getattr(section(index, ref), VARIANTS[params["variant"]])
    root = effective_root(container)
    if root is None or not list(root.iter(qn("w:p"))):
        raise FormatError("Set header content before setting its format", "needs_clarification")
    container = editable_header(index, ref, params["variant"])
    for p in container._element.iter(qn("w:p")):
        if name in FONT_TAGS:
            for run in p.iter(qn("w:r")):
                write_font(element(run, "rPr"), name, params["value"])
            write_font(element(element(p, "pPr"), "rPr"), name, params["value"])
        else:
            write_paragraph(element(p, "pPr"), name, {k:v for k,v in params.items() if k!="variant"} if name=="line_spacing" or name.endswith("indent") else params["value"])


def set_east_asia(index, ref, params): set_format(index, ref, params, "font_east_asia")
def set_latin(index, ref, params): set_format(index, ref, params, "font_latin")
def set_size(index, ref, params): set_format(index, ref, params, "font_size_pt")
def set_alignment(index, ref, params): set_format(index, ref, params, "alignment")
def set_bold(index, ref, params): set_format(index, ref, params, "bold")
def set_italic(index, ref, params): set_format(index, ref, params, "italic")
def set_color(index, ref, params): set_format(index, ref, params, "color")
def set_spacing_before(index, ref, params): set_format(index, ref, params, "space_before_pt")
def set_spacing_after(index, ref, params): set_format(index, ref, params, "space_after_pt")
def set_line_spacing(index, ref, params): set_format(index, ref, params, "line_spacing")
def set_first_indent(index, ref, params): set_format(index, ref, params, "first_line_indent")
def set_left_indent(index, ref, params): set_format(index, ref, params, "left_indent")
def set_right_indent(index, ref, params): set_format(index, ref, params, "right_indent")
def set_hanging_indent(index, ref, params): set_format(index, ref, params, "hanging_indent")
def set_keep_next(index, ref, params): set_format(index, ref, params, "keep_with_next")
def set_keep_together(index, ref, params): set_format(index, ref, params, "keep_together")
def set_page_break(index, ref, params): set_format(index, ref, params, "page_break_before")
def set_widow(index, ref, params): set_format(index, ref, params, "widow_control")


def read_header(index, ref, cap):
    name = cap.property.removeprefix("header:")
    if name == "odd_even":
        return index.document.settings.odd_and_even_pages_header_footer
    if name == "first_page":
        return section(index, ref).different_first_page_header_footer
    variants = {}
    for variant, attr in VARIANTS.items():
        container = getattr(section(index, ref), attr)
        if name == "link":
            variants[variant] = container.is_linked_to_previous
            continue
        root = effective_root(container)
        if name == "content":
            literal=[]
            in_field=[False]
            def visit(node):
                for child in node:
                    if child.tag==qn("w:fldSimple"): continue
                    if child.tag==qn("w:fldChar"):
                        kind=child.get(qn("w:fldCharType"))
                        if kind=="begin": in_field[0]=True
                        elif kind=="end": in_field[0]=False
                    elif child.tag==qn("w:t") and not in_field[0]: literal.append(child.text or "")
                    elif child.tag in {qn("w:br"), qn("w:cr")} and not in_field[0]: literal.append("\n")
                    elif child.tag==qn("w:tab") and not in_field[0]: literal.append("\t")
                    visit(child)
            if root is not None: visit(root)
            variants[variant] = {"text": "".join(literal),
                                 "fields": ([n.get(qn("w:instr"), "") for n in root.iter(qn("w:fldSimple"))] +
                                            [n.text or "" for n in root.iter(qn("w:instrText"))]) if root is not None else []}
        else:
            variants[variant] = []
            if root is not None:
                for p in root.iter(qn("w:p")):
                    values = [read_font(index.document, p, r, name) for r in p.iter(qn("w:r"))] if name in FONT_TAGS else [read_paragraph(index.document, p, name)]
                    variants[variant].extend(values or [read_font(index.document, p, None, name)])
    return variants


def verify_header(actual, params, cap):
    from .registry import equal_value
    name = cap.property.removeprefix("header:")
    if name in {"first_page", "odd_even"}:
        return actual is params["value"]
    selected = actual[params["variant"]]
    if name == "content":
        if params["mode"] == "text":
            return selected == {"text": params["text"], "fields": []}
        return selected == {"text": params["prefix"], "fields": [f'STYLEREF "{params["style_name"]}"']}
    if name == "link":
        return selected is params["value"]
    expected={k:v for k,v in params.items() if k!="variant"} if name=="line_spacing" or name.endswith("indent") else params["value"]
    return bool(selected) and all(equal_value(v, expected) or (name.endswith("indent") and expected["value"]==0 and abs(v["value"])<0.025) for v in selected)


def header_projection(index, cap, params, refs):
    """Compare effective unselected headers and every footer, even after unlinking."""
    selected = {index.elements[r["id"]] for r in refs}
    name = cap.property.removeprefix("header:")
    body = copy.deepcopy(index.document.element)
    originals = list(index.document.element.iter(qn("w:sectPr")))
    for pos, s in enumerate(body.iter(qn("w:sectPr"))):
        for node in list(s):
            if node.tag == qn("w:headerReference") or (name == "first_page" and originals[pos] in selected and node.tag == qn("w:titlePg")):
                s.remove(node)
    settings = copy.deepcopy(index.document.settings.element)
    if name == "odd_even":
        for n in settings.findall(qn("w:evenAndOddHeaders")): settings.remove(n)
    values = [xml_value(body), xml_value(settings), xml_value(index.document.styles.element)]
    for s in index.document.sections:
        for variant, attr in VARIANTS.items():
            container = getattr(s, attr)
            while container is not None and not container._has_definition:
                container = container._prior_headerfooter
            root = container._element if container is not None else None
            clone = copy.deepcopy(root)
            if clone is not None:
                for node in clone.iter():
                    for attr_name, value in list(node.attrib.items()):
                        if attr_name.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}") and value in container.part.rels:
                            node.set(attr_name, container.part.rels[value].target_ref)
            if s._sectPr in selected and variant == params.get("variant"):
                if name in {"content", "link"}: clone = None
                elif clone is not None:
                    for pr in list(clone.iter(qn("w:rPr"))) + list(clone.iter(qn("w:pPr"))):
                        from .properties import FONT_TAGS, PARAGRAPH_TAGS, ATTRIBUTES
                        for tag in FONT_TAGS.get(name, PARAGRAPH_TAGS.get(name, ())):
                            node = pr.find(qn("w:" + tag))
                            if node is None: continue
                            if name in ATTRIBUTES:
                                for a in ATTRIBUTES[name].get(tag, ()): node.attrib.pop(qn("w:" + a), None)
                                if not node.attrib and not len(node): pr.remove(node)
                            else: pr.remove(node)
                    for pr in list(clone.iter(qn("w:rPr"))):
                        if not len(pr) and not pr.attrib: pr.getparent().remove(pr)
                    for pr in list(clone.iter(qn("w:pPr"))):
                        if not len(pr) and not pr.attrib: pr.getparent().remove(pr)
            if clone is not None and not clone.attrib and all(n.tag==qn("w:p") and not len(n) and not n.attrib for n in clone):
                clone=None
            values.append(xml_value(clone))
        for attr in ("footer", "first_page_footer", "even_page_footer"):
            footer = getattr(s, attr)
            while footer is not None and not footer._has_definition:
                footer = footer._prior_headerfooter
            values.append((xml_value(footer._element) if footer is not None else "",
                           sorted((rel.rId, rel.reltype, rel.target_ref, rel.is_external)
                                  for rel in footer.part.rels.values()) if footer is not None else []))
    values.append({str(p.partname): p.blob for p in index.document.part.package.parts
                   if not re.fullmatch(r"/word/(?:header|footer)\d+\.xml", str(p.partname)) and
                   str(p.partname) not in {"/word/document.xml", "/word/styles.xml", "/word/settings.xml"}})
    return values
