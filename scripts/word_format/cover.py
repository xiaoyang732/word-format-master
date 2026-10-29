"""Render cover layouts from explicit data; appearance belongs to template files."""

from __future__ import annotations

import json
from pathlib import Path

from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Mm

from .contracts import FormatError, text_value
from .namespaces import qn
from .properties import element, write_font
from .structures import set_paragraph_format, set_run_font


def apply_east_asia_font(run, font_name: str, ascii_font: str):
    write_font(element(run._r, "rPr"), "font_east_asia", font_name)
    write_font(element(run._r, "rPr"), "font_latin", ascii_font)


def _paragraph(container, config, data):
    value = data.get(config["field"])
    if not value:
        return
    paragraph = container.add_paragraph()
    set_paragraph_format(paragraph, config["tokens"])
    set_run_font(paragraph.add_run(str(value)), config["tokens"])


def render_cover(doc, layout, data, *, fields=None, use_section_break=True):
    if layout.get("kind") != "word-format-cover-layout" or layout.get("version") != "1.0":
        raise FormatError("Unsupported cover layout")
    text_value(data.get("title_cn"), "cover title")
    if "different_first_page" in layout:
        doc.sections[0].different_first_page_header_footer = layout["different_first_page"]
    for block in layout["blocks"]:
        if not block.get("when") or data.get(block["when"]):
            _paragraph(doc, block, data)
    if fields is None:
        fields = [(label, data[key]) for label, key in layout["labels"] if data.get(key)]
    if fields:
        config = layout["field_table"]
        grid = doc.add_table(rows=len(fields), cols=2)
        grid.alignment = {
            "left": WD_TABLE_ALIGNMENT.LEFT,
            "center": WD_TABLE_ALIGNMENT.CENTER,
            "right": WD_TABLE_ALIGNMENT.RIGHT,
        }[config["alignment"]]
        grid.autofit = False
        element(grid._tbl.tblPr, "tblLayout").set(qn("w:type"), "fixed")
        borders = element(grid._tbl.tblPr, "tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element(borders, side).set(qn("w:val"), "nil")
        for col, width in zip(grid.columns, config["widths_mm"], strict=True):
            col.width = Mm(width)
        for row, (label, value) in zip(grid.rows, fields, strict=True):
            for i, (content, key) in enumerate(((label, "label"), (value, "value"))):
                cell = row.cells[i]
                cell.width = Mm(config["widths_mm"][i])
                paragraph = cell.paragraphs[0]
                set_paragraph_format(paragraph, config[key])
                set_run_font(paragraph.add_run(str(content)), config[key])
            if config.get("value_border"):
                border = element(element(row.cells[1]._tc.get_or_add_tcPr(), "tcBorders"), "bottom")
                for key, value in config["value_border"].items():
                    border.set(qn("w:" + key), str(value))
    if layout.get("date"):
        _paragraph(doc, layout["date"], data)
    if use_section_break:
        section = doc.add_section(WD_SECTION_START.NEW_PAGE)
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
    else:
        doc.add_page_break()


def add_standard_thesis_cover(
    doc, *, title_cn, fields=None, use_section_break=True, layout=None, **data
):
    """Compatibility entry point selecting the independent standard layout explicitly."""
    if layout is None:
        path = Path(__file__).resolve().parents[2] / "templates/builtin/thesis-standard/cover.json"
        layout = json.loads(path.read_text(encoding="utf-8"))
    render_cover(
        doc,
        layout,
        {"title_cn": title_cn, **data},
        fields=fields,
        use_section_break=use_section_break,
    )
