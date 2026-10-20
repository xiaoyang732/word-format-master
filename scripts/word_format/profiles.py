"""Evidence-bearing format profiles compiled through the existing operation registry."""

from __future__ import annotations

import copy
import io
import re
import zipfile
from collections import defaultdict
from pathlib import Path

from docx import Document

from .contracts import (
    FormatError,
    json_hash,
    number,
    output_path,
    require_keys,
    sha256,
    source_path,
    text_value,
)
from .inspect import DocumentIndex
from .namespaces import qn
from .properties import FONT_TAGS, PARAGRAPH_TAGS, ppr_sources, read_paragraph, rpr_sources
from .registry import REGISTRY
from .requirement_rules import parse_blocks
from .requirement_sources import extract_requirements
from .spec import FONT, PAGE, PARA
from .template_structure import compare_template_structure, extract_template_structure


def _profile(source):
    return {
        "schema_version": "1.0",
        "kind": "word-format-profile",
        "source": source,
        "rules": [],
        "issues": [],
        "coverage": {},
        "snapshot": {},
    }


def _issue(profile, reason, **details):
    profile["issues"].append({"id": f"i{len(profile['issues']) + 1}", "reason": reason, **details})


def _rule(profile, candidate):
    cap = REGISTRY.get(candidate.get("action"))
    if cap is None:
        raise FormatError("Profile rule must use a registered action", "unsupported")
    if candidate.get("target", {}).get("type") not in cap.targets:
        raise FormatError(
            "Profile rule target is not supported by its registered action", "unsupported"
        )
    candidate = copy.deepcopy(candidate)
    candidate["params"] = cap.validator(candidate["params"])
    candidate["id"] = f"r{len(profile['rules']) + 1}"
    profile["rules"].append(candidate)


def requirements_profile(path):
    extracted = extract_requirements(path)
    profile = _profile(extracted["source"])
    candidates, unresolved = parse_blocks(extracted["blocks"])
    grouped = defaultdict(list)
    for candidate in candidates:
        grouped[(json_hash(candidate["target"]), candidate["action"])].append(candidate)
    for group in grouped.values():
        choices = {json_hash(c["params"]): c for c in group}
        if len(choices) > 1:
            _issue(profile, "Conflicting requirements for one property", candidates=group)
            continue
        candidate = copy.deepcopy(group[0])
        candidate["evidence"] = [e for c in group for e in c["evidence"]]
        try:
            _rule(profile, candidate)
        except FormatError as exc:
            _issue(profile, str(exc), candidates=group)
    for item in extracted["issues"] + unresolved:
        reason = item.pop("reason")
        _issue(profile, reason, **item)
    profile["snapshot"] = {"blocks": extracted["blocks"]}
    profile["coverage"] = {
        "represented_rules": len(profile["rules"]),
        "read_blocks": len(extracted["blocks"]),
        "unresolved": len(profile["issues"]),
        "policy": "No preset defaults; every unparsed phrase requires an explicit resolution",
    }
    return profile


def _template_document(path):
    if path.suffix.lower() == ".docx":
        return Document(path)
    # python-docx does not accept DOTX main content types. Read a converted
    # in-memory package; the user's template is never rewritten.
    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(buffer, "w") as target:
        for item in source.infolist():
            blob = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                blob = blob.replace(
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                )
            target.writestr(item, blob)
    buffer.seek(0)
    return Document(buffer)


def _selector(record):
    target = {"type": "paragraph", "role": record["role"], "all": True}
    if record.get("module"):
        target["module"] = record["module"]
    if record.get("heading_level"):
        target["level"] = record["heading_level"]
    if record["role"] == "table_text":
        target["part"] = "word/document.xml"
    if record["role"] == "caption":
        match = re.match(r"\s*(图|表|Figure|Fig\.|Table)", record["text"], re.I)
        if match:
            target["contains"] = match[1]
    return target


def _atomic_candidate(target, name, value, evidence):
    action = (FONT | PARA).get(name, "paragraph." + name + ".set")
    if name.endswith("indent"):
        if value["value"] < 0:
            if name != "first_line_indent":
                raise FormatError(
                    "Negative effective indent cannot be represented by this operation"
                )
            action = "paragraph.hanging_indent.set"
            value = {**value, "value": -value["value"]}
        params = value
    else:
        params = value if name == "line_spacing" else {"value": value}
    return {"target": target, "action": action, "params": params, "evidence": evidence}


def template_profile(path):
    from analyze_docx import analyze_docx

    source = Path(path).resolve()
    if source.suffix.lower() not in {".docx", ".dotx"}:
        raise FormatError("Formatting samples must be DOCX or DOTX", "unsupported")
    analysis = analyze_docx(source)
    if analysis["package"]["ooxml_conformance"] == "strict":
        raise FormatError(
            "Convert Strict OOXML to Transitional before extracting an executable profile",
            "unsupported",
        )
    profile = _profile(
        {"kind": "template", "path": str(source), "sha256": analysis["source"]["sha256"]}
    )
    document = _template_document(source)
    index = DocumentIndex(document)
    # Use the same effective readers as inspect, rather than inferred_spec's
    # representative paragraph or first-section heuristics.
    records = index.public()
    groups = defaultdict(list)
    snapshot = []
    unknown = set()
    for record in records:
        if record["kind"] != "paragraph" or record["part"] != "word/document.xml":
            continue
        p = index.elements[record["id"]]
        values = {n: index.font_values(record["id"], n) for n in FONT_TAGS}
        values.update({n: read_paragraph(document, p, n) for n in PARAGRAPH_TAGS})
        snapshot.append(
            {
                **record,
                "format": values,
                "direct_pPr": str(p.pPr.xml) if p.pPr is not None else None,
                "direct_rPr": [str(r.rPr.xml) for r in p.iter(qn("w:r")) if r.rPr is not None],
            }
        )
        properties = list(ppr_sources(document, p)) + list(rpr_sources(document, p))
        for run in p.iter(qn("w:r")):
            properties.extend(rpr_sources(document, p, run))
        for pr in properties:
            if pr is None:
                continue
            known = {
                "pStyle",
                "rStyle",
                "rFonts",
                "sz",
                "szCs",
                "b",
                "bCs",
                "i",
                "iCs",
                "color",
                "jc",
                "spacing",
                "ind",
                "keepNext",
                "keepLines",
                "pageBreakBefore",
                "widowControl",
                "sectPr",
                "outlineLvl",
                "lang",
                "noProof",
                "numPr",
                "rPr",
            }
            unknown.update(n.tag.split("}")[-1] for n in pr if n.tag.split("}")[-1] not in known)
        if record.get("unsupported_container") or not record.get("text", "").strip():
            continue
        selected = _selector(record)
        groups[json_hash(selected)].append((selected, record, values))
    for group in groups.values():
        target = group[0][0]
        for name in list(FONT_TAGS) + list(PARAGRAPH_TAGS):
            # first-line is the signed offset; the opposite zero is redundant.
            if name == "hanging_indent":
                continue
            variants = defaultdict(list)
            raw = {}
            for _, record, values in group:
                value = values[name]
                if name in FONT_TAGS:
                    if len(value) != 1:
                        value = {"mixed": value}
                    else:
                        value = value[0]
                key = json_hash(value)
                raw[key] = value
                variants[key].append(
                    {"anchor": record["id"], "style_id": record.get("style_id"), "confidence": 1.0}
                )
            choices = []
            unrepresentable = []
            for key, evidence in variants.items():
                value = raw[key]
                if value is None:
                    unrepresentable.append({"value": None, "evidence": evidence})
                    continue
                try:
                    candidate = _atomic_candidate(target, name, value, evidence)
                    candidate["params"] = REGISTRY[candidate["action"]].validator(
                        candidate["params"]
                    )
                    choices.append(candidate)
                except (FormatError, TypeError, KeyError):
                    unrepresentable.append({"value": value, "evidence": evidence})
            if len(choices) == 1 and not unrepresentable:
                _rule(profile, choices[0])
            else:
                _issue(
                    profile,
                    "Nonuniform or unrepresentable effective sample formatting",
                    target=target,
                    property=name,
                    candidates=choices,
                    observed=unrepresentable,
                )
    for section in analysis["document"]["sections"]:
        target = {"type": "section", "number": section["index"]}
        values = {k: section.get(k) for k in ("width_mm", "height_mm", "orientation")}
        for key, value in section["margins_mm"].items():
            name = {
                "header": "header_distance_mm",
                "footer": "footer_distance_mm",
                "gutter": "gutter_mm",
            }.get(key, "margin_" + key + "_mm")
            values[name] = value
        for name, value in values.items():
            if value is not None:
                _rule(
                    profile,
                    {
                        "target": target,
                        "action": PAGE[name],
                        "params": {"value": value},
                        "evidence": [{"anchor": f"section:{section['index']}", "confidence": 1.0}],
                    },
                )
        _rule(
            profile,
            {
                "target": target,
                "action": "header.first_page_different.set",
                "params": {"value": section["title_page"]},
                "evidence": [{"anchor": f"section:{section['index']}", "confidence": 1.0}],
            },
        )
        if section["columns"] != 1 or section["page_numbering"]:
            _issue(
                profile,
                "Section columns or page numbering require explicit configuration",
                section=section,
            )
        if section["section_type"] != "nextPage":
            _issue(
                profile, "Sample section boundary type requires explicit review", section=section
            )
    for i, section in enumerate(document.sections, 1):
        represented = {
            "pgSz",
            "pgMar",
            "titlePg",
            "headerReference",
            "footerReference",
            "type",
            "cols",
            "pgNumType",
        }
        extra = sorted({n.tag.split("}")[-1] for n in section._sectPr} - represented)
        if extra:
            _issue(
                profile,
                "Additional section layout properties require explicit review",
                anchor=f"section:{i}",
                properties=extra,
            )
    _rule(
        profile,
        {
            "target": {"type": "document"},
            "action": "header.odd_even_different.set",
            "params": {"value": document.settings.odd_and_even_pages_header_footer},
            "evidence": [{"anchor": "word/settings.xml#evenAndOddHeaders", "confidence": 1.0}],
        },
    )
    # Table readers can safely represent alignment and repeated header rows.
    for record in records:
        if record["kind"] != "table":
            continue
        for action in ("table.alignment.set", "table.header_repeat.set"):
            cap = REGISTRY[action]
            value = cap.reader(index, index.ref(record["id"]), cap)
            try:
                _rule(
                    profile,
                    {
                        "target": {"type": "table", "ordinal": int(record["id"][1:]) + 1},
                        "action": action,
                        "params": {"value": value},
                        "evidence": [{"anchor": record["id"], "confidence": 1.0}],
                    },
                )
            except FormatError:
                _issue(
                    profile,
                    "Table property requires explicit interpretation",
                    property=action,
                    anchor=record["id"],
                )
        cap = REGISTRY["table.three_line.apply"]
        if cap.reader(index, index.ref(record["id"]), cap):
            _rule(
                profile,
                {
                    "target": {"type": "table", "ordinal": int(record["id"][1:]) + 1},
                    "action": cap.action,
                    "params": {"value": True},
                    "evidence": [{"anchor": record["id"], "confidence": 1.0}],
                },
            )
        _issue(
            profile,
            "Table borders, widths, styles and cell geometry require review",
            anchor=record["id"],
        )
    # Empty sample variants also need a decision before target content is kept.
    _issue(
        profile,
        "Map header/footer content, formatting and fields explicitly, including empty variants; sample-specific text is not copied",
        fields=analysis["document"]["fields"],
        headers_footers=analysis["document"]["headers_footers"],
    )
    for feature in analysis["unsupported"]:
        _issue(
            profile,
            "Detected sample feature requires review or a supported explicit operation",
            feature=feature,
        )
    for feature in ("footnotes", "endnotes", "comments"):
        if analysis["features"].get(feature):
            _issue(
                profile,
                "Sample annotation formatting is recorded but not automatically migrated",
                feature=feature,
            )
    if unknown:
        _issue(
            profile,
            "Additional direct or inherited OOXML properties are recorded but not automatically migrated",
            properties=sorted(unknown),
        )
    used_ids = {r.get("style_id") for r in snapshot}
    used_styles = [s for s in document.styles if s.style_id in used_ids]
    if any(s.element.find(".//" + qn("w:numPr")) is not None for s in used_styles) or any(
        index.elements[r["id"]].find("./" + qn("w:pPr") + "/" + qn("w:numPr")) is not None
        for r in snapshot
    ):
        _issue(
            profile,
            "Used Word numbering requires semantic mapping to target headings or lists",
            numbering=analysis["document"]["numbering"],
        )
    profile["snapshot"] = {
        "paragraphs": snapshot,
        "styles": analysis["styles"],
        "used_style_ids": sorted(used_ids),
        "sections": analysis["document"]["sections"],
        "document": analysis["document"],
        "features": analysis["features"],
        "package": analysis["package"],
        "styles_xml": str(document.styles.element.xml),
        "settings_xml": str(document.settings.element.xml),
        "sections_xml": [str(s._sectPr.xml) for s in document.sections],
        "section_state": [r for r in records if r["kind"] == "section"],
        "modules": analysis["inferred_spec"].get("document_structure", {}),
        "structure_contract": extract_template_structure(document),
    }
    profile["coverage"] = {
        "represented_rules": len(profile["rules"]),
        "unresolved": len(profile["issues"]),
        "sections": len(document.sections),
        "paragraphs": len(snapshot),
        "policy": "Uniform effective role properties only; mixed values and special structures require resolution",
    }
    validate_profile_source(profile["source"])
    return profile


def validate_profile(profile):
    require_keys(
        profile,
        {"schema_version", "kind", "source", "rules", "issues", "coverage", "snapshot"},
        {"schema_version", "kind", "source", "rules", "issues"},
        "profile",
    )
    if profile["schema_version"] != "1.0" or profile["kind"] != "word-format-profile":
        raise FormatError("Unsupported profile version")
    validate_profile_source(profile["source"])
    if not isinstance(profile["rules"], list) or not isinstance(profile["issues"], list):
        raise FormatError("Profile rules and issues must be arrays")
    identifiers = set()
    for item in profile["rules"] + profile["issues"]:
        if not isinstance(item, dict):
            raise FormatError("Profile rules and issues must be objects")
        identifier = text_value(item.get("id"), "profile item ID")
        if identifier in identifiers:
            raise FormatError("Duplicate profile item ID")
        identifiers.add(identifier)
    for rule in profile["rules"]:
        require_keys(
            rule,
            {"id", "target", "action", "params", "evidence"},
            {"id", "target", "action", "params", "evidence"},
            "profile rule",
        )
        text_value(rule["action"], "profile action")
        cap = REGISTRY.get(rule["action"])
        if cap is None:
            raise FormatError("Unregistered profile action", "unsupported")
        if not isinstance(rule["target"], dict) or rule["target"].get("type") not in cap.targets:
            raise FormatError("Invalid profile rule target", "unsupported")
        cap.validator(rule["params"])
        if not isinstance(rule["evidence"], list) or not rule["evidence"]:
            raise FormatError("Profile rules require evidence")
        for evidence in rule["evidence"]:
            if not isinstance(evidence, dict):
                raise FormatError("Rule evidence must be an object")
            text_value(evidence.get("anchor"), "evidence anchor")
            number(evidence.get("confidence"), 0, 1, "evidence confidence")
    for issue in profile["issues"]:
        text_value(issue.get("reason"), "issue reason")
    for key in ("coverage", "snapshot"):
        if key in profile and not isinstance(profile[key], dict):
            raise FormatError(f"Profile {key} must be an object")


def validate_profile_source(source):
    require_keys(source, {"kind", "path", "sha256"}, {"kind", "path", "sha256"}, "profile source")
    if source["kind"] not in {"requirements", "template"}:
        raise FormatError("Profile source kind must be requirements or template")
    path = Path(text_value(source["path"], "profile source path"))
    if not path.is_absolute() or not re.fullmatch(
        r"[0-9a-f]{64}", text_value(source["sha256"], "source SHA-256")
    ):
        raise FormatError("Profile source requires an absolute path and SHA-256")
    if not path.is_file() or sha256(path) != source["sha256"]:
        raise FormatError("Profile source changed or is missing; extract a fresh profile")


def profile_plan(
    input_path,
    destination,
    profile,
    *,
    resolutions=None,
    mapping=None,
    request_text="",
    match_structure=False,
):
    from .api import build_plan
    from .selectors import resolve_targets

    validate_profile(profile)
    source = Path(profile["source"]["path"]).resolve()
    target_path = source_path(input_path)
    target_hash = sha256(target_path)
    output = output_path(target_path, destination)
    if output == source or (output.exists() and output.samefile(source)):
        raise FormatError("Output must not overwrite the profile source")
    resolutions = {} if resolutions is None else resolutions
    mapping = {} if mapping is None else mapping
    require_keys(mapping, {"sections", "rules", "preserve"}, label="mapping")
    if not isinstance(resolutions, dict):
        raise FormatError("Resolutions must map issue IDs to explicit decisions")
    unknown = set(resolutions) - {i["id"] for i in profile["issues"]}
    if unknown:
        raise FormatError(f"Unknown resolution IDs: {sorted(unknown)}")
    rules = copy.deepcopy(profile["rules"])
    pending = []
    structure_contract = None
    if match_structure:
        structure_result = check_template_structure(target_path, profile)
        structure_contract = structure_result["contract"]
        if structure_result["status"] != "passed":
            pending.append(
                {
                    "id": "template-structure",
                    "reason": "Author the draft with exactly the sample's parts before formatting; mapping.preserve cannot waive structural differences",
                    "structure": structure_result,
                }
            )
    decisions = []
    for issue in profile["issues"]:
        resolution = resolutions.get(issue["id"])
        if resolution is None:
            pending.append(issue)
            continue
        require_keys(
            resolution,
            {"decision", "reason", "candidate", "rules"},
            {"decision", "reason"},
            "resolution",
        )
        text_value(resolution["reason"], "resolution reason")
        decision = resolution["decision"]
        if decision == "choose":
            candidate = resolution.get("candidate")
            choices = issue.get("candidates", [])
            if (
                isinstance(candidate, bool)
                or not isinstance(candidate, int)
                or not 0 <= candidate < len(choices)
            ):
                raise FormatError("Resolution candidate must be a valid zero-based index")
            additions = [choices[candidate]]
        elif decision == "replace":
            additions = resolution.get("rules")
            if not isinstance(additions, list) or not additions:
                raise FormatError("Replacement resolution requires nonempty rules")
        elif decision == "ignore":
            additions = []
        else:
            raise FormatError("Resolution decision must be choose, replace or ignore")
        for i, addition in enumerate(additions):
            require_keys(
                addition,
                {"id", "action", "target", "params", "evidence"},
                {"action", "target", "params"},
                "resolution rule",
            )
            rules.append(
                {
                    **addition,
                    "id": f"{issue['id']}-resolved-{i + 1}",
                    "evidence": issue.get("evidence", [])
                    or [{"issue_id": issue["id"], "resolution": resolution}],
                }
            )
        decisions.append({"issue": issue, "resolution": resolution})
    document = Document(target_path)
    index = DocumentIndex(document)
    section_count = profile.get("coverage", {}).get("sections", 0)
    sections = mapping.get("sections", {})
    if not isinstance(sections, dict):
        raise FormatError("mapping.sections must be an object")
    if set(sections) - {str(i) for i in range(1, section_count + 1)}:
        raise FormatError("mapping.sections contains unknown sample section numbers")
    for target_number in sections.values():
        number(target_number, 1, len(document.sections), "target section number", True)
    if section_count > 1 and set(sections) != {str(i) for i in range(1, section_count + 1)}:
        pending.append(
            {
                "id": "section-mapping",
                "reason": "Map every sample section explicitly using mapping.sections",
                "sample_sections": section_count,
                "target_sections": len(document.sections),
            }
        )
    operations, matched = [], []
    preserved_rules = []
    rule_mapping = mapping.get("rules", {})
    if not isinstance(rule_mapping, dict) or set(rule_mapping) - {r["id"] for r in rules}:
        raise FormatError("mapping.rules contains unknown rule IDs")
    for rule in rules:
        cap = REGISTRY.get(rule.get("action"))
        if cap is None:
            raise FormatError("Unregistered profile action", "unsupported")
        params = cap.validator(rule["params"])
        target = copy.deepcopy(rule_mapping.get(rule["id"], rule["target"]))
        if not isinstance(target, dict):
            raise FormatError("Rule target mapping must be an object")
        if isinstance(target, dict) and set(target) == {"preserve"}:
            text_value(target["preserve"], "preserved rule reason")
            preserved_rules.append({"rule": rule, "reason": target["preserve"]})
            continue
        if (
            rule["id"] not in rule_mapping
            and target.get("type") == "section"
            and "number" in target
            and str(target["number"]) in sections
        ):
            target["number"] = sections[str(target["number"])]
        if rule["id"] not in rule_mapping and any(
            e.get("confidence", 1) < 1 for e in rule.get("evidence", [])
        ):
            pending.append(
                {
                    "id": rule["id"],
                    "reason": "Low-confidence evidence requires an explicit mapping.rules decision",
                    "evidence": rule["evidence"],
                }
            )
            continue
        try:
            refs = resolve_targets(index, target, cap.targets)
        except FormatError as exc:
            pending.append(
                {
                    "id": rule["id"],
                    "reason": str(exc),
                    "target": target,
                    "evidence": rule.get("evidence", []),
                }
            )
            continue
        operations.append(
            {"id": rule["id"], "action": cap.action, "target": target, "params": params}
        )
        matched.append(
            {
                "rule_id": rule["id"],
                "target_ids": [r["id"] for r in refs],
                "evidence": rule.get("evidence", []),
            }
        )
    preserve = mapping.get("preserve", {})
    if not isinstance(preserve, dict):
        raise FormatError("mapping.preserve must map target paragraph IDs to reasons")
    for oid, reason in preserve.items():
        if index.records.get(oid, {}).get("kind") != "paragraph":
            raise FormatError("Unknown preserved target paragraph ID")
        text_value(reason, "preserved paragraph reason")
    if profile["source"]["kind"] == "template":
        covered = {oid for m in matched for oid in m["target_ids"]}
        uncovered = [
            r["id"]
            for r in index.records.values()
            if r["kind"] == "paragraph"
            and r["part"] == "word/document.xml"
            and r.get("text", "").strip()
            and r["id"] not in covered
            and r["id"] not in preserve
        ]
        if uncovered:
            pending.append(
                {
                    "id": "target-coverage",
                    "reason": "Target paragraphs have no sample role; map rules explicitly or document preservation in mapping.preserve",
                    "target_ids": uncovered,
                }
            )
    context = {
        "source": profile["source"],
        "profile_sha256": json_hash(profile),
        "coverage": profile.get("coverage", {}),
        "mapping": matched,
        "decisions": decisions,
        "preserved_rules": preserved_rules,
        "preserved_targets": preserve,
    }
    if structure_contract is not None:
        context["structure_contract"] = structure_contract
    if pending:
        return {
            "schema_version": "1.0",
            "kind": "word-format-profile-review",
            "status": "needs_clarification",
            "issues": pending,
            "profile_context": context,
            "note": "No executable plan was generated; resolve issues or supply explicit target mappings",
        }
    if not operations:
        raise FormatError("Profile has no executable formatting rules", "needs_clarification")
    if sha256(target_path) != target_hash:
        raise FormatError("Target changed during profile matching; inspect and plan again")
    request = {
        "origin": "direct",
        "operations": operations,
        "request_text": request_text,
        "profile_context": context,
    }
    return build_plan(input_path, destination, request)


def template_structure_contract(profile):
    """Read authority again; stale or edited snapshots cannot redefine its parts."""
    validate_profile_source(profile["source"])
    if profile["source"]["kind"] != "template":
        raise FormatError("Exact thesis structure requires a template profile")
    contract = extract_template_structure(_template_document(Path(profile["source"]["path"])))
    recorded = profile.get("snapshot", {}).get("structure_contract")
    if recorded is not None and recorded != contract:
        raise FormatError("Template structure snapshot changed; extract a fresh profile")
    return contract


def check_template_structure(input_path, profile):
    """Read-only preflight/final check for drafts authored from a sample."""
    validate_profile(profile)
    target = source_path(input_path)
    target_hash = sha256(target)
    contract = template_structure_contract(profile)
    result = compare_template_structure(Document(target), contract)
    validate_profile_source(profile["source"])
    if sha256(target) != target_hash:
        raise FormatError("Draft changed during structure checking; inspect and check again")
    return {
        "kind": "word-template-structure-check",
        "source": copy.deepcopy(profile["source"]),
        "target": {"path": str(target), "sha256": target_hash},
        "contract": contract,
        **result,
    }
