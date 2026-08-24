#!/usr/bin/env python3
"""Build and apply a deterministic AI citation-placement exchange format."""

from __future__ import annotations

import copy
import hashlib
import re
from pathlib import Path
from typing import Any

from docx.document import Document as DocumentType
from docx.text.paragraph import Paragraph
from docx.text.run import Run


REFERENCE_HEADING_RE = re.compile(
    r"^\s*(?:references|bibliography|works cited|参考文献|参考资料)\s*$", re.I
)
TERMINAL_PUNCTUATION = set(".。!！?？;；:：,，、")


def normalized_text(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def paragraph_sha256(text: str) -> str:
    return hashlib.sha256(normalized_text(text).encode("utf-8")).hexdigest()


def _is_heading(paragraph: Paragraph) -> bool:
    style = paragraph.style.name.lower() if paragraph.style else ""
    return style.startswith("heading")


def extract_citation_data(document: DocumentType) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Return body paragraph anchors and references using stable top-level indexes."""
    paragraphs: list[dict[str, Any]] = []
    references: list[dict[str, Any]] = []
    in_references = False
    reference_index = 0
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if REFERENCE_HEADING_RE.match(text):
            in_references = True
            continue
        if in_references:
            if text and not _is_heading(paragraph):
                reference_index += 1
                label_match = re.match(r"^\s*(\[\s*\d+\s*\]|\d+[.)])", text)
                label = label_match.group(1) if label_match else f"[{reference_index}]"
                references.append({
                    "id": f"ref-{reference_index}",
                    "label": label,
                    "paragraph_index": index,
                    "text": text,
                    "paragraph_sha256": paragraph_sha256(text),
                })
            continue
        if text and not _is_heading(paragraph):
            paragraphs.append({
                "paragraph_index": index,
                "text": text,
                "paragraph_sha256": paragraph_sha256(text),
            })
    return paragraphs, references


def build_citation_task(document: DocumentType, *, filename: str = "document.docx", source_sha256: str | None = None) -> dict[str, Any]:
    paragraphs, references = extract_citation_data(document)
    return {
        "schema_version": "1.0",
        "kind": "word-citation-placement-task",
        "source": {
            "filename": Path(filename).name,
            "sha256": source_sha256,
            "paragraph_index_scope": "top-level document paragraphs",
        },
        "instructions": {
            "goal": "Identify claims in body paragraphs that require a cited reference.",
            "insert_options": ["before-terminal-punctuation", "after-terminal-punctuation", "end", "start"],
            "spacing_options": ["none", "space-before", "space-after", "spaces-around"],
            "return_only": "JSON with citations.placements; do not rewrite paragraph text or references.",
            "fingerprint_rule": "Copy paragraph_index and paragraph_sha256 exactly from candidates.",
        },
        "paragraphs": paragraphs,
        "references": references,
        "citations": {"placements": []},
    }


def _citation_offset(text: str, placement: dict[str, Any]) -> int:
    position = str(placement.get("insert") or "before-terminal-punctuation")
    if position == "start":
        return 0
    if position == "char-offset":
        offset = int(placement.get("char_offset", -1))
        if offset < 0 or offset > len(text):
            raise ValueError("citation char_offset is outside the paragraph")
        return offset
    end = len(text.rstrip())
    if position == "end":
        return len(text)
    if position in {"before-terminal-punctuation", "after-terminal-punctuation"}:
        index = end - 1
        while index >= 0 and text[index] in TERMINAL_PUNCTUATION:
            index -= 1
        punctuation_start = index + 1
        return punctuation_start if position.startswith("before") else end
    raise ValueError(f"unsupported citation insert position: {position}")


def _insert_run_at_offset(paragraph: Paragraph, offset: int, text: str) -> None:
    runs = list(paragraph.runs)
    if not runs:
        paragraph.add_run(text)
        return
    cursor = 0
    for run in runs:
        run_text = run.text or ""
        next_cursor = cursor + len(run_text)
        if offset <= next_cursor:
            relative = max(0, offset - cursor)
            original_element = copy.deepcopy(run._r)
            before, after = run_text[:relative], run_text[relative:]
            run.text = before
            citation = paragraph.add_run()
            if original_element.rPr is not None:
                citation._r.insert(0, copy.deepcopy(original_element.rPr))
            citation.text = text
            run._r.addnext(citation._r)
            if after:
                suffix = Run(original_element, paragraph)
                suffix.text = after
                citation._r.addnext(suffix._r)
            return
        cursor = next_cursor
    paragraph.add_run(text)


def apply_citations(document: DocumentType, citations: dict[str, Any] | None) -> list[str]:
    if not citations:
        return []
    placements = citations.get("placements", [])
    if not isinstance(placements, list):
        raise ValueError("citations.placements must be an array")
    _, references = extract_citation_data(document)
    reference_ids = {item["id"] for item in references}
    indexed = list(document.paragraphs)
    validated: list[tuple[Paragraph, dict[str, Any], str]] = []
    seen_paragraphs: set[int] = set()
    for placement in placements:
        if not isinstance(placement, dict):
            raise ValueError("each citation placement must be an object")
        index = int(placement.get("paragraph_index", -1))
        if index < 0 or index >= len(indexed):
            raise ValueError(f"citation paragraph_index {index} is outside the document")
        if index in seen_paragraphs:
            raise ValueError(f"multiple citation placements for paragraph {index} are not supported")
        seen_paragraphs.add(index)
        paragraph = indexed[index]
        expected = str(placement.get("paragraph_sha256") or "")
        actual = paragraph_sha256(paragraph.text)
        if not expected or expected != actual:
            raise ValueError(f"citation paragraph fingerprint mismatch at index {index}")
        ids = placement.get("reference_ids", [])
        if not isinstance(ids, list) or not ids or any(str(item) not in reference_ids for item in ids):
            raise ValueError(f"citation references are invalid at paragraph {index}")
        citation_text = str(placement.get("citation_text") or "").strip()
        if not citation_text or len(citation_text) > 200:
            raise ValueError(f"citation_text is missing or too long at paragraph {index}")
        if citation_text in paragraph.text:
            raise ValueError(f"citation text already exists at paragraph {index}")
        spacing = str(placement.get("spacing") or "none")
        if spacing not in {"none", "space-before", "space-after", "spaces-around"}:
            raise ValueError(f"citation spacing is invalid at paragraph {index}")
        insertion_text = (
            (" " if spacing in {"space-before", "spaces-around"} else "")
            + citation_text
            + (" " if spacing in {"space-after", "spaces-around"} else "")
        )
        validated.append((paragraph, placement, insertion_text))

    for paragraph, placement, citation_text in validated:
        offset = _citation_offset(paragraph.text, placement)
        _insert_run_at_offset(paragraph, offset, citation_text)
    return [f"Inserted {len(validated)} AI-approved citation placement(s) with paragraph fingerprints"]
