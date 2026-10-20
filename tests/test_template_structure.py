"""Sample-based drafting must neither invent nor omit thesis parts."""

import copy
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from verify_output import verify_structure
from word_format import apply_plan, build_plan, verify_plan
from word_format.contracts import FormatError, sha256
from word_format.profiles import check_template_structure, profile_plan, template_profile
from word_format.template_structure import extract_template_structure


class TemplateStructureTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.sample = self.root / "样例.docx"
        self.draft = self.root / "草稿.docx"
        self.output = self.root / "输出.docx"

    def write(self, path, parts, *, sample=False):
        doc = Document()
        doc.styles["Normal"].font.size = Pt(12)
        doc.styles["Normal"].font.name = "Times New Roman"
        for text in parts:
            if text == "cover":
                doc.add_paragraph("样例题目" if sample else "新论文题目", style="Title")
                doc.add_paragraph("样例姓名" if sample else "新作者")
                doc.add_page_break()
            elif text.startswith("chapter"):
                doc.add_heading("第一章 样例研究" if sample else "第一章 新研究", 1)
                doc.add_heading("1.1 样例方法" if sample else "1.1 新方法", 2)
                doc.add_paragraph("样例正文" if sample else "新的研究正文")
            else:
                doc.add_paragraph(text)
                if text not in {"目录", "关键词：测试", "Key words: test"}:
                    doc.add_paragraph("样例内容" if sample else "新的模块内容")
        doc.save(path)
        return doc

    def profile(self, parts):
        self.write(self.sample, parts, sample=True)
        profile = template_profile(self.sample)
        # These synthetic fixtures leave theme fonts unspecified. The tests only
        # preserve those undefined formatting properties, never structure gaps.
        resolutions = {
            issue["id"]: {"decision": "ignore", "reason": "测试样例未定义的格式保持原值"}
            for issue in profile["issues"]
        }
        return profile, resolutions

    def test_same_parts_pass_with_new_text_and_no_english_or_declaration(self):
        parts = ["cover", "摘要", "关键词：测试", "chapter", "参考文献"]
        profile, decisions = self.profile(parts)
        self.write(self.draft, parts)
        self.assertEqual(check_template_structure(self.draft, profile)["status"], "passed")
        plan = profile_plan(self.draft, self.output, profile,
                            resolutions=decisions, match_structure=True)
        self.assertEqual(plan["request"]["verification_spec"]["document_structure"]["match"], "exact")
        report = apply_plan(plan)
        self.assertEqual(verify_plan(plan, report)["status"], "passed")
        texts = "\n".join(p.text for p in Document(self.output).paragraphs)
        self.assertIn("新论文题目", texts)
        self.assertNotIn("样例", texts)
        self.assertNotIn("Abstract", texts)
        self.assertNotIn("承诺书", texts)

    def test_missing_cover_or_extra_optional_parts_block_plan_even_when_preserved(self):
        parts = ["cover", "摘要", "关键词：测试", "chapter", "参考文献"]
        profile, decisions = self.profile(parts)
        variants = [parts[1:]]
        for extra in ("英文摘要", "Abstract", "诚信承诺书", "致谢", "附录 A", "目录", "符号说明"):
            variants.append(parts[:1] + [extra] + parts[1:])
        for actual in variants:
            with self.subTest(parts=actual):
                doc = self.write(self.draft, actual)
                mapping = {"preserve": {f"p{i}": "保留目标段落格式" for i in range(len(doc.paragraphs))}}
                result = profile_plan(self.draft, self.output, profile, resolutions=decisions,
                                      mapping=mapping, match_structure=True)
                self.assertEqual(result["status"], "needs_clarification")
                self.assertTrue(any(i["id"] == "template-structure" for i in result["issues"]))
                self.assertFalse(self.output.exists())

    def test_missing_present_parts_and_different_languages_fail(self):
        parts = ["cover", "诚信承诺书", "摘要", "关键词：测试", "Abstract",
                 "Key words: test", "chapter", "参考文献"]
        profile, _ = self.profile(parts)
        for missing in ("cover", "诚信承诺书", "摘要", "Abstract", "Key words: test", "chapter"):
            with self.subTest(missing=missing):
                self.write(self.draft, [p for p in parts if p != missing])
                self.assertEqual(check_template_structure(self.draft, profile)["status"], "failed")
        # Equal numbers of abstract/keyword modules cannot conceal a language swap.
        self.write(self.draft, ["摘要" if p == "Abstract" else p for p in parts])
        result = check_template_structure(self.draft, profile)
        self.assertEqual(result["status"], "failed")
        self.assertTrue(result["differences"])

    def test_duplicate_parts_chapters_and_changed_hierarchy_fail(self):
        parts = ["摘要", "关键词：测试", "chapter", "参考文献"]
        profile, _ = self.profile(parts)
        for actual in (parts[:1] + ["摘要"] + parts[1:],
                       parts[:3] + ["chapter2"] + parts[3:]):
            self.write(self.draft, actual)
            self.assertEqual(check_template_structure(self.draft, profile)["status"], "failed")
        doc = self.write(self.draft, parts)
        doc.add_heading("1.2 不应额外出现的节", 2)
        # Put the extra subsection inside the chapter, before references.
        doc.element.body.insert(len(doc.element.body) - 4, doc.paragraphs[-1]._p)
        doc.save(self.draft)
        self.assertEqual(check_template_structure(self.draft, profile)["status"], "failed")

    def test_actual_sample_order_wins_over_generic_conventional_order(self):
        parts = ["摘要", "chapter", "参考文献", "诚信承诺书"]
        profile, _ = self.profile(parts)
        self.write(self.draft, parts)
        self.assertEqual(check_template_structure(self.draft, profile)["status"], "passed")
        self.write(self.draft, ["诚信承诺书", *parts[:-1]])
        self.assertEqual(check_template_structure(self.draft, profile)["status"], "failed")

    def test_declaration_forms_and_appendix_titles_have_distinct_occurrences(self):
        parts = ["cover", "本科毕业论文（设计）诚信承诺书", "学位论文原创性声明",
                 "学位论文使用授权书", "摘要", "chapter", "附录A 调查问卷", "附录 B 数据说明"]
        profile, _ = self.profile(parts)
        modules = profile["snapshot"]["structure_contract"]["ordered_modules"]
        self.assertEqual([m["variant"] for m in modules if m["kind"] == "declaration"],
                         ["integrity", "originality", "authorization"])
        self.assertEqual([m["occurrence"] for m in modules if m["kind"] == "appendix"], [1, 2])
        self.write(self.draft, ["诚 信 承 诺 书" if p == "本科毕业论文（设计）诚信承诺书" else p for p in parts])
        self.assertEqual(check_template_structure(self.draft, profile)["status"], "passed")
        self.write(self.draft, ["诚信承诺书" if p == "学位论文原创性声明" else p for p in parts])
        self.assertEqual(check_template_structure(self.draft, profile)["status"], "failed")

    def test_table_cover_and_cached_toc_entries_do_not_invent_parts(self):
        doc = Document()
        doc.add_table(rows=1, cols=2).cell(0, 0).text = "论文题目：表格封面"
        doc.add_paragraph("摘要")
        doc.add_paragraph("目录")
        doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
        doc.add_paragraph("参考文献", style="TOC 1")
        doc.add_heading("第一章 绪论", 1)
        structure = extract_template_structure(doc)
        self.assertEqual([m["kind"] for m in structure["ordered_modules"]],
                         ["cover", "abstract", "toc", "chapters"])

    def test_image_cover_detected_and_plain_body_is_not_a_cover(self):
        from PIL import Image
        image = self.root / "封面.png"
        Image.new("RGB", (16, 16), "white").save(image)
        doc = Document()
        doc.add_picture(str(image))
        doc.add_heading("第一章 绪论", 1)
        self.assertEqual(extract_template_structure(doc)["ordered_modules"][0]["kind"], "cover")
        plain = Document()
        plain.add_paragraph("普通正文")
        self.assertEqual(extract_template_structure(plain)["ordered_modules"][0]["kind"], "body")

    def test_inline_article_title_cannot_become_a_separate_cover(self):
        doc = Document()
        doc.add_paragraph("文章标题", style="Title")
        doc.add_paragraph("摘要")
        doc.add_paragraph("摘要内容")
        doc.add_heading("第一章 研究问题", 1)
        doc.save(self.sample)
        profile = template_profile(self.sample)
        module = profile["snapshot"]["structure_contract"]["ordered_modules"][0]
        self.assertEqual(module["kind"], "title_block")
        self.assertFalse(module["separate_page"])
        doc.paragraphs[0].text = "新的文章题目"
        doc.save(self.draft)
        self.assertEqual(check_template_structure(self.draft, profile)["status"], "passed")
        doc.paragraphs[1].paragraph_format.page_break_before = True
        doc.save(self.draft)
        self.assertEqual(check_template_structure(self.draft, profile)["status"], "failed")

    def test_cover_boundary_detected_from_a_section_break(self):
        from docx.enum.section import WD_SECTION_START
        doc = Document()
        doc.add_paragraph("论文新标题", style="Title")
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        doc.add_paragraph("摘要")
        modules = extract_template_structure(doc)["ordered_modules"]
        self.assertEqual(modules[0]["kind"], "cover")
        self.assertTrue(modules[0]["separate_page"])

    def test_format_only_preserves_existing_parts_without_draft_mode(self):
        profile, decisions = self.profile(["chapter"])
        self.write(self.draft, ["chapter", "诚信承诺书"])
        doc = Document(self.draft)
        mapping = {"preserve": {f"p{i}": "用户仅排版，保留正文" for i in range(len(doc.paragraphs))}}
        plan = profile_plan(self.draft, self.output, profile, resolutions=decisions, mapping=mapping)
        self.assertEqual(verify_plan(plan, apply_plan(plan))["status"], "passed")
        self.assertIn("诚信承诺书", [p.text for p in Document(self.output).paragraphs])

    def test_added_toc_during_apply_cannot_publish(self):
        parts = ["摘要", "chapter"]
        profile, decisions = self.profile(parts)
        self.write(self.draft, parts)
        plan = profile_plan(self.draft, self.output, profile,
                            resolutions=decisions, match_structure=True)
        request = copy.deepcopy(plan["request"])
        request["operations"].append({"action": "toc.configure", "target": {"type": "document"},
                                      "params": {"enabled": True}})
        plan = build_plan(self.draft, self.output, request)
        with self.assertRaisesRegex(FormatError, "structure|Thesis parts"):
            apply_plan(plan)
        self.assertFalse(self.output.exists())

    def test_changed_structure_snapshot_or_context_is_rejected(self):
        parts = ["cover", "摘要", "chapter"]
        profile, decisions = self.profile(parts)
        self.write(self.draft, parts)
        bad = copy.deepcopy(profile)
        bad["snapshot"]["structure_contract"]["ordered_modules"].pop(0)
        with self.assertRaises(FormatError):
            check_template_structure(self.draft, bad)
        plan = profile_plan(self.draft, self.output, profile,
                            resolutions=decisions, match_structure=True)
        request = copy.deepcopy(plan["request"])
        request["verification_spec"]["document_structure"]["ordered_modules"].pop(0)
        with self.assertRaises(FormatError):
            build_plan(self.draft, self.output, request)

    def test_legacy_structure_does_not_allow_unspecified_standard_parts(self):
        self.write(self.draft, ["摘要", "chapter", "诚信承诺书"])
        result = verify_structure(self.draft, {"document_structure": {
            "ordered_modules": [{"kind": "abstract"}, {"kind": "chapters"}]}}, check_format=False)
        self.assertEqual(result["status"], "failed")

    def test_cli_checks_structure_and_protects_input_and_sample(self):
        parts = ["cover", "摘要", "chapter"]
        profile, _ = self.profile(parts)
        self.write(self.draft, parts)
        profile_path = self.root / "profile.json"
        profile_path.write_text(json.dumps(profile, ensure_ascii=False), encoding="utf-8")
        report_path = self.root / "structure.json"
        command = [sys.executable, str(ROOT / "scripts/format_cli.py"), "check-structure",
                   str(self.draft), "--profile", str(profile_path), "--output"]
        result = subprocess.run([*command, str(report_path)], capture_output=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr.decode("utf-8"))
        self.assertEqual(json.loads(report_path.read_text(encoding="utf-8"))["status"], "passed")
        for protected in (self.sample, self.draft, profile_path):
            before = sha256(protected)
            result = subprocess.run([*command, str(protected)], capture_output=True, timeout=30)
            self.assertEqual(result.returncode, 2)
            self.assertEqual(sha256(protected), before)
        self.write(self.draft, parts + ["Abstract"])
        result = subprocess.run([*command, str(report_path)], capture_output=True, timeout=30)
        self.assertEqual(result.returncode, 2)
        self.assertEqual(json.loads(report_path.read_text(encoding="utf-8"))["status"], "failed")
        decisions_path = self.root / "resolutions.json"
        decisions_path.write_text(json.dumps({
            issue["id"]: {"decision": "ignore", "reason": "测试中保留未定义字体"}
            for issue in profile["issues"]
        }), encoding="utf-8")
        result = subprocess.run([
            sys.executable, str(ROOT / "scripts/format_cli.py"), "plan", str(self.draft),
            str(self.output), "--profile", str(profile_path), "--match-structure",
            "--resolutions", str(decisions_path), "--output", str(report_path),
        ], capture_output=True, timeout=30)
        self.assertEqual(result.returncode, 2)
        review = json.loads(report_path.read_text(encoding="utf-8"))
        self.assertTrue(any(i["id"] == "template-structure" for i in review["issues"]))
        self.assertFalse(self.output.exists())


if __name__ == "__main__":
    unittest.main()
