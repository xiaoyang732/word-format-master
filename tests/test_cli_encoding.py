"""CLI reports stay UTF-8 even when Windows pipes default to cp1252."""

import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from apply_spec import SUPPORTED_APPLICATION_PATHS
from dashboard_session import build_ai_handoff, prepare_dashboard_session
from word_format import build_plan
from word_format.contracts import sha256


class CliEncodingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.source = self.folder / "论文表格.docx"
        self.output = self.folder / "排版结果.docx"
        doc = Document()
        doc.add_paragraph("正文中文表格")
        doc.save(self.source)
        self.source_hash = sha256(self.source)
        # Reproduce GitHub's non-UTF-8 redirected streams on any host locale.
        self.environment = {**os.environ, "PYTHONUTF8": "0", "PYTHONIOENCODING": "cp1252"}

    def run_cli(self, script, *args):
        return subprocess.run(
            [sys.executable, "-B", str(ROOT / "scripts" / script), *map(str, args)],
            capture_output=True, text=True, encoding="utf-8", errors="strict",
            env=self.environment, timeout=60,
        )

    def write_json(self, name, data):
        path = self.folder / name
        path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        return path

    def assert_json_success(self, result):
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertNotIn("\\u", result.stdout)
        return json.loads(result.stdout)

    def test_confirmed_handoff_report_stdout_and_file_preserve_chinese(self):
        for to_file in (False, True):
            with self.subTest(to_file=to_file):
                output = self.output if not to_file else self.folder / "报告文件模式.docx"
                session = prepare_dashboard_session(self.source, output)
                handoff = build_ai_handoff(
                    session, {"spec": {"body": {"font_size_pt": 12}}, "verification": {"visual_enabled": False}},
                    skill_dir=ROOT, registry=SUPPORTED_APPLICATION_PATHS,
                )
                path = self.write_json("handoff.json", handoff)
                args = [self.source, output, "--handoff", path]
                if to_file:
                    report_path = self.folder / "报告.json"
                    result = self.run_cli("apply_spec.py", *args, "--report", report_path)
                    self.assertEqual(result.returncode, 0, result.stderr)
                    report = json.loads(report_path.read_text(encoding="utf-8"))
                else:
                    result = self.run_cli("apply_spec.py", *args)
                    report = self.assert_json_success(result)
                    self.assertIn("排版结果", result.stdout)
                self.assertEqual(report["status"], "passed")
                self.assertEqual(report["output"], str(output))
                self.assertEqual(Document(output).paragraphs[0].text, "正文中文表格")
                self.assertEqual(sha256(self.source), self.source_hash)
                output.unlink()

    def test_direct_cli_inspect_apply_verify_utf8(self):
        inspected = self.assert_json_success(self.run_cli("format_cli.py", "inspect", self.source))
        self.assertIn("正文中文表格", json.dumps(inspected, ensure_ascii=False))
        request = {"operations": [{"action": "font.size.set", "target": {"type": "paragraph", "id": "p0"}, "params": {"value": 12}}]}
        plan = build_plan(self.source, self.output, request)
        plan_path = self.write_json("plan.json", plan)
        report = self.assert_json_success(self.run_cli("format_cli.py", "apply", "--plan", plan_path))
        report_path = self.write_json("report.json", report)
        verified = self.assert_json_success(self.run_cli("format_cli.py", "verify", "--plan", plan_path, "--report", report_path))
        self.assertEqual(verified["status"], "passed")
        self.assertEqual(sha256(self.source), self.source_hash)

    def test_analysis_requirements_verification_json_utf8(self):
        analysis = self.assert_json_success(self.run_cli("analyze_docx.py", self.source))
        self.assertEqual(analysis["source"]["filename"], self.source.name)
        parsed = self.assert_json_success(self.run_cli("parse_requirements.py", "--text", "正文宋体，小四。"))
        self.assertEqual(parsed["spec_patch"]["body"]["font_east_asia"], "宋体")
        spec = self.write_json("spec.json", {})
        result = self.run_cli("verify_output.py", self.source, spec)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout)["status"], "passed")

    def test_chinese_errors_on_stderr_utf8(self):
        missing = self.folder / "不存在的文档.docx"
        invalid_request = self.write_json("invalid-request.json", {"operations": [{
            "action": "不存在的格式操作", "target": {"type": "paragraph", "id": "p0"}, "params": {},
        }]})
        for script, args in (
            ("format_cli.py", ["plan", self.source, self.output, "--request", invalid_request]),
            ("analyze_docx.py", [missing]),
            ("apply_spec.py", [self.source, self.output, "--handoff", self.folder / "不存在的交接.json"]),
            ("render_docx.py", [self.source, "--output-dir", self.folder / "pages", "--dpi", "1"]),
        ):
            with self.subTest(script=script):
                result = self.run_cli(script, *args)
                self.assertNotEqual(result.returncode, 0)
                self.assertNotIn("UnicodeEncodeError", result.stderr)
                if script != "analyze_docx.py":
                    self.assertTrue(any("\u4e00" <= char <= "\u9fff" for char in result.stderr))


if __name__ == "__main__":
    unittest.main()
