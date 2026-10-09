"""Header requirements, isolation, explicit boundaries and failed-writer tests."""

import copy
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from word_format import build_plan, apply_plan, verify_plan, inspect_document
from word_format.contracts import FormatError, sha256
from word_format.registry import REGISTRY
from word_format.headers import effective_root
from analyze_docx import analyze_docx
from parse_requirements import parse_requirements
from word_format.spec import legacy_paths


class HeaderTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / "source.docx"
        self.output = Path(self.temp.name) / "output.docx"
        doc = Document()
        doc.add_paragraph("Abstract", "Heading 1")
        doc.add_paragraph("Front matter")
        doc.sections[0].header.paragraphs[0].text = "original"
        doc.sections[0].header.paragraphs[0].runs[0].bold = True
        doc.sections[0].footer.paragraphs[0].text = "footer preserved"
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        doc.add_paragraph("Chapter One", "Heading 1")
        doc.add_paragraph("Body")
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        doc.add_paragraph("References", "Heading 1")
        doc.add_paragraph("Reference entry")
        doc.save(self.source)

    def op(self, action, params, number=2):
        return {"action": action, "target": {"type": "section", "number": number}, "params": params}

    def execute(self, operations=None, spec=None):
        source_hash = sha256(self.source)
        request = {"spec": spec} if spec is not None else {"operations": operations}
        plan = build_plan(self.source, self.output, request)
        report = apply_plan(plan)
        self.assertEqual(verify_plan(plan, report)["status"], "passed")
        self.assertEqual(sha256(self.source), source_hash)
        return Document(self.output), report

    def text(self, header):
        root = effective_root(header)
        return "".join(n.text or "" for n in root.iter(qn("w:t"))) if root is not None else ""

    def test_section_content_isolates_following_inherited_header(self):
        doc, _ = self.execute([self.op("header.content.set", {"mode": "text", "text": "Chapter One"})])
        self.assertEqual([self.text(s.header) for s in doc.sections], ["original", "Chapter One", "original"])
        self.assertEqual([self.text(s.footer) for s in doc.sections], ["footer preserved"] * 3)
        self.assertTrue(doc.sections[0].header.paragraphs[0].runs[0].bold)
        self.assertTrue(doc.sections[1].header.paragraphs[0].runs[0].bold)

    def test_chapter_fields_and_odd_even_variants(self):
        doc, _ = self.execute([
            {"action": "header.odd_even_different.set", "target": {"type": "document"}, "params": {"value": True}},
            self.op("header.content.set", {"mode": "chapter_title", "style_name": "Heading 1", "prefix": "Thesis: "}),
            self.op("header.content.set", {"mode": "text", "text": "Thesis title", "variant": "even"}),
            self.op("header.alignment.set", {"value": "right"}),
            self.op("header.alignment.set", {"value": "left", "variant": "even"}),
            self.op("header.font.size.set", {"value": 10.5}),
            self.op("header.font.east_asia.set", {"value": "SimSun"}),
            self.op("header.font.latin.set", {"value": "Times New Roman"}),
        ])
        self.assertTrue(doc.settings.odd_and_even_pages_header_footer)
        field = doc.sections[1].header._element.find(".//" + qn("w:fldSimple"))
        self.assertEqual(field.get(qn("w:instr")), 'STYLEREF "Heading 1"')
        self.assertEqual(self.text(doc.sections[1].even_page_header), "Thesis title")

    def test_first_page_blank_only_selected_section(self):
        doc, _ = self.execute([
            self.op("header.first_page_different.set", {"value": True}),
            self.op("header.content.set", {"mode": "text", "text": "", "variant": "first"}),
        ])
        self.assertEqual([s.different_first_page_header_footer for s in doc.sections], [False, True, False])
        self.assertEqual(self.text(doc.sections[1].first_page_header), "")

    def test_local_font_keeps_text_bold_alignment_and_following_header(self):
        doc, _ = self.execute([self.op("header.font.size.set", {"value": 9})])
        self.assertEqual([self.text(s.header) for s in doc.sections], ["original"] * 3)
        self.assertEqual(doc.sections[1].header.paragraphs[0].runs[0].font.size.pt, 9)
        self.assertTrue(doc.sections[1].header.paragraphs[0].runs[0].bold)
        self.assertIsNone(doc.sections[2].header.paragraphs[0].runs[0].font.size)

    def test_wrong_writer_and_footer_pollution_are_rejected(self):
        plan = build_plan(self.source, self.output, {"operations": [self.op("header.font.size.set", {"value": 9})]})
        cap = REGISTRY["header.font.size.set"]
        def bad(index, ref, params):
            cap.handler(index, ref, params)
            index.document.sections[0].footer.paragraphs[0].text = "unexpected"
        with patch.dict(REGISTRY, {cap.action: type(cap)(**{**cap.__dict__, "handler": bad})}):
            with self.assertRaises(FormatError): apply_plan(plan)
        self.assertFalse(self.output.exists())

    def test_existing_header_relationships_survive_unlink(self):
        doc = Document(self.source)
        p = doc.sections[0].header.paragraphs[0]
        rel = p.part.relate_to("https://example.com", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", True)
        link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), rel)
        run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = "linked"; run.append(text); link.append(run); p._p.append(link)
        doc.save(self.source)
        result, _ = self.execute([self.op("header.font.size.set", {"value": 9})])
        for s in result.sections:
            node = effective_root(s.header).find(".//" + qn("w:hyperlink"))
            self.assertIsNotNone(node)
            self.assertEqual(s.header.part.rels[node.get(qn("r:id"))].target_ref, "https://example.com")

    def test_footer_relationship_pollution_is_rejected(self):
        doc=Document(self.source)
        p=doc.sections[0].footer.paragraphs[0]
        rid=p.part.relate_to("https://example.com", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", True)
        link=OxmlElement("w:hyperlink"); link.set(qn("r:id"),rid)
        run=OxmlElement("w:r"); text=OxmlElement("w:t"); text.text="footer link"
        run.append(text); link.append(run); p._p.append(link)
        doc.save(self.source)
        cap=REGISTRY["header.font.size.set"]
        plan=build_plan(self.source,self.output,{"operations":[self.op(cap.action,{"value":9})]})
        def bad(index,ref,params):
            cap.handler(index,ref,params)
            index.document.sections[0].footer.part.rels[rid]._target="https://changed.example"
        with patch.dict(REGISTRY,{cap.action:type(cap)(**{**cap.__dict__,"handler":bad})}):
            with self.assertRaises(FormatError): apply_plan(plan)
        self.assertFalse(self.output.exists())

    def test_spec_expands_web_fields_to_same_handlers(self):
        spec = {"headers_footers": {"different_odd_even": True,
            "header": {"enabled": True, "mode": "chapter_title", "style_name": "Heading 1", "font_size_pt": 9, "alignment": "right"},
            "even_header": {"enabled": True, "text": "Thesis title", "font_size_pt": 9, "alignment": "left"},
            "sections": [{"number": 1, "header": {"enabled": True, "text": "Abstract"}}]}}
        doc, report = self.execute(spec=spec)
        self.assertEqual(self.text(doc.sections[0].header), "Abstract")
        self.assertTrue(any(o["action"] == "header.font.size.set" for o in report["operations"]))

    def test_explicit_next_page_break_preserves_body(self):
        before = Document(self.source)
        doc, _ = self.execute([{"action": "section.break.next_page.insert", "target": {"type": "paragraph", "id": "p1"}, "params": {}}])
        self.assertEqual(len(doc.sections), len(before.sections) + 1)
        self.assertEqual([p.text for p in doc.paragraphs], [p.text for p in before.paragraphs])
        data = inspect_document(self.output)
        self.assertTrue(any(r["kind"] == "section" and "headers" in r for r in data["objects"]))

    def test_bad_variant_missing_style_and_duplicate_boundary(self):
        for operation in [self.op("header.font.size.set", {"value": 9, "variant": "unknown"}),
                          self.op("header.content.set", {"mode": "chapter_title", "style_name": 'Heading "1"'})]:
            with self.assertRaises(FormatError): build_plan(self.source, self.output, {"operations": [operation]})
        plan = build_plan(self.source, self.output, {"operations": [self.op("header.content.set", {"mode": "chapter_title", "style_name": "Missing"})]})
        with self.assertRaises(FormatError): apply_plan(plan)
        self.assertFalse(self.output.exists())
        # Empty paragraphs already carrying a section boundary are not reusable.
        doc=Document(self.source)
        boundary_index=next(i for i,p in enumerate(doc.paragraphs) if p._p.find("./"+qn("w:pPr")+"/"+qn("w:sectPr")) is not None)
        plan=build_plan(self.source,self.output,{"operations":[{"action":"section.break.next_page.insert","target":{"type":"paragraph","id":f"p{boundary_index}"},"params":{}}]})
        with self.assertRaises(FormatError): apply_plan(plan)
        self.assertFalse(self.output.exists())

    def test_all_header_property_handlers_readback_and_isolation(self):
        parameters={"font_east_asia":{"value":"SimSun"},"font_latin":{"value":"Arial"},"font_size_pt":{"value":10.5},
                    "bold":{"value":False},"italic":{"value":True},"color":{"value":"AABBCC"},"alignment":{"value":"right"},
                    "space_before_pt":{"value":3},"space_after_pt":{"value":7},"line_spacing":{"kind":"at_least","value":14,"unit":"pt"}}
        for action,cap in REGISTRY.items():
            if not cap.property.startswith("header:") or cap.property in {"header:content","header:link","header:first_page","header:odd_even"}: continue
            name=cap.property[7:]
            params=parameters.get(name,{"value":1.5,"unit":"chars"} if name.endswith("indent") else {"value":True})
            with self.subTest(action=action):
                doc,_=self.execute([self.op(action,params)])
                self.assertEqual([self.text(s.header) for s in doc.sections],["original"]*3)
                self.output.unlink()

    def test_link_previous_retains_other_sections(self):
        doc=Document(self.source)
        doc.sections[1].header.is_linked_to_previous=False
        doc.sections[1].header.paragraphs[0].text="middle"
        doc.save(self.source)
        doc,_=self.execute([self.op("header.link_previous.set",{"value":True})])
        self.assertTrue(doc.sections[1].header.is_linked_to_previous)
        self.assertEqual([self.text(s.header) for s in doc.sections],["original","original","middle"])

    def test_template_keeps_sections_variants_fields_and_prefix(self):
        doc,_=self.execute([
            self.op("header.content.set",{"mode":"chapter_title","style_name":"Heading 1","prefix":"Thesis: "}),
            self.op("header.content.set",{"mode":"text","text":"Thesis title","variant":"even"}),
        ])
        tokens=analyze_docx(self.output)["inferred_spec"]["headers_footers"]
        self.assertEqual(tokens["sections"][0]["header"]["text"],"original")
        self.assertEqual(tokens["sections"][1]["header"]["prefix"],"Thesis: ")
        self.assertEqual(tokens["sections"][1]["header"]["mode"],"chapter_title")
        self.assertEqual(tokens["sections"][1]["even_header"]["text"],"Thesis title")
        self.assertNotIn("original\nThesis title",tokens["header"]["text"])
        # Reapply only the inferred header contract; body structure is independent.
        self.output.unlink()
        doc,_=self.execute(spec={"headers_footers":tokens})
        self.assertEqual(self.text(doc.sections[0].header),"original")
        self.assertEqual(self.text(doc.sections[1].even_page_header),"Thesis title")

    def test_selected_web_section_and_conflicting_indents(self):
        doc,_=self.execute(spec={"page":{"header_distance_mm":17},"headers_footers":{"section_number":2,"preserve_existing":True,
            "header":{"enabled":True,"text":"selected"}}})
        self.assertEqual([self.text(s.header) for s in doc.sections],["original","selected","original"])
        self.assertAlmostEqual(doc.sections[1].header_distance.mm,17,places=1)
        self.assertEqual(doc.sections[0].header_distance,Document(self.source).sections[0].header_distance)
        self.output.unlink()
        with self.assertRaises(FormatError):
            build_plan(self.source,self.output,{"operations":[self.op("header.first_line_indent.set",{"value":2,"unit":"chars"}),
                                                               self.op("header.hanging_indent.set",{"value":1,"unit":"chars"})]})

    def test_template_field_suffix_and_extra_paragraph_are_preserved(self):
        for extra_paragraph in (False, True):
            with self.subTest(extra_paragraph=extra_paragraph):
                doc=Document(self.source)
                header=doc.sections[1].header
                header.is_linked_to_previous=False
                p=header.paragraphs[0]
                p.text="Thesis: "
                field=OxmlElement("w:fldSimple")
                field.set(qn("w:instr"),'STYLEREF "Heading 1"')
                run=OxmlElement("w:r"); text=OxmlElement("w:t"); text.text="cached heading"
                run.append(text); field.append(run); p._p.append(field)
                if extra_paragraph:
                    header.add_paragraph("Second line")
                else:
                    p.add_run(" — suffix")
                doc.save(self.source)
                before=copy.deepcopy(header._element)
                analysis=analyze_docx(self.source)
                tokens=analysis["inferred_spec"]["headers_footers"]
                self.assertTrue(any("unsupported STYLEREF" in item for item in analysis["unsupported"]))
                self.assertNotIn("mode",tokens["sections"][1]["header"])
                result,_=self.execute(spec={"headers_footers":tokens})
                self.assertEqual(result.sections[1].header._element.xml,before.xml)
                self.output.unlink()

    def test_continuous_section_becomes_next_page_after_explicit_break(self):
        doc=Document(self.source)
        doc.sections[0].start_type=WD_SECTION_START.CONTINUOUS
        doc.save(self.source)
        doc,_=self.execute([{"action":"section.break.next_page.insert","target":{"type":"paragraph","id":"p1"},"params":{}}])
        self.assertEqual(doc.sections[0].start_type,WD_SECTION_START.CONTINUOUS)
        self.assertEqual(doc.sections[1].start_type,WD_SECTION_START.NEW_PAGE)

    def test_wrong_font_writer_fails_independent_readback(self):
        cap=REGISTRY["header.font.size.set"]
        plan=build_plan(self.source,self.output,{"operations":[self.op(cap.action,{"value":9})]})
        with patch.dict(REGISTRY,{cap.action:type(cap)(**{**cap.__dict__,"handler":lambda *args:None})}):
            with self.assertRaises(FormatError): apply_plan(plan)
        self.assertFalse(self.output.exists())

    def test_requirements_parser_and_web_paths_have_exact_handlers(self):
        parsed=parse_requirements('奇偶页不同。奇数页页眉显示章节标题，宋体小五号，右对齐。偶数页页眉文字：论文题目。偶数页页眉宋体五号，左对齐。')["spec_patch"]
        self.assertTrue(parsed["headers_footers"]["different_odd_even"])
        self.assertEqual(parsed["headers_footers"]["header"]["mode"],"chapter_title")
        self.assertEqual(parsed["headers_footers"]["header"]["font_size_pt"],9)
        self.assertEqual(parsed["headers_footers"]["even_header"]["font_size_pt"],10.5)
        self.assertFalse(parsed.get("body"))
        paths=legacy_paths()
        self.assertEqual(paths["headers_footers.header.space_after_pt"],"header.spacing_after.set")
        self.assertEqual(paths["headers_footers.even_header.font_size_pt"],"header.font.size.set")
        caps=[c for c in REGISTRY.values() if c.property.startswith("header:")]
        self.assertEqual(len(caps),len({c.handler for c in caps}))

    def test_header_only_spec_font_edit_keeps_content(self):
        doc,_=self.execute(spec={"headers_footers":{"section_number":2,"header":{"font_size_pt":9}}})
        self.assertEqual([self.text(s.header) for s in doc.sections],["original"]*3)
        self.assertTrue(doc.sections[1].header.paragraphs[0].runs[0].bold)

    def test_literal_header_preserves_breaks_and_tabs(self):
        value="University\nThesis\tTitle"
        doc,_=self.execute([self.op("header.content.set",{"mode":"text","text":value}),
                            self.op("header.font.size.set",{"value":9})])
        self.assertEqual(doc.sections[1].header.paragraphs[0].text,value)
        self.assertEqual(self.text(doc.sections[2].header),"original")


if __name__ == "__main__": unittest.main()
