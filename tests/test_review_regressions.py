"""Regression cases from the project review; all documents are synthetic."""

import copy
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from xml.etree import ElementTree
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from render_docx import render_document
from word_format import apply_plan, build_plan, verify_plan
from word_format.contracts import FormatError, sha256
from verify_output import validate_visual_report


class ReviewRegressionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.source = self.folder / "论文.docx"
        self.output = self.folder / "结果.docx"
        document = Document()
        document.add_heading("第一章 绪论", level=1)
        document.add_paragraph("正文一").runs[0].font.size = Pt(10)
        document.add_paragraph("正文二").runs[0].font.size = Pt(12)
        document.save(self.source)
        self.request = {"operations": [{
            "action": "font.size.set",
            "target": {"type": "paragraph", "role": "body", "all": True},
            "params": {"value": 12},
        }]}

    def applied(self):
        plan = build_plan(self.source, self.output, self.request)
        return plan, apply_plan(plan)

    def test_verify_rejects_missing_empty_and_truncated_checks(self):
        plan, report = self.applied()
        for mode in ("missing", "empty", "truncated", "duplicate"):
            with self.subTest(mode=mode):
                partial = copy.deepcopy(report)
                if mode == "missing":
                    partial.pop("checks")
                elif mode == "empty":
                    partial["checks"] = []
                elif mode == "truncated":
                    partial["checks"] = partial["checks"][:1]
                else:
                    partial["checks"][1] = copy.deepcopy(partial["checks"][0])
                with self.assertRaises(FormatError):
                    verify_plan(plan, partial)

    def test_verify_rejects_changed_check_contract(self):
        plan, report = self.applied()
        for field, value in (("action", "font.bold.set"), ("params", {"value": 10}),
                             ("target", {"id": "p0"}), ("output_target", None),
                             ("output_target.id", "p0")):
            with self.subTest(field=field):
                changed = copy.deepcopy(report)
                if field == "output_target.id":
                    changed["checks"][0]["output_target"]["id"] = value
                else:
                    changed["checks"][0][field] = value
                with self.assertRaises(FormatError):
                    verify_plan(plan, changed)

    def test_verify_supports_reordered_checks_and_remapped_toc_targets(self):
        self.request["operations"].insert(0, {
            "action": "toc.configure", "target": {"type": "document"},
            "params": {"enabled": True},
        })
        plan, report = self.applied()
        report["checks"].reverse()
        self.assertEqual(verify_plan(plan, report)["status"], "passed")

    def test_bad_target_and_action_return_format_errors(self):
        for target in (None, [], 4, {"type": []}):
            with self.subTest(target=target):
                request = copy.deepcopy(self.request)
                request["operations"][0]["target"] = target
                with self.assertRaises(FormatError):
                    build_plan(self.source, self.output, request)
        request = copy.deepcopy(self.request)
        request["operations"][0]["action"] = []
        with self.assertRaises(FormatError):
            build_plan(self.source, self.output, request)

    def test_visual_review_requires_each_page_as_an_integer_once(self):
        report = {"status": "passed", "model": {"supports_image_input": True}}
        for pages in ([True], [1.5], ["1"], [1, 1], []):
            with self.subTest(pages=pages):
                self.assertEqual(validate_visual_report({**report, "reviewed_pages": pages}, 1)["status"], "failed")
        self.assertEqual(validate_visual_report({**report, "reviewed_pages": [2, 1]}, 2)["status"], "passed")

    def test_verify_rechecks_structure_independently_of_saved_status(self):
        plan, report = self.applied()
        with patch("verify_output.verify_structure", return_value={"status": "failed", "errors": ["broken structure"], "checks": []}):
            result = verify_plan(plan, report)
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["structure"]["errors"], ["broken structure"])

    def test_audit_reports_each_effective_value_without_writing(self):
        from word_format import audit_document
        before = sha256(self.source)
        with patch("word_format.api.apply_plan", side_effect=AssertionError("audit wrote a document")):
            report = audit_document(self.source, self.request)
        self.assertEqual(report["status"], "failed")
        self.assertEqual(report["summary"], {"checked": 2, "passed": 1, "failed": 1})
        self.assertEqual([c["target"]["id"] for c in report["checks"]], ["p1", "p2"])
        self.assertEqual(sha256(self.source), before)
        self.assertFalse(self.output.exists())
        self.assertEqual(sorted(p.name for p in self.folder.iterdir()), [self.source.name])

    def test_audit_accepts_partial_spec_and_cli_protects_inputs(self):
        from word_format import audit_document
        report = audit_document(self.source, {"spec": {"body": {"font_size_pt": 12}}})
        self.assertEqual(report["status"], "failed")
        self.assertGreater(report["summary"]["failed"], 0)
        request_file = self.folder / "request.json"
        request_file.write_text(json.dumps(self.request), encoding="utf-8")
        before = sha256(self.source)
        command = [sys.executable, str(ROOT / "scripts" / "format_cli.py"), "audit",
                   str(self.source), "--request", str(request_file)]
        result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", timeout=30)
        self.assertEqual(result.returncode, 2, result.stderr)
        self.assertEqual(json.loads(result.stdout)["summary"]["failed"], 1)
        for protected in (self.source, request_file):
            result = subprocess.run(command + ["--output", str(protected)], capture_output=True,
                                    text=True, encoding="utf-8", timeout=30)
            self.assertEqual(result.returncode, 2)
            self.assertIn("must not overwrite", result.stderr)
        self.assertEqual(sha256(self.source), before)

    def runtime(self):
        return {"preferred": "word", "word": {"available": True, "name": "Microsoft Word"}}

    @staticmethod
    def fake_pages(pdf, folder, dpi):
        page = folder / "page-1.png"
        page.write_bytes(b"synthetic-page-evidence")
        return [page]

    def test_render_uses_snapshot_and_hashes_published_evidence(self):
        before = sha256(self.source)
        rendered_sources = []
        def convert(source, output):
            rendered_sources.append(source)
            self.assertNotEqual(source, self.source)
            self.assertEqual(sha256(source), before)
            output.write_bytes(b"synthetic-pdf")
        with patch("render_docx.detect_renderers", return_value=self.runtime()), \
             patch("render_docx._word_to_pdf", side_effect=convert), \
             patch("render_docx._pdf_to_pngs", side_effect=self.fake_pages):
            report = render_document(self.source, self.folder / "render", emit_pdf=True)
        manifest = json.loads(Path(report["manifest"]).read_text(encoding="utf-8"))
        self.assertEqual(manifest["source_sha256"], before)
        self.assertEqual(Path(manifest["source"]), self.source)
        self.assertEqual(manifest["pages"][0]["sha256"], sha256(report["pages"][0]))
        self.assertTrue(Path(report["pdf"]).is_file())
        self.assertFalse(rendered_sources[0].exists())

    def test_render_rejects_source_changed_during_conversion(self):
        folder = self.folder / "render"
        folder.mkdir()
        manifest = folder / "render-manifest.json"
        manifest.write_text('{"stale": true}', encoding="utf-8")
        def convert(source, output):
            document = Document(self.source)
            document.add_paragraph("concurrent edit")
            document.save(self.source)
            output.write_bytes(b"synthetic-pdf")
        with patch("render_docx.detect_renderers", return_value=self.runtime()), \
             patch("render_docx._word_to_pdf", side_effect=convert), \
             patch("render_docx._pdf_to_pngs", side_effect=self.fake_pages):
            with self.assertRaisesRegex(RuntimeError, "changed|变化"):
                render_document(self.source, folder)
        self.assertFalse(manifest.exists())
        self.assertFalse((folder / "page-1.png").exists())

    def test_failed_render_invalidates_manifest_and_keeps_previous_pages(self):
        folder = self.folder / "render"
        folder.mkdir()
        manifest = folder / "render-manifest.json"
        manifest.write_text('{"stale": true}', encoding="utf-8")
        page = folder / "page-1.png"
        page.write_bytes(b"previous-page")
        with patch("render_docx.detect_renderers", return_value=self.runtime()), \
             patch("render_docx._word_to_pdf", side_effect=RuntimeError("conversion failed")):
            with self.assertRaises(RuntimeError):
                render_document(self.source, folder)
        self.assertFalse(manifest.exists())
        self.assertEqual(page.read_bytes(), b"previous-page")

    def test_render_rejects_relative_external_links_before_staging(self):
        with ZipFile(self.source) as package:
            parts = {name: package.read(name) for name in package.namelist()}
        rel_path = "word/_rels/document.xml.rels"
        root = ElementTree.fromstring(parts[rel_path])
        ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
        ElementTree.SubElement(root, ns + "Relationship", {
            "Id": "rId900", "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
            "Target": "assets/figure.png", "TargetMode": "External",
        })
        parts[rel_path] = ElementTree.tostring(root, encoding="utf-8")
        with ZipFile(self.source, "w", ZIP_DEFLATED) as package:
            for name, data in parts.items():
                package.writestr(name, data)
        with patch("render_docx.detect_renderers", side_effect=AssertionError("unexpected renderer detection")):
            with self.assertRaisesRegex(ValueError, "relative external links"):
                render_document(self.source, self.folder / "render")
        self.assertFalse((self.folder / "render" / "render-manifest.json").exists())


if __name__ == "__main__":
    unittest.main()
