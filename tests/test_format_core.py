"""Regression contracts for both entrances and their shared format writers."""

import copy
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"scripts"))
from word_format import apply_plan, build_plan, inspect_document, verify_plan
from word_format.contracts import FormatError, sha256
from word_format.registry import REGISTRY
from apply_spec import apply_document
from parse_requirements import parse_requirements
from verify_output import verify_structure


class FormatCoreTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder=Path(self.temp.name)
        self.source=self.folder/"source.docx"
        self.output=self.folder/"output.docx"
        doc=Document()
        doc.add_paragraph("Chapter One",style="Heading 1")
        p=doc.add_paragraph("alpha beta gamma")
        p.runs[0].font.size=Pt(20)
        p.runs[0].bold=True
        p.paragraph_format.space_after=Pt(3)
        doc.add_paragraph("alpha beta gamma")
        doc.add_paragraph("Chapter Two",style="Heading 1")
        doc.add_paragraph("second chapter body")
        table=doc.add_table(rows=2,cols=2)
        table.cell(0,0).text="cell"
        doc.save(self.source)
        self.original_hash=sha256(self.source)

    def operation(self,action,params,target=None):
        return {"action":action,"target":target or {"type":"paragraph","id":"p1"},"params":params}

    def execute(self,operations):
        plan=build_plan(self.source,self.output,{"operations":operations})
        report=apply_plan(plan)
        self.assertEqual(verify_plan(plan,report)["status"],"passed")
        self.assertEqual(sha256(self.source),self.original_hash)
        return Document(self.output),report

    def test_all_atomic_writers_round_trip_and_preserve_package(self):
        parameters={"font_east_asia":{"value":"SimSun"},"font_latin":{"value":"Arial"},
            "font_size_pt":{"value":12,"unit":"pt"},"bold":{"value":False},"italic":{"value":True},"color":{"value":"12AB34"},
            "alignment":{"value":"right"},"space_before_pt":{"value":7,"unit":"pt"},"space_after_pt":{"value":11,"unit":"pt"},
            "line_spacing":{"kind":"at_least","value":18,"unit":"pt"},"style":{"name":"Local Custom"},
            "table_alignment":{"value":"center"},"repeat_header":{"value":True}}
        for action,cap in REGISTRY.items():
            if cap.structural: continue
            with self.subTest(action=action):
                if cap.property.startswith("page:"):
                    attr=cap.property[5:]
                    params={"value":"landscape"} if attr=="orient" else {"value":310 if attr=="h" else 220 if attr=="w" else 20,"unit":"mm"}
                    target={"type":"section","number":1}
                elif cap.property in {"table_alignment","repeat_header"}:
                    params=parameters[cap.property];target={"type":"table","id":"t0"}
                else:
                    params=parameters.get(cap.property,{"value":1.5,"unit":"chars"} if cap.property.endswith("indent") else {"value":False})
                    target={"type":"paragraph","id":"p1"}
                doc,report=self.execute([self.operation(action,params,target)])
                self.assertEqual([p.text for p in doc.paragraphs],["Chapter One","alpha beta gamma","alpha beta gamma","Chapter Two","second chapter body"])
                self.assertEqual(doc.paragraphs[2].text,"alpha beta gamma")
                self.assertEqual(report["verification"]["preservation"]["status"],"passed")
                self.output.unlink()

    def test_local_font_change_preserves_spacing_bold_and_other_paragraph(self):
        doc,_=self.execute([self.operation("font.size.set",{"value":12})])
        self.assertEqual(doc.paragraphs[1].runs[0].font.size.pt,12)
        self.assertTrue(doc.paragraphs[1].runs[0].bold)
        self.assertEqual(doc.paragraphs[1].paragraph_format.space_after.pt,3)
        self.assertIsNone(doc.paragraphs[2].runs[0].font.size)

    def test_scoped_cleanup_does_not_erase_table_size(self):
        report=apply_document(self.source,self.output,{"body":{"font_size_pt":12},"tables":{"font_size_pt":10.5}},clear_direct=True)
        self.assertEqual(report["status"],"passed")
        self.assertEqual(Document(self.output).tables[0].cell(0,0).paragraphs[0].runs[0].font.size.pt,10.5)

    def test_partial_spec_does_not_normalize_headings_or_add_indent(self):
        apply_document(self.source,self.output,{"page":{"margin_left_mm":26}})
        a=Document(self.source);b=Document(self.output)
        self.assertEqual(a.paragraphs[0]._p.xml,b.paragraphs[0]._p.xml)
        self.assertIsNone(b.paragraphs[1].paragraph_format.first_line_indent)

    def test_chapter_selector_and_ordinal(self):
        doc,_=self.execute([self.operation("font.size.set",{"value":14},{"type":"paragraph","chapter_text":"Chapter Two","role":"body","all":True})])
        self.assertEqual(doc.paragraphs[4].runs[0].font.size.pt,14)
        self.assertEqual(doc.paragraphs[1].runs[0].font.size.pt,20)

    def test_ambiguous_selection_and_conflicts(self):
        cases=[
            [self.operation("font.size.set",{"value":12},{"type":"paragraph","text":"alpha beta gamma"})],
            [self.operation("font.size.set",{"value":12}),self.operation("font.size.set",{"value":14})],
            [self.operation("font.size.set",{"value":12},{"type":"text_range","id":"p1","start":0,"end":8}),self.operation("font.size.set",{"value":14},{"type":"text_range","id":"p1","start":5,"end":10})],
            [self.operation("paragraph.first_line_indent.set",{"value":2,"unit":"chars"}),self.operation("paragraph.hanging_indent.set",{"value":2,"unit":"chars"})],
        ]
        for ops in cases:
            with self.subTest(ops=ops),self.assertRaises(FormatError): build_plan(self.source,self.output,{"operations":ops})
        self.assertFalse(self.output.exists())

    def test_range_across_runs_and_hyperlink_keeps_bookmarks(self):
        doc=Document(self.source);p=doc.paragraphs[1]
        p.runs[0].text="alpha "
        bookmark=OxmlElement("w:bookmarkStart");bookmark.set(qn("w:id"),"7");bookmark.set(qn("w:name"),"test");p._p.append(bookmark)
        link=OxmlElement("w:hyperlink");link.set(qn("w:anchor"),"test")
        run=OxmlElement("w:r");text=OxmlElement("w:t");text.text="beta";run.append(text);link.append(run);p._p.append(link)
        p.add_run(" gamma")
        end=OxmlElement("w:bookmarkEnd");end.set(qn("w:id"),"7");p._p.append(end)
        doc.save(self.source);self.original_hash=sha256(self.source)
        out,_=self.execute([self.operation("font.size.set",{"value":12},{"type":"text_range","id":"p1","quote":"ha beta g"})])
        self.assertEqual(len(list(out.paragraphs[1]._p.iter(qn("w:hyperlink")))),1)
        self.assertEqual(len(list(out.paragraphs[1]._p.iter(qn("w:bookmarkStart")))),1)
        objects=inspect_document(self.output)["objects"]
        self.assertEqual(next(p for p in objects if p["id"]=="p1")["text"],"alpha beta gamma")

    def test_range_rejects_fields(self):
        doc=Document(self.source)
        node=OxmlElement("w:fldChar");node.set(qn("w:fldCharType"),"begin");doc.paragraphs[1].add_run()._r.append(node);doc.save(self.source)
        with self.assertRaises(FormatError): build_plan(self.source,self.output,{"operations":[self.operation("font.bold.set",{"value":True},{"type":"text_range","id":"p1","quote":"alpha"})]})

    def test_source_and_plan_fingerprints(self):
        plan=build_plan(self.source,self.output,{"operations":[self.operation("font.size.set",{"value":12})]})
        changed=copy.deepcopy(plan);changed["operations"][0]["params"]["value"]=14
        with self.assertRaises(FormatError): apply_plan(changed)
        doc=Document(self.source);doc.add_paragraph("changed");doc.save(self.source)
        with self.assertRaises(FormatError): apply_plan(plan)
        self.assertFalse(self.output.exists())

    def test_wrong_writer_cannot_publish(self):
        cap=REGISTRY["font.size.set"]
        from dataclasses import replace
        def wrong(index,ref,params): index.elements[ref["id"]].append(OxmlElement("w:bookmarkStart"))
        with patch.dict(REGISTRY,{cap.action:replace(cap,handler=wrong)}):
            with self.assertRaises(FormatError): self.execute([self.operation(cap.action,{"value":12})])
        self.assertFalse(self.output.exists())

    def test_no_toc_heading_means_no_output(self):
        doc=Document();doc.add_paragraph("body");doc.save(self.source)
        with self.assertRaises(FormatError): apply_document(self.source,self.output,{"table_of_contents":{"enabled":True}})
        self.assertFalse(self.output.exists())

    def test_repeat_header_can_be_disabled_and_borders_restored(self):
        doc,_=self.execute([self.operation("table.header_repeat.set",{"value":True},{"type":"table","id":"t0"}),self.operation("table.three_line.apply",{"value":True},{"type":"table","id":"t0"})])
        self.source=self.output;self.output=self.folder/"restored.docx";self.original_hash=sha256(self.source)
        out,_=self.execute([self.operation("table.header_repeat.set",{"value":False},{"type":"table","id":"t0"}),self.operation("table.three_line.apply",{"value":False},{"type":"table","id":"t0"})])
        self.assertIsNone(out.tables[0]._tbl.find("./"+qn("w:tr")+"/"+qn("w:trPr")+"/"+qn("w:tblHeader")))
        self.assertIsNone(out.tables[0]._tbl.find("./"+qn("w:tblPr")+"/"+qn("w:tblBorders")))

    def test_parser_does_not_leak_heading_spacing_or_invent_numbering(self):
        patch=parse_requirements("一级标题黑体，1.5倍行距。正文宋体，首行缩进0字符。")["spec_patch"]
        self.assertNotIn("line_spacing",patch.get("body",{}))
        self.assertEqual(patch["body"]["first_line_indent_chars"],0)
        self.assertNotIn("headings",patch.get("lists",{}))
        self.assertEqual(patch["headings"][0]["line_spacing"]["value"],1.5)

    def test_independent_verifier_detects_direct_override(self):
        result=verify_structure(self.source,{"body":{"font_size_pt":12}})
        self.assertEqual(result["status"],"failed")

    def test_module_does_not_change_shared_normal_style(self):
        doc=Document();doc.add_paragraph("摘要");doc.add_paragraph("abstract body");doc.add_paragraph("第一章 绪论");doc.add_paragraph("chapter body");doc.save(self.source)
        spec={"document_structure":{"ordered_modules":[{"kind":"abstract","style_roles":{"body":{"tokens":{"font_size_pt":10}}}},{"kind":"chapters"}]}}
        apply_document(self.source,self.output,spec)
        out=Document(self.output)
        self.assertEqual(out.paragraphs[1].runs[0].font.size.pt,10)
        self.assertIsNone(out.paragraphs[3].runs[0].font.size)
        self.assertEqual(doc.styles["Normal"].element.xml,out.styles["Normal"].element.xml)

    def test_direct_cli_and_preflight_report_paths(self):
        request=self.folder/"request.json";planfile=self.folder/"plan.json";reportfile=self.folder/"report.json"
        request.write_text(json.dumps({"operations":[self.operation("font.size.set",{"value":12})]}),encoding="utf-8")
        cli=[sys.executable,"-B",str(ROOT/"scripts/format_cli.py")]
        def run(*args): return subprocess.run(cli+list(args),capture_output=True,text=True,encoding="utf-8")
        result=run("plan",str(self.source),str(self.output),"--request",str(request),"--output",str(planfile));self.assertEqual(result.returncode,0,result.stderr)
        denied=run("apply","--plan",str(planfile),"--report",str(self.source));self.assertEqual(denied.returncode,2);self.assertFalse(self.output.exists())
        result=run("apply","--plan",str(planfile),"--report",str(reportfile));self.assertEqual(result.returncode,0,result.stderr)
        result=run("verify","--plan",str(planfile),"--report",str(reportfile));self.assertEqual(result.returncode,0,result.stderr)
        self.assertEqual(sha256(self.source),self.original_hash)

    def test_character_style_is_rejected(self):
        with self.assertRaises(FormatError): build_plan(self.source,self.output,{"operations":[self.operation("font.size.set",{"value":12},{"type":"style","name":"Emphasis"})]})


if __name__=="__main__": unittest.main()
